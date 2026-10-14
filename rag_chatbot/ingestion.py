"""Bounded, source-aware document extraction and deterministic token chunking.

The optional vision callback is supplied by the service layer. This module never
makes network calls and never changes an uploaded original.
"""

import csv
import hashlib
import io
import json
import os
import tempfile
import threading
import uuid
import zipfile
from pathlib import Path

import html2text
from docx import Document
from docx.text.paragraph import Paragraph
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from PIL import Image, ImageOps
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pypdf import PdfReader

from chatbot.log import logger
from chatbot.readers import EmptyDocumentError, read_text_file
from .schemas import Chunk, SourceElement

MAX_FILE_BYTES = 100 * 1024 * 1024
MAX_PAGES = 1000
MAX_IMAGE_FRAMES = 100
MAX_IMAGE_PIXELS = 40_000_000
MAX_SPREADSHEET_CELLS = 1_000_000
MAX_EXPANDED_OFFICE_BYTES = 512 * 1024 * 1024
MAX_EXTRACTED_CHARS = 20_000_000
MAX_IMAGE_EDGE = 2048
# PDFium is not thread-safe, including across separate document instances.
# Share this lock with the on-demand viewer; never hold it during vision calls.
PDFIUM_LOCK = threading.RLock()
TABLE_ROWS_PER_ELEMENT = 25
IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".webp", ".gif", ".bmp", ".tif", ".tiff"}
TEXT_EXTENSIONS = {".txt", ".md", ".markdown", ".rst"}
SUPPORTED_EXTENSIONS = IMAGE_EXTENSIONS | TEXT_EXTENSIONS | {
    ".pdf", ".docx", ".xlsx", ".xls", ".pptx", ".html", ".htm", ".csv", ".tsv",
}
_RASTER_CONTENT_TYPES = {
    "image/png", "image/jpeg", "image/jpg", "image/gif", "image/tiff",
    "image/bmp", "image/x-ms-bmp", "image/webp",
}


class IngestionError(ValueError):
    """A source cannot be completely and safely ingested with these options."""


def _check_office_archive(path):
    with zipfile.ZipFile(path) as archive:
        entries = archive.infolist()
        if len(entries) > 10_000 or sum(item.file_size for item in entries) > MAX_EXPANDED_OFFICE_BYTES:
            raise IngestionError("The Office document expands beyond the ingestion size limit.")


def _asset_directory(path):
    directory = path.parent / f".{path.name}.rag-assets"
    directory.mkdir(exist_ok=True)
    if directory.resolve().parent != path.parent:
        raise IngestionError("The document image directory must stay beside its original file.")
    return directory


def _save_png(image, path):
    """Normalize a derived image and commit it without leaving partial PNGs."""
    oriented = ImageOps.exif_transpose(image)
    try:
        if "A" in oriented.getbands() or "transparency" in oriented.info:
            converted = Image.new("RGB", oriented.size, "white")
            with oriented.convert("RGBA") as rgba:
                with rgba.getchannel("A") as alpha:
                    converted.paste(rgba, mask=alpha)
        else:
            converted = oriented.convert("RGB")
    finally:
        oriented.close()
    converted.thumbnail((MAX_IMAGE_EDGE, MAX_IMAGE_EDGE))
    temporary = None
    try:
        with tempfile.NamedTemporaryFile(dir=path.parent, suffix=".png", delete=False) as file:
            temporary = Path(file.name)
        converted.save(temporary, format="PNG")
        os.replace(temporary, path)
    finally:
        converted.close()
        if temporary is not None:
            temporary.unlink(missing_ok=True)


def _describe(vision_describe, image_path, locator):
    text = vision_describe(str(image_path))
    if not isinstance(text, str) or not text.strip():
        raise IngestionError(f"Vision extraction returned no usable text or description for {locator}.")
    return SourceElement(text.strip(), locator, "image", str(image_path))


def _embedded_image(blob, content_type, path, vision_describe, locator, cache):
    """Describe each distinct raster once while retaining every source location."""
    if content_type not in _RASTER_CONTENT_TYPES:
        logger.warning("Skipping unsupported embedded image format %s at %s", content_type, locator)
        return None
    digest = hashlib.sha256(blob).hexdigest()
    if digest not in cache:
        image_path = _asset_directory(path) / f"embedded-{digest}.png"
        with Image.open(io.BytesIO(blob)) as image:
            if image.width * image.height > MAX_IMAGE_PIXELS:
                raise IngestionError(f"Embedded image at {locator} exceeds the {MAX_IMAGE_PIXELS:,}-pixel limit.")
            _save_png(image, image_path)
        element = _describe(vision_describe, image_path, locator)
        cache[digest] = (element.text, element.image_path)
    text, image_path = cache[digest]
    return SourceElement(text, locator, "image", image_path)


def _paragraph_images(paragraph, path, vision_describe, locator, cache):
    for number, relationship_id in enumerate(paragraph._p.xpath(".//a:blip/@r:embed"), start=1):
        part = paragraph.part.related_parts[relationship_id]
        element = _embedded_image(
            part.blob, part.content_type, path, vision_describe,
            f"{locator} · image {number}", cache,
        )
        if element is not None:
            yield element


def _cell_images(cell, path, vision_describe, locator, cache):
    for number, block in enumerate(cell.iter_inner_content(), start=1):
        if isinstance(block, Paragraph):
            yield from _paragraph_images(block, path, vision_describe, f"{locator} · paragraph {number}", cache)
        else:
            for row_number, row in enumerate(block.rows, start=1):
                seen = set()
                for column, child in enumerate(row.cells, start=1):
                    if child._tc in seen:
                        continue
                    seen.add(child._tc)
                    child_locator = f"{locator} · nested table {number} · cell {get_column_letter(column)}{row_number}"
                    yield from _cell_images(child, path, vision_describe, child_locator, cache)


def _pdf_elements(path, vision_describe):
    reader = PdfReader(path)
    if reader.is_encrypted and not reader.decrypt(""):
        raise IngestionError("This PDF is password-protected. Upload an unlocked copy.")
    if len(reader.pages) > MAX_PAGES:
        raise IngestionError(f"PDFs are limited to {MAX_PAGES:,} pages; split this document first.")
    texts = []
    total_characters = 0
    for page in reader.pages:
        text = page.extract_text() or ""
        total_characters += len(text)
        if total_characters > MAX_EXTRACTED_CHARS:
            raise IngestionError("The extracted PDF text exceeds the ingestion size limit.")
        texts.append(text)
    missing = [index + 1 for index, text in enumerate(texts) if not text.strip()]
    if missing and vision_describe is None:
        locations = ", ".join(map(str, missing[:10])) + (", ..." if len(missing) > 10 else "")
        raise IngestionError(
            f"PDF page(s) {locations} have no extractable text. Enable image/OCR processing "
            "or OCR the PDF first; remove genuinely blank pages if needed. No partial index was created."
        )
    labels = reader.page_labels
    renderer = None
    try:
        if missing:
            import pypdfium2 as pdfium

            with PDFIUM_LOCK:
                renderer = pdfium.PdfDocument(str(path))
        for index, text in enumerate(texts):
            locator = f"Page {index + 1}"
            if labels[index] != str(index + 1):
                locator += f" (printed label {labels[index]})"
            if text.strip():
                yield SourceElement(text.strip(), locator)
                continue
            image_path = _asset_directory(path) / f"page-{index + 1:04d}.png"
            with PDFIUM_LOCK:
                page = renderer[index]
                bitmap = None
                try:
                    width, height = page.get_size()
                    if width <= 0 or height <= 0:
                        raise IngestionError(f"{locator} has invalid page dimensions.")
                    bitmap = page.render(scale=min(2.0, MAX_IMAGE_EDGE / max(width, height)))
                    image = bitmap.to_pil()
                    try:
                        _save_png(image, image_path)
                    finally:
                        image.close()
                finally:
                    if bitmap is not None:
                        bitmap.close()
                    page.close()
            yield _describe(vision_describe, image_path, locator)
    finally:
        if renderer is not None:
            with PDFIUM_LOCK:
                renderer.close()


def _image_elements(path, vision_describe):
    if vision_describe is None:
        raise IngestionError("Images require image/OCR processing. Enable it and retry this document.")
    with Image.open(path) as image:
        frames = getattr(image, "n_frames", 1)
        if frames > MAX_IMAGE_FRAMES:
            raise IngestionError(f"Images are limited to {MAX_IMAGE_FRAMES} frames; split this image first.")
        for index in range(frames):
            image.seek(index)
            if image.width * image.height > MAX_IMAGE_PIXELS:
                raise IngestionError(f"Image frame {index + 1} exceeds the {MAX_IMAGE_PIXELS:,}-pixel limit.")
            image_path = _asset_directory(path) / f"image-{index + 1:04d}.png"
            _save_png(image, image_path)
            locator = "Image" if frames == 1 else f"Frame {index + 1}"
            yield _describe(vision_describe, image_path, locator)


def _cell_text(cell):
    parts = []
    for block in cell.iter_inner_content():
        if isinstance(block, Paragraph):
            parts.append(block.text)
        else:
            parts.extend(" | ".join(_cell_text(child) for child in row.cells) for row in block.rows)
    return "\n".join(part for part in parts if part.strip())


def _word_elements(path, vision_describe):
    headings = []
    image_cache = {}
    paragraph_number = table_number = 0
    for block in Document(path).iter_inner_content():
        if isinstance(block, Paragraph):
            paragraph_number += 1
            text = block.text.strip()
            style = block.style.name if block.style is not None else ""
            if style.startswith("Heading ") and style.removeprefix("Heading ").isdigit():
                level = int(style.removeprefix("Heading "))
                while headings and headings[-1][0] >= level:
                    headings.pop()
                headings.append((level, text))
            section = " / ".join(title for _, title in headings)
            locator = f"Section {section} · paragraph {paragraph_number}" if section else f"Paragraph {paragraph_number}"
            if text:
                content = f"Section: {section}\n\n{text}" if section and text != headings[-1][1] else text
                yield SourceElement(content, locator)
            if vision_describe is not None:
                yield from _paragraph_images(block, path, vision_describe, locator, image_cache)
        else:
            table_number += 1
            section = " / ".join(title for _, title in headings)
            first_row = " | ".join(_cell_text(cell) for cell in block.rows[0].cells) if block.rows else ""
            for index, row in enumerate(block.rows, start=1):
                values = [_cell_text(cell) for cell in row.cells]
                text = " | ".join(f"{get_column_letter(column)}: {value}" for column, value in enumerate(values, start=1))
                locator = f"Table {table_number} · row {index}"
                if section:
                    locator = f"Section {section} · {locator}"
                if any(value.strip() for value in values):
                    if index > 1 and first_row.strip():
                        text = f"First table row: {first_row}\n\n{text}"
                    if section:
                        text = f"Section: {section}\n\n{text}"
                    yield SourceElement(text, locator, "table")
                if vision_describe is not None:
                    seen = set()
                    for column, cell in enumerate(row.cells, start=1):
                        if cell._tc in seen:
                            continue
                        seen.add(cell._tc)
                        yield from _cell_images(
                            cell, path, vision_describe,
                            f"{locator} · cell {get_column_letter(column)}", image_cache,
                        )


def _table_elements(rows, sheet_name):
    """Group rows while keeping original cell addresses in the evidence text."""
    group = []
    first_row = last_row = None
    context = None
    total_cells = 0
    for row_number, values in rows:
        total_cells += len(values)
        if total_cells > MAX_SPREADSHEET_CELLS:
            raise IngestionError(f"Tables are limited to {MAX_SPREADSHEET_CELLS:,} cells.")
        cells = [f"{get_column_letter(column)}{row_number}: {value}" for column, value in enumerate(values, start=1) if value is not None and str(value).strip()]
        if not cells:
            continue
        text = " | ".join(cells)
        if context is None:
            context = text
        if first_row is None:
            first_row = row_number
        last_row = row_number
        group.append(text)
        if len(group) == TABLE_ROWS_PER_ELEMENT:
            yield SourceElement(f"Sheet: {sheet_name}\nFirst populated row: {context}\n\n" + "\n".join(group), f"Sheet {sheet_name} · rows {first_row}-{last_row}", "table")
            group, first_row = [], None
    if group:
        yield SourceElement(f"Sheet: {sheet_name}\nFirst populated row: {context}\n\n" + "\n".join(group), f"Sheet {sheet_name} · rows {first_row}-{last_row}", "table")


def _excel_elements(path):
    workbook = load_workbook(path, read_only=True, data_only=False, keep_links=False)
    cached = None
    try:
        if sum((sheet.max_row or 0) * (sheet.max_column or 0) for sheet in workbook) > MAX_SPREADSHEET_CELLS:
            raise IngestionError(f"Workbooks are limited to {MAX_SPREADSHEET_CELLS:,} cells, including their used ranges.")
        cached = load_workbook(path, read_only=True, data_only=True, keep_links=False)
        processed_cells = 0
        for sheet in workbook:
            def rows():
                nonlocal processed_cells
                for number, (formulas, values) in enumerate(zip(sheet.iter_rows(), cached[sheet.title].iter_rows()), start=1):
                    processed_cells += len(formulas)
                    if processed_cells > MAX_SPREADSHEET_CELLS:
                        raise IngestionError(f"Workbooks are limited to {MAX_SPREADSHEET_CELLS:,} cells.")
                    row = []
                    for cell, value in zip(formulas, values):
                        if cell.data_type == "f":
                            cached_value = f"cached value: {value.value}" if value.value is not None else "cached value unavailable"
                            row.append(f"{cell.value} ({cached_value})")
                        else:
                            row.append(cell.value)
                    yield number, row
            yield from _table_elements(rows(), sheet.title)
    finally:
        workbook.close()
        if cached is not None:
            cached.close()


def _legacy_excel_elements(path):
    import xlrd

    workbook = xlrd.open_workbook(path, on_demand=True)
    try:
        cells = 0
        for sheet_name in workbook.sheet_names():
            sheet = workbook.sheet_by_name(sheet_name)
            cells += sheet.nrows * sheet.ncols
            if cells > MAX_SPREADSHEET_CELLS:
                raise IngestionError(f"Workbooks are limited to {MAX_SPREADSHEET_CELLS:,} cells.")
            def rows():
                for index in range(sheet.nrows):
                    values = []
                    for cell in sheet.row(index):
                        if cell.ctype == xlrd.XL_CELL_DATE:
                            values.append(xlrd.xldate_as_datetime(cell.value, workbook.datemode).isoformat())
                        else:
                            values.append(cell.value)
                    yield index + 1, values
            yield from _table_elements(rows(), sheet.name)
            workbook.unload_sheet(sheet_name)
    finally:
        workbook.release_resources()


def _slide_elements(shapes, locator, title, path, vision_describe, image_cache):
    for shape in shapes:
        shape_locator = f"{locator} · shape {shape.shape_id}"
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _slide_elements(shape.shapes, shape_locator, title, path, vision_describe, image_cache)
        elif shape.shape_type == MSO_SHAPE_TYPE.PICTURE and vision_describe is not None:
            element = _embedded_image(
                shape.image.blob, shape.image.content_type, path, vision_describe,
                f"{shape_locator} · image", image_cache,
            )
            if element is not None:
                yield element
        elif shape.has_table:
            first_row = " | ".join(cell.text for cell in shape.table.rows[0].cells) if len(shape.table.rows) else ""
            for number, row in enumerate(shape.table.rows, start=1):
                text = " | ".join(f"{get_column_letter(column)}: {cell.text}" for column, cell in enumerate(row.cells, start=1))
                if any(cell.text.strip() for cell in row.cells):
                    if number > 1 and first_row.strip():
                        text = f"First table row: {first_row}\n\n{text}"
                    if title:
                        text = f"Slide title: {title}\n\n{text}"
                    yield SourceElement(text, f"{shape_locator} · table row {number}", "table")
        elif shape.has_text_frame and shape.text.strip():
            text = shape.text.strip()
            if title and title != text:
                text = f"Slide title: {title}\n\n{text}"
            yield SourceElement(text, shape_locator)


def _powerpoint_elements(path, vision_describe):
    presentation = Presentation(path)
    image_cache = {}
    if len(presentation.slides) > MAX_PAGES:
        raise IngestionError(f"Presentations are limited to {MAX_PAGES:,} slides.")
    for number, slide in enumerate(presentation.slides, start=1):
        locator = f"Slide {number}"
        title = slide.shapes.title.text.strip() if slide.shapes.title is not None else ""
        yield from _slide_elements(slide.shapes, locator, title, path, vision_describe, image_cache)
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame
            if notes is not None and notes.text.strip():
                yield SourceElement(notes.text.strip(), f"{locator} · speaker notes", "notes")


def extract_elements(path, *, vision_describe=None):
    """Extract complete source elements, raising if required OCR is unavailable.

    ``vision_describe`` receives a persistent PNG filename and returns extracted
    text plus a useful image description. Its caller owns billing and retries.
    """
    path = Path(path).resolve(strict=True)
    if not path.is_file():
        raise IngestionError("Select a file to ingest.")
    if path.stat().st_size > MAX_FILE_BYTES:
        raise IngestionError(f"Files are limited to {MAX_FILE_BYTES // (1024 * 1024)} MiB; split this source first.")
    extension = path.suffix.lower()
    if extension not in SUPPORTED_EXTENSIONS:
        raise IngestionError(f"Unsupported document type {extension or '(no extension)'}. Convert legacy Word/PowerPoint files to DOCX/PPTX.")
    if extension in {".docx", ".xlsx", ".pptx"}:
        _check_office_archive(path)
    if extension == ".pdf":
        extracted = _pdf_elements(path, vision_describe)
    elif extension in IMAGE_EXTENSIONS:
        extracted = _image_elements(path, vision_describe)
    elif extension == ".docx":
        extracted = _word_elements(path, vision_describe)
    elif extension == ".xlsx":
        extracted = _excel_elements(path)
    elif extension == ".xls":
        extracted = _legacy_excel_elements(path)
    elif extension == ".pptx":
        extracted = _powerpoint_elements(path, vision_describe)
    elif extension in {".csv", ".tsv"}:
        rows = csv.reader(io.StringIO(read_text_file(path)), delimiter="\t" if extension == ".tsv" else ",")
        extracted = _table_elements(enumerate(rows, start=1), path.name)
    else:
        text = read_text_file(path)
        if extension in {".html", ".htm"}:
            text = html2text.html2text(text)
        extracted = [SourceElement(text, "Document text")]
    elements = []
    total_characters = 0
    for element in extracted:
        if not element.text.strip():
            continue
        total_characters += len(element.text)
        if total_characters > MAX_EXTRACTED_CHARS:
            raise IngestionError("Extracted text exceeds the ingestion size limit; split this source first.")
        elements.append(element)
    if not elements:
        raise EmptyDocumentError("No usable text or image description could be extracted from this document.")
    logger.info("Extracted %d source elements from %s", len(elements), path.name)
    return elements


def chunk_elements(elements, document_id, chunk_size=600, chunk_overlap=80):
    """Split within source boundaries; stable UUIDs make repeated upserts safe."""
    from langchain_text_splitters import RecursiveCharacterTextSplitter

    if not isinstance(chunk_size, int) or chunk_size <= 0:
        raise ValueError("chunk_size must be a positive token count.")
    if not isinstance(chunk_overlap, int) or not 0 <= chunk_overlap < chunk_size:
        raise ValueError("chunk_overlap must be non-negative and smaller than chunk_size.")
    splitter = RecursiveCharacterTextSplitter.from_tiktoken_encoder(
        encoding_name="cl100k_base", chunk_size=chunk_size, chunk_overlap=chunk_overlap,
        allowed_special=set(), disallowed_special=(),
    )
    chunks = []
    for element_number, element in enumerate(elements):
        for part_number, text in enumerate(splitter.split_text(element.text)):
            if not text.strip():
                continue
            identity = json.dumps([
                "rag-chunk-v1", document_id, chunk_size, chunk_overlap,
                element_number, part_number, element.locator, element.kind, text,
            ], ensure_ascii=False, separators=(",", ":"))
            chunks.append(Chunk(
                str(uuid.uuid5(uuid.NAMESPACE_URL, identity)), str(document_id), text,
                element.locator, element.kind, element.image_path,
            ))
    return chunks
