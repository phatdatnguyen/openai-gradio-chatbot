"""Bounded local document previews. Originals are never modified or executed."""

from dataclasses import dataclass, field
import base64
import csv
from html import escape
from html.parser import HTMLParser
import io
import math
from pathlib import Path
import struct
import warnings
import zipfile


MAX_FILE_BYTES = 100 * 1024 * 1024
MAX_EXPANDED_OFFICE_BYTES = 64 * 1024 * 1024
MAX_ARCHIVE_ENTRIES = 10_000
MAX_PAGES = 1000
MAX_IMAGE_FRAMES = 100
MAX_IMAGE_PIXELS = 40_000_000
MAX_IMAGE_EDGE = 1800
MAX_EMBEDDED_BYTES = 8 * 1024 * 1024
MAX_EMBEDDED_IMAGES = 6
MAX_TEXT_BYTES = 2 * 1024 * 1024
MAX_TEXT_CHARS = 200_000
TEXT_CHARS_PER_PAGE = 12_000
ROWS_PER_PAGE = 50
MAX_SHEET_ROWS = 10_000
MAX_SHEET_COLUMNS = 50
MAX_SHEETS = 256
BLOCKS_PER_PAGE = 20
MAX_CELL_CHARS = 1000
IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".webp", ".gif", ".bmp", ".tif", ".tiff"}
TEXT_EXTENSIONS = {".txt", ".md", ".markdown", ".rst", ".json", ".xml", ".log"}
OFFICE_EXTENSIONS = {".doc", ".docx", ".ppt", ".pptx"}


@dataclass
class Preview:
    image: object | None = None
    html: str = ""
    table: list[list[object]] = field(default_factory=list)
    headers: list[str] = field(default_factory=list)
    text: str = ""
    unit_label: str = "Page"
    unit_count: int = 1
    unit: int = 1
    sheets: list[str] = field(default_factory=list)
    sheet: str | None = None
    status: str = ""


class PreviewError(ValueError):
    """A file cannot be safely previewed with the current limits."""


def _source_path(path, data_dir):
    root = Path(data_dir).resolve()
    sources = root / "sources"
    if sources.resolve() != sources or not sources.is_dir():
        raise PreviewError("The workspace sources directory is unavailable or unsafe.")
    candidate = Path(path)
    candidate = (root / candidate).resolve() if not candidate.is_absolute() else candidate.resolve()
    if not candidate.is_relative_to(sources):
        raise PreviewError("The preview must belong to this workspace's sources directory.")
    if not candidate.is_file():
        raise PreviewError("The original file is unavailable. It may have been removed.")
    if candidate.stat().st_size > MAX_FILE_BYTES:
        raise PreviewError("Preview files are limited to 100 MiB. Download the original to review it.")
    return candidate


def _unit(value, count):
    try:
        number = int(value if value is not None else 1)
    except (ValueError, TypeError, OverflowError) as exc:
        raise PreviewError("Enter a valid page number.") from exc
    return min(max(number, 1), max(count, 1))


def _archive_limit(path):
    with zipfile.ZipFile(path) as archive:
        entries = archive.infolist()
        if len(entries) > MAX_ARCHIVE_ENTRIES or sum(item.file_size for item in entries) > MAX_EXPANDED_OFFICE_BYTES:
            raise PreviewError("The expanded Office file exceeds the preview limit. Download the original to review it.")
        if any(item.flag_bits & 1 for item in entries):
            raise PreviewError("Password-protected Office files cannot be previewed.")


def _pdf_preview(path, unit):
    import pypdfium2 as pdfium
    from .ingestion import PDFIUM_LOCK

    # PDFium is not thread-safe; extraction shares this same lock. Copy the
    # pixels before releasing native resources, so Gradio owns an ordinary PIL image.
    with PDFIUM_LOCK:
        try:
            with pdfium.PdfDocument(str(path)) as document:
                count = len(document)
                if not 1 <= count <= MAX_PAGES:
                    raise PreviewError(f"PDF previews support 1 to {MAX_PAGES:,} pages.")
                number = _unit(unit, count)
                page = document[number - 1]
                bitmap = None
                try:
                    width, height = page.get_size()
                    if not all(math.isfinite(size) and 0 < size <= 1_000_000 for size in (width, height)):
                        raise PreviewError("The PDF page has unsupported dimensions.")
                    scale = min(2.0, MAX_IMAGE_EDGE / max(width, height))
                    bitmap = page.render(scale=scale)
                    rendered = bitmap.to_pil()
                    try:
                        image = rendered.convert("RGB")
                    finally:
                        rendered.close()
                finally:
                    if bitmap is not None:
                        bitmap.close()
                    page.close()
                return Preview(image=image, unit=number, unit_count=count,
                               status=f"Page {number} of {count}.")
        except pdfium.PdfiumError as exc:
            raise PreviewError("This PDF could not be rendered. It may be damaged or password-protected.") from exc


def _normalized_image(image, edge=MAX_IMAGE_EDGE):
    from PIL import Image, ImageOps

    if image.width * image.height > MAX_IMAGE_PIXELS:
        raise PreviewError(f"Images are limited to {MAX_IMAGE_PIXELS:,} pixels for preview.")
    with ImageOps.exif_transpose(image) as oriented:
        with oriented.convert("RGBA") as rgba:
            result = Image.new("RGB", rgba.size, "white")
            result.paste(rgba, mask=rgba.getchannel("A"))
    result.thumbnail((edge, edge), Image.Resampling.LANCZOS)
    return result


def _image_preview(path, unit):
    from PIL import Image

    with warnings.catch_warnings():
        warnings.simplefilter("error", Image.DecompressionBombWarning)
        try:
            with Image.open(path) as original:
                count = getattr(original, "n_frames", 1)
                if count > MAX_IMAGE_FRAMES:
                    raise PreviewError(f"Image previews support at most {MAX_IMAGE_FRAMES} frames.")
                number = _unit(unit, count)
                original.seek(number - 1)
                image = _normalized_image(original)
        except (Image.DecompressionBombError, Image.DecompressionBombWarning) as exc:
            raise PreviewError("The image dimensions exceed the preview limit.") from exc
    return Preview(image=image, unit=number, unit_count=count, unit_label="Frame",
                   status=f"Frame {number} of {count}." if count > 1 else "Image preview.")


def _cell(value):
    if value is None:
        return ""
    text = str(value)
    return text if len(text) <= MAX_CELL_CHARS else text[:MAX_CELL_CHARS] + "…"


def _column(index):
    letters = ""
    while index:
        index, remainder = divmod(index - 1, 26)
        letters = chr(65 + remainder) + letters
    return letters


def _table_preview(rows, row_count, column_count, unit, *, sheets=None, sheet=None):
    count = max(1, math.ceil(min(row_count, MAX_SHEET_ROWS) / ROWS_PER_PAGE))
    number = _unit(unit, count)
    first = (number - 1) * ROWS_PER_PAGE + 1
    width = min(column_count, MAX_SHEET_COLUMNS)
    values = [[first + index] + [_cell(value) for value in list(row)[:width]] for index, row in enumerate(rows)]
    for row in values:
        row.extend([""] * (width + 1 - len(row)))
    last = first + len(values) - 1
    status = f"Rows {first}–{last} of {row_count:,}." if values else "This sheet is empty."
    if row_count > MAX_SHEET_ROWS or column_count > MAX_SHEET_COLUMNS:
        status += f" Preview limited to {MAX_SHEET_ROWS:,} rows and {MAX_SHEET_COLUMNS} columns; download the original for the rest."
    return Preview(table=values, headers=["Row"] + [_column(index) for index in range(1, width + 1)],
                   unit_label="Rows page", unit_count=count, unit=number,
                   sheets=sheets or [], sheet=sheet, status=status)


def _spreadsheet_preview(path, unit, sheet):
    if path.suffix.lower() == ".xls":
        import xlrd
        from xlrd.compdoc import CompDocError
        try:
            book = xlrd.open_workbook(str(path), on_demand=True, logfile=io.StringIO())
        except (xlrd.XLRDError, CompDocError, struct.error) as exc:
            raise PreviewError("The spreadsheet could not be opened. It may be damaged or password-protected.") from exc
        try:
            sheets = book.sheet_names()
            if not sheets or len(sheets) > MAX_SHEETS:
                raise PreviewError(f"Spreadsheet previews support 1 to {MAX_SHEETS} sheets.")
            selected = sheet if sheet is not None else sheets[0]
            if selected not in sheets:
                raise PreviewError("Select an available worksheet.")
            tab = book.sheet_by_name(selected)
            count = max(1, math.ceil(min(tab.nrows, MAX_SHEET_ROWS) / ROWS_PER_PAGE))
            number = _unit(unit, count)
            start = (number - 1) * ROWS_PER_PAGE
            rows = []
            for row in range(start, min(start + ROWS_PER_PAGE, tab.nrows, MAX_SHEET_ROWS)):
                values = []
                for column in range(min(tab.ncols, MAX_SHEET_COLUMNS)):
                    cell = tab.cell(row, column)
                    if cell.ctype == xlrd.XL_CELL_DATE:
                        try:
                            value = xlrd.xldate.xldate_as_datetime(cell.value, book.datemode)
                        except xlrd.xldate.XLDateError:
                            value = cell.value
                    elif cell.ctype == xlrd.XL_CELL_BOOLEAN:
                        value = bool(cell.value)
                    elif cell.ctype == xlrd.XL_CELL_ERROR:
                        value = xlrd.error_text_from_code.get(cell.value, "#ERROR")
                    else:
                        value = cell.value
                    values.append(value)
                rows.append(values)
            return _table_preview(rows, tab.nrows, tab.ncols, number, sheets=sheets, sheet=selected)
        except (xlrd.XLRDError, CompDocError, struct.error) as exc:
            raise PreviewError("The spreadsheet contains damaged workbook data.") from exc
        finally:
            book.release_resources()
    from openpyxl import load_workbook
    from openpyxl.utils.exceptions import InvalidFileException

    _archive_limit(path)
    try:
        book = load_workbook(path, read_only=True, data_only=False, keep_links=False)
    except InvalidFileException as exc:
        raise PreviewError("The spreadsheet format could not be read.") from exc
    try:
        sheets = book.sheetnames
        if not sheets or len(sheets) > MAX_SHEETS:
            raise PreviewError(f"Spreadsheet previews support 1 to {MAX_SHEETS} sheets.")
        selected = sheet if sheet is not None else sheets[0]
        if selected not in sheets:
            raise PreviewError("Select an available worksheet.")
        tab = book[selected]
        row_count, column_count = tab.max_row or 0, tab.max_column or 0
        scanned = None
        if not row_count or not column_count:
            # Streaming XLSX exporters may omit the dimension element. Infer a
            # bounded shape rather than treating real cells as an empty sheet.
            scanned = list(tab.iter_rows(max_row=MAX_SHEET_ROWS + 1,
                                         max_col=MAX_SHEET_COLUMNS + 1, values_only=True))
            row_count = len(scanned)
            column_count = max((index + 1 for row in scanned for index, value in enumerate(row)
                                if value is not None), default=0)
        number = _unit(unit, max(1, math.ceil(min(row_count, MAX_SHEET_ROWS) / ROWS_PER_PAGE)))
        start = (number - 1) * ROWS_PER_PAGE + 1
        end = min(start + ROWS_PER_PAGE - 1, row_count, MAX_SHEET_ROWS)
        rows = (scanned[start - 1:end] if scanned is not None else
                list(tab.iter_rows(min_row=start, max_row=end, max_col=min(column_count, MAX_SHEET_COLUMNS),
                                   values_only=True)) if row_count else [])
        result = _table_preview(rows, row_count, column_count, number, sheets=sheets, sheet=selected)
        if scanned is not None and row_count > MAX_SHEET_ROWS:
            result.status = result.status.replace(f"of {row_count:,}.", f"of at least {row_count:,}.", 1)
        result.status += " Formula cells show their formulas."
        return result
    finally:
        book.close()


def _read_text(path):
    with path.open("rb") as handle:
        raw = handle.read(MAX_TEXT_BYTES + 1)
    truncated = len(raw) > MAX_TEXT_BYTES
    raw = raw[:MAX_TEXT_BYTES]
    encoding = "utf-16" if raw.startswith((b"\xff\xfe", b"\xfe\xff")) else "utf-8-sig"
    text = raw.decode(encoding, errors="replace")
    return text, truncated


class _PlainHTML(HTMLParser):
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.parts = []
        self.hidden = 0

    def handle_starttag(self, tag, attrs):
        if tag in {"script", "style", "noscript", "template"}:
            self.hidden += 1
        elif tag in {"p", "div", "br", "li", "tr", "h1", "h2", "h3", "h4"} and not self.hidden:
            self.parts.append("\n")

    def handle_endtag(self, tag):
        if tag in {"script", "style", "noscript", "template"} and self.hidden:
            self.hidden -= 1

    def handle_data(self, data):
        if not self.hidden:
            self.parts.append(data)


def _text_preview(path, unit):
    text, truncated = _read_text(path)
    if path.suffix.lower() in {".html", ".htm"}:
        parser = _PlainHTML()
        parser.feed(text)
        text = "".join(parser.parts).strip()
    truncated |= len(text) > MAX_TEXT_CHARS
    text = text[:MAX_TEXT_CHARS]
    count = max(1, math.ceil(len(text) / TEXT_CHARS_PER_PAGE))
    number = _unit(unit, count)
    offset = (number - 1) * TEXT_CHARS_PER_PAGE
    return Preview(text=text[offset:offset + TEXT_CHARS_PER_PAGE], unit=number, unit_count=count,
                   unit_label="Text page", status="Text preview." + (
                       " Content exceeds the text preview limit; download the original for the rest." if truncated else ""))


def _csv_preview(path, unit):
    text, truncated = _read_text(path)
    rows = []
    reader = csv.reader(io.StringIO(text), delimiter="\t" if path.suffix.lower() == ".tsv" else ",")
    for row in reader:
        if len(rows) == MAX_SHEET_ROWS:
            truncated = True
            break
        rows.append(row[:MAX_SHEET_COLUMNS + 1])
    width = max((len(row) for row in rows), default=0)
    count = max(1, math.ceil(len(rows) / ROWS_PER_PAGE))
    number = _unit(unit, count)
    start = (number - 1) * ROWS_PER_PAGE
    result = _table_preview(rows[start:start + ROWS_PER_PAGE], len(rows), width, number)
    if truncated:
        result.status += " Content exceeds the CSV preview limit; download the original for the rest."
    return result


def _html_table(rows):
    result = ["<table style=\"border-collapse:collapse;max-width:100%;display:block;overflow:auto\">"]
    wide = False
    for index, row in enumerate(rows):
        if index == 30:
            result.append("<tr><td>Additional rows omitted from this structural preview.</td></tr>")
            break
        cells = list(row)
        wide |= len(cells) > 20
        result.append("<tr>" + "".join("<td style=\"border:1px solid #aaa;padding:6px;vertical-align:top\">" +
                      escape(_cell(cell)) + "</td>" for cell in cells[:20]) + "</tr>")
    if wide:
        result.append("<tr><td>Additional columns omitted from this structural preview.</td></tr>")
    result.append("</table>")
    return "".join(result)


def _embedded_html(blob):
    from PIL import Image

    if len(blob) > MAX_EMBEDDED_BYTES:
        return "<p>[Embedded image exceeds the preview size limit.]</p>"
    with warnings.catch_warnings():
        warnings.simplefilter("error", Image.DecompressionBombWarning)
        try:
            with Image.open(io.BytesIO(blob)) as original:
                image = _normalized_image(original, edge=1024)
            with image, io.BytesIO() as buffer:
                image.save(buffer, format="JPEG", quality=80)
                encoded = base64.b64encode(buffer.getvalue()).decode("ascii")
            return '<img alt="Embedded document image" style="max-width:100%;height:auto" src="data:image/jpeg;base64,' + encoded + '">'
        except (OSError, ValueError, Image.DecompressionBombError, Image.DecompressionBombWarning):
            return "<p>[Embedded image format cannot be previewed.]</p>"


def _structural_html(parts):
    return '<div style="line-height:1.6;overflow-wrap:anywhere;padding:16px">' + "".join(parts) + "</div>"


def _docx_preview(path, unit):
    from docx import Document
    from docx.text.paragraph import Paragraph

    document = Document(path)
    blocks = list(document.iter_inner_content())
    count = max(1, math.ceil(len(blocks) / BLOCKS_PER_PAGE))
    if count > MAX_PAGES:
        raise PreviewError("This Word document exceeds the structural preview page limit.")
    number = _unit(unit, count)
    parts = []
    image_count = 0
    for block in blocks[(number - 1) * BLOCKS_PER_PAGE:number * BLOCKS_PER_PAGE]:
        if isinstance(block, Paragraph):
            style = block.style.name if block.style else ""
            tag = "h" + style[-1] if style.startswith("Heading ") and style[-1:] in "123456" else "p"
            text = block.text
            parts.append(f"<{tag}>" + escape(text[:12_000]) + f"</{tag}>")
            if len(text) > 12_000:
                parts.append("<p>[Long paragraph shortened in this preview.]</p>")
            for relationship_id in block._p.xpath(".//a:blip/@r:embed"):
                if image_count >= MAX_EMBEDDED_IMAGES:
                    parts.append("<p>[Additional embedded images omitted from this preview.]</p>")
                    break
                part = block.part.related_parts.get(relationship_id)
                if part is not None and part.content_type.startswith("image/"):
                    parts.append(_embedded_html(part.blob))
                    image_count += 1
        else:
            parts.append(_html_table(([cell.text for cell in row.cells[:21]] for row in block.rows)))
    return Preview(html=_structural_html(parts), unit=number, unit_count=count, unit_label="Section page",
                   status="Structural Word preview; section pages do not match printed page numbers.")


def _pptx_preview(path, unit):
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    document = Presentation(path)
    count = len(document.slides)
    if not 1 <= count <= MAX_PAGES:
        raise PreviewError(f"Presentation previews support 1 to {MAX_PAGES:,} slides.")
    number = _unit(unit, count)
    parts = []
    images = 0
    omitted_images = False
    pending = list(reversed(list(document.slides[number - 1].shapes)))
    seen = 0
    while pending and seen < 200:
        shape = pending.pop()
        seen += 1
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            pending.extend(reversed(list(shape.shapes)))
        elif shape.has_table:
            parts.append(_html_table(([cell.text for cell in row.cells] for row in shape.table.rows)))
        elif shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
            if images < MAX_EMBEDDED_IMAGES:
                parts.append(_embedded_html(shape.image.blob))
                images += 1
            else:
                omitted_images = True
        elif shape.has_text_frame:
            text = shape.text
            parts.append("<p>" + escape(text[:12_000]).replace("\n", "<br>") + "</p>")
            if len(text) > 12_000:
                parts.append("<p>[Long text block shortened in this preview.]</p>")
    if pending or omitted_images:
        parts.append("<p>[Additional slide objects omitted from this preview.]</p>")
    return Preview(html=_structural_html(parts), unit=number, unit_count=count, unit_label="Slide",
                   status=f"Structural slide {number} of {count}; original layout, charts, and animations may differ.")


def _office_preview(path, data_dir, unit):
    from .office_preview import get_pdf_preview, PreviewUnavailable

    if path.suffix.lower() in {".docx", ".pptx"}:
        _archive_limit(path)
    try:
        converted = _source_path(get_pdf_preview(path, data_dir=data_dir), data_dir)
        result = _pdf_preview(converted, unit)
        result.status += " Converted Office preview."
        return result
    except PreviewUnavailable as exc:
        reason = str(exc)
    if path.suffix.lower() == ".docx":
        from docx.opc.exceptions import OpcError
        from docx.exceptions import InvalidXmlError
        try:
            result = _docx_preview(path, unit)
        except (OpcError, InvalidXmlError) as exc:
            raise PreviewError("The Word document contains damaged or unsupported data.") from exc
    elif path.suffix.lower() == ".pptx":
        from pptx.exc import PythonPptxError
        try:
            result = _pptx_preview(path, unit)
        except PythonPptxError as exc:
            raise PreviewError("The presentation contains damaged or unsupported data.") from exc
    else:
        return Preview(status=f"Office preview unavailable: {reason} Download the original to open it in Office.")
    result.status += f" PDF conversion unavailable: {reason}"
    return result


def preview_document(path, *, data_dir, unit=1, sheet=None):
    """Preview one page/frame/sheet segment from this workspace's sources only.

    No file content is executed or loaded as raw HTML. Heavy format libraries
    and optional Office conversion are imported only when that format is opened.
    Expected malformed/unsupported file errors produce an empty preview/status.
    """
    try:
        source = _source_path(path, data_dir)
        extension = source.suffix.lower()
        if extension == ".pdf":
            return _pdf_preview(source, unit)
        if extension in IMAGE_EXTENSIONS:
            return _image_preview(source, unit)
        if extension in OFFICE_EXTENSIONS:
            return _office_preview(source, data_dir, unit)
        if extension in {".xlsx", ".xls"}:
            return _spreadsheet_preview(source, unit, sheet)
        if extension in {".csv", ".tsv"}:
            return _csv_preview(source, unit)
        if extension in TEXT_EXTENSIONS | {".html", ".htm"}:
            return _text_preview(source, unit)
        return Preview(status="This file type has no built-in preview. Download the original to review it.")
    except (OSError, ValueError, TypeError, LookupError, ArithmeticError, SyntaxError, zipfile.BadZipFile, csv.Error) as exc:
        return Preview(status=f"Preview unavailable: {exc}")
