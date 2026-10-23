"""Trusted slide previews and editable PowerPoint export; no remote rendering."""

import html
from io import BytesIO
import math
import unicodedata

from .mindmap import _clusters
from .presentation import validate_presentation


SLIDE_WIDTH = 13 + 1 / 3
SLIDE_HEIGHT = 7.5
_TITLE_X, _TITLE_Y, _TITLE_WIDTH = 0.65, 0.58, SLIDE_WIDTH - 1.3
_BODY_X, _BODY_WIDTH, _BODY_BOTTOM = 0.93, SLIDE_WIDTH - 1.58, 6.88


def _slides(payload):
    if not isinstance(payload, dict) or payload.get("tool_id") != "presentation":
        raise ValueError("Choose a Presentation output to render slides.")
    sources = payload.get("sources")
    if not isinstance(sources, list) or any(not isinstance(source, dict) or not isinstance(source.get("id"), str)
                                            for source in sources):
        raise ValueError("Presentation sources must have identifiers.")
    return validate_presentation(payload.get("content"), allowed_source_ids={source["id"] for source in sources})["slides"]


def _citation_numbers(payload, slide):
    numbers = {source["id"]: number for number, source in enumerate(payload["sources"], 1)}
    return [numbers[source_id] for source_id in slide["source_ids"]]


def _footer(payload, slide, *, details="speaker notes"):
    numbers = _citation_numbers(payload, slide)
    text = "Sources: " + " ".join(f"[{number}]" for number in numbers[:8])
    if len(numbers) > 8:
        text += f" +{len(numbers) - 8} more"
    return text + f" · Details in {details}"


def render_slide_html(payload, index=0):
    """Return one escaped, bounded HTML preview at a zero-based slide index."""
    slides = _slides(payload)
    if type(index) is not int or not 0 <= index < len(slides):
        raise ValueError("Choose a valid Presentation slide index.")
    slide = slides[index]
    escaped = html.escape
    bullets = "".join(f'<li style="margin:0 0 .65em;padding-left:.15em;overflow-wrap:anywhere">{escaped(bullet)}</li>'
                      for bullet in slide["bullets"])
    return (
        '<div class="rag-presentation-viewport" style="max-height:580px;overflow:auto;padding:4px">'
        '<article class="rag-presentation-slide" style="aspect-ratio:16/9;min-width:420px;width:100%;'
        'box-sizing:border-box;display:flex;flex-direction:column;background:#f8fafc;color:#15243b;'
        'border:1px solid #d7e0eb;border-top:6px solid #2563eb;border-radius:10px;'
        'padding:clamp(18px,3vw,34px);font-family:Arial,sans-serif;box-shadow:0 2px 8px #17255412">'
        f'<div style="font-size:12px;font-weight:700;letter-spacing:.08em;color:#526581">SLIDE {index + 1} / {len(slides)}</div>'
        f'<h2 style="font-size:clamp(23px,2.5vw,34px);line-height:1.12;margin:12px 0 20px;'
        f'white-space:pre-wrap;overflow-wrap:anywhere">{escaped(slide["title"])}</h2>'
        f'<div style="flex:1;min-height:0;overflow:auto"><ul style="font-size:clamp(17px,1.7vw,23px);'
        f'line-height:1.4;margin:0;padding-left:1.2em;white-space:pre-wrap">{bullets}</ul></div>'
        f'<footer style="font-size:11px;color:#526581;border-top:1px solid #d7e0eb;padding-top:10px;margin-top:14px;'
        f'overflow-wrap:anywhere">{escaped(_footer(payload, slide, details="References"))}</footer></article></div>'
    )


def _units(cluster):
    # Account for wide glyphs and also reserve room when a joined emoji falls
    # back to separate glyphs. Combining marks do not add horizontal width.
    # One unit is half an em. Ordinary lowercase/digits reserve .6em (Arial
    # 'e' is wider than .5em); capitals, punctuation, m/w and Unicode get 1em.
    return sum(1.2 if character in "abcdefghijklnopqrstuvxyz0123456789 " else 2
               for character in cluster if character != "\u200d"
               and not unicodedata.category(character).startswith("M")
               and not "\U0001f3fb" <= character <= "\U0001f3ff")


def _wrap_text(text, width, *, prefer_words=True):
    """Wrap normalized display text without dropping graphemes or long tokens."""
    lines, current, used = [], [], 0
    for cluster in _clusters(" ".join(text.split())):
        size = _units(cluster)
        while current and used + size > width:
            boundary = (next((i + 1 for i in range(len(current) - 1, -1, -1) if current[i].isspace()), len(current))
                        if prefer_words else len(current))
            lines.append("".join(current[:boundary]))
            current = current[boundary:]
            used = sum(_units(value) for value in current)
        current.append(cluster)
        used += size
    if current:
        lines.append("".join(current))
    return lines


def slide_layout(slide):
    """Compute a conservative layout in points without system font discovery.

    Text uses explicit soft line breaks, so PowerPoint/LibreOffice do not need
    the font or OS that generated the file to reproduce our line placement.
    Dense text can use character wrapping; all original text remains in notes.
    """
    for prefer_words in (True, False):
        for title_size in (36, 32, 28):
            title_width = math.floor(_TITLE_WIDTH * 72 / (title_size * 0.5))
            title_lines = _wrap_text(slide["title"], title_width, prefer_words=prefer_words)
            if any(sum(_units(cluster) for cluster in _clusters(line)) > title_width for line in title_lines):
                continue  # An unusually long joined grapheme cannot be split safely.
            title_height = len(title_lines) * title_size * 1.05 + 2
            body_top = _TITLE_Y * 72 + title_height + 14.4
            for body_size in (24, 22, 20, 18):
                body_width = math.floor(_BODY_WIDTH * 72 / (body_size * 0.5))
                lines = [_wrap_text(bullet, body_width, prefer_words=prefer_words)
                         for bullet in slide["bullets"]]
                if any(sum(_units(cluster) for cluster in _clusters(line)) > body_width for group in lines for line in group):
                    continue
                heights = [len(group) * body_size * 1.04 + 1 for group in lines]
                body_height = sum(heights) + 3 * (len(heights) - 1)
                if body_top + body_height <= _BODY_BOTTOM * 72:
                    return {"title_size": title_size, "title_lines": title_lines, "title_height": title_height,
                            "body_size": body_size, "body_top": body_top, "body_lines": lines, "body_heights": heights}
    raise ValueError("This slide cannot fit at a readable size. Shorten its title or bullets before exporting PPTX.")


def _speaker_notes(payload, slide):
    from .exports import _provenance
    numbers = {source["id"]: (number, source) for number, source in enumerate(payload["sources"], 1)}
    parts = ["Speaker notes", slide["notes"], "Original slide text", slide["title"], *slide["bullets"], "References"]
    for source_id in slide["source_ids"]:
        number, source = numbers[source_id]
        parts.append(f"[{number}] {source['filename']} — {source['locator']}\n{source['text']}")
    parts.extend(["Provenance", *_provenance(payload)])
    return "\n\n".join(parts)


def presentation_pptx_bytes(payload):
    """Build native text slides with citations and notes; import PPTX lazily."""
    slides = _slides(payload)
    layouts = [slide_layout(slide) for slide in slides]
    try:
        from pptx import Presentation
        from pptx.dml.color import RGBColor
        from pptx.enum.shapes import MSO_SHAPE
        from pptx.enum.text import MSO_AUTO_SIZE, MSO_ANCHOR, PP_ALIGN
        from pptx.util import Inches, Pt
    except ImportError as exc:
        raise RuntimeError("PPTX export requires the optional python-pptx package. Markdown and JSON remain available.") from exc
    from .exports import _xml_text
    presentation = Presentation()
    presentation.slide_width, presentation.slide_height = Inches(SLIDE_WIDTH), Inches(SLIDE_HEIGHT)
    presentation.core_properties.title = _xml_text(payload["title"])[:255]
    presentation.core_properties.author = "RAG Studio"

    def textbox(slide, name, x, y, width, height, text, size, color, *, bold=False, align=PP_ALIGN.LEFT):
        box = slide.shapes.add_textbox(Inches(x), Pt(y), Inches(width), Pt(height))
        box.name = name
        frame = box.text_frame
        frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
        frame.word_wrap = False
        frame.auto_size = MSO_AUTO_SIZE.NONE
        frame.vertical_anchor = MSO_ANCHOR.TOP
        # Soft breaks are inserted by our layout code; generated model text
        # itself has already been checked and cannot contain this control.
        frame.text = "\v".join(_xml_text(line) for line in text.split("\v"))
        for paragraph in frame.paragraphs:
            paragraph.alignment = align
            paragraph.line_spacing = Pt(size * (1.05 if name == "Slide title" else 1.04))
            paragraph.space_before = paragraph.space_after = Pt(0)
            paragraph.font.name = "Arial"
            paragraph.font.size = Pt(size)
            paragraph.font.bold = bold
            paragraph.font.color.rgb = RGBColor.from_string(color)
        return box

    for index, (content, layout) in enumerate(zip(slides, layouts), 1):
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor.from_string("F8FAFC")
        accent = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, presentation.slide_width, Inches(0.08))
        accent.name = "Accent bar"
        accent.fill.solid()
        accent.fill.fore_color.rgb = RGBColor.from_string("2563EB")
        accent.line.fill.background()
        textbox(slide, "Slide number", _TITLE_X, 16, _TITLE_WIDTH, 14, f"SLIDE {index} / {len(slides)}", 10, "526581", bold=True)
        textbox(slide, "Slide title", _TITLE_X, _TITLE_Y * 72, _TITLE_WIDTH, layout["title_height"],
                "\v".join(layout["title_lines"]), layout["title_size"], "15243B", bold=True)
        top = layout["body_top"]
        for number, (lines, height) in enumerate(zip(layout["body_lines"], layout["body_heights"]), 1):
            textbox(slide, f"Bullet marker {number}", _TITLE_X, top, .22, height, "•", layout["body_size"], "2563EB")
            textbox(slide, f"Bullet {number}", _BODY_X, top, _BODY_WIDTH, height,
                    "\v".join(lines), layout["body_size"], "20314B")
            top += height + 3
        textbox(slide, "References", _TITLE_X, 7.03 * 72, _TITLE_WIDTH, 18, _footer(payload, content), 10, "526581")
        slide.notes_slide.notes_text_frame.text = _xml_text(_speaker_notes(payload, content))
    buffer = BytesIO()
    presentation.save(buffer)
    return buffer.getvalue()


def presentation_markdown(payload):
    from .exports import _learning_markdown_text, _provenance
    from .artifacts import citation_label
    plain = _learning_markdown_text
    parts = [f"# {plain(payload['title'])}"]
    for number, slide in enumerate(_slides(payload), 1):
        parts.append(f"## Slide {number}: {plain(slide['title'])}")
        parts.extend(f"- {plain(bullet)}" for bullet in slide["bullets"])
        parts.append(f"**References:** {citation_label(payload, slide['source_ids'])}")
        parts.extend(["### Speaker notes", plain(slide["notes"])])
    parts.append("## Sources")
    for number, source in enumerate(payload["sources"], 1):
        parts.extend([f"### [{number}] {plain(source['filename'])}", f"**Location:** {plain(source['locator'])}", plain(source["text"])])
    parts.extend(["## Provenance", *[plain(line) for line in _provenance(payload)]])
    return "\n\n".join(parts) + "\n"
