"""Local Studio exporters. Optional Office libraries are imported on demand."""

import csv
import html
from io import BytesIO, StringIO
import json
import re
import unicodedata
from xml.etree import ElementTree

from .artifacts import citation_label, render_summary, source_appendix, spreadsheet_table


_INVALID_XML = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\ud800-\udfff\ufffe\uffff]")
_INVALID_FILENAME = re.compile(r'[<>:"/\\|?*\x00-\x1f\x7f]')
_RESERVED_FILENAME = re.compile(r"^(?:CON|PRN|AUX|NUL|COM[1-9¹²³]|LPT[1-9¹²³])(?:\.|$)", re.IGNORECASE)


def summary_export_filename(filename, extension):
    """Validate a portable basename, replacing an optional known export suffix."""
    if extension not in ("md", "json", "docx"):
        raise ValueError("Unsupported Summary export format.")
    return _export_filename(filename, extension, ("md", "markdown", "json", "docx"))


def spreadsheet_export_filename(filename, extension):
    """Use the same portable filename rules for direct Spreadsheet downloads."""
    if extension not in ("csv", "xlsx", "json"):
        raise ValueError("Unsupported Spreadsheet export format.")
    return _export_filename(filename, extension, ("csv", "xlsx", "json"))


def mind_map_export_filename(filename, extension):
    if extension not in ("svg", "md", "json"):
        raise ValueError("Unsupported Mind Map export format.")
    return _export_filename(filename, extension, ("svg", "md", "markdown", "json"))


def learning_export_filename(filename, extension):
    """Use portable basenames for Flashcards and Quizzes downloads."""
    if extension not in ("md", "csv", "json", "txt", "xml"):
        raise ValueError("Unsupported learning export format.")
    return _export_filename(filename, extension, ("md", "markdown", "csv", "json", "txt", "xml"))


def presentation_export_filename(filename, extension):
    if extension not in ("pptx", "md", "json"):
        raise ValueError("Unsupported Presentation export format.")
    return _export_filename(filename, extension, ("pptx", "md", "markdown", "json"))


def graphics_export_filename(filename, extension):
    if extension not in ("png", "jpeg", "webp", "md", "json"):
        raise ValueError("Unsupported Graphics export format.")
    return _export_filename(filename, extension, ("png", "jpg", "jpeg", "webp", "md", "markdown", "json"))


def plot_export_filename(filename, extension):
    if extension not in ("png", "svg", "pdf", "csv", "json"):
        raise ValueError("Unsupported Plot export format.")
    return _export_filename(filename, extension, ("png", "svg", "pdf", "csv", "json"))


def review_export_filename(filename, extension):
    """Portable direct downloads for Compare and Evidence Check."""
    if extension not in ("md", "csv", "json"):
        raise ValueError("Unsupported review export format.")
    return _export_filename(filename, extension, ("md", "markdown", "csv", "json"))


def _export_filename(filename, extension, known_suffixes):
    if not isinstance(filename, str) or not filename or filename != filename.strip():
        raise ValueError("Enter an export filename without leading or trailing spaces.")
    filename = unicodedata.normalize("NFC", filename)
    if (filename.startswith(".") or filename.endswith(".") or _INVALID_FILENAME.search(filename)
            or _RESERVED_FILENAME.match(filename)):
        raise ValueError("Use an export filename without path separators, reserved Windows names, or trailing dots.")
    stem, separator, suffix = filename.rpartition(".")
    if separator and suffix.lower() in known_suffixes:
        filename = stem
    if not filename or filename.endswith((".", " ")):
        raise ValueError("Enter a filename before the export extension, without trailing dots or spaces.")
    result = f"{filename}.{extension}"
    try:
        size = len(result.encode("utf-8"))
    except UnicodeError as exc:
        raise ValueError("The export filename contains unsupported Unicode characters.") from exc
    if len(filename) > 100 or size > 255:
        raise ValueError("Export filenames must have a basename of at most 100 characters and fit 255 UTF-8 bytes including the extension.")
    return result


def _xml_text(value, *, spreadsheet=False):
    text = str(value)
    if _INVALID_XML.search(text):
        raise ValueError("Office export contains unsupported control characters. Edit the text or export JSON instead.")
    if spreadsheet and len(text) > 32767:
        raise ValueError("XLSX cells support at most 32,767 characters. Shorten the text or export CSV/JSON instead.")
    return text


def _provenance(payload):
    text = [f"Model: {payload['model']}", f"Created: {payload.get('created_at', '')}"]
    if payload.get("updated_at"):
        text.append(f"Updated: {payload['updated_at']}")
    if "edited_text" in payload or "edited_rows" in payload:
        text.append("Edited by the user; original citations are retained and may need review.")
    coverage = payload.get("coverage", {})
    text.append("Coverage: " + json.dumps(coverage, ensure_ascii=False))
    text.append("Settings: " + json.dumps(payload.get("settings", {}), ensure_ascii=False))
    return text


def _csv_safe(value):
    # Spreadsheet programs can interpret text from a CSV as a formula, even
    # behind whitespace. Prefix it explicitly without altering saved evidence.
    if value.lstrip(" \t\r\n").startswith(("=", "+", "-", "@")) or value.startswith(("\t", "\r", "\n")):
        return "'" + value
    return value


def _spreadsheet_export_table(payload):
    headers, values = spreadsheet_table(payload)
    headers = ["Source document", *headers, *[f"{column} — citations" for column in headers]]
    rows = []
    for original, edited in zip(payload["content"]["rows"], values):
        rows.append([original["filename"], *edited,
                     *[citation_label(payload, cell["source_ids"]) for cell in original["cells"]]])
    return headers, rows


def _learning_markdown_text(value):
    # Generated content and saved provenance are plain text, including strings
    # that look like HTML entities, links, tables, or Markdown block directives.
    text = html.escape(value, quote=False)
    return re.sub(r"([\\`*_{}\[\]()#+.!|~=\-])", r"\\\1", text)


def _learning_markdown(payload):
    plain = _learning_markdown_text
    parts = [f"# {plain(payload['title'])}"]
    if payload["tool_id"] == "flashcards":
        for number, card in enumerate(payload["content"]["cards"], 1):
            parts.extend([f"## Card {number}", f"**Front:** {plain(card['front'])}",
                          f"**Back:** {plain(card['back'])}",
                          f"**References:** {citation_label(payload, card['source_ids'])}"])
    else:
        for number, question in enumerate(payload["content"]["questions"], 1):
            parts.extend([f"## Question {number}", plain(question["question"])])
            for letter, option in zip("ABCD", question["options"]):
                parts.append(f"- **{letter}:** {plain(option)}")
            answer = question["correct_option"]
            parts.extend([f"**Correct answer:** {'ABCD'[answer]} — {plain(question['options'][answer])}",
                          f"**Explanation:** {plain(question['explanation'])}",
                          f"**References:** {citation_label(payload, question['source_ids'])}"])
    parts.append("## Sources")
    for number, source in enumerate(payload["sources"], 1):
        parts.extend([f"### [{number}] {plain(source['filename'])}",
                      f"**Location:** {plain(source['locator'])}", plain(source["text"])])
    parts.extend(["## Provenance", *[plain(line) for line in _provenance(payload)]])
    return "\n\n".join(parts) + "\n"


def _learning_export_table(payload):
    if payload["tool_id"] == "flashcards":
        return ["Front", "Back", "References"], [
            [card["front"], card["back"], citation_label(payload, card["source_ids"])]
            for card in payload["content"]["cards"]]
    return ["Question", "A", "B", "C", "D", "Correct answer", "Explanation", "References"], [
        [question["question"], *question["options"], "ABCD"[question["correct_option"]],
         question["explanation"], citation_label(payload, question["source_ids"])]
        for question in payload["content"]["questions"]]


def _validated_review(payload):
    from .review import validate_compare, validate_evidence_check

    if payload.get("tool_id") == "compare":
        return validate_compare(payload["content"], documents=payload["documents"], sources=payload["sources"])
    if payload.get("tool_id") == "evidence_check":
        return validate_evidence_check(payload["content"], claims=payload["settings"].get("claims"),
                                       sources=payload["sources"])
    raise ValueError("This output is not Compare or Evidence Check.")


def _review_document_names(payload):
    # Stable IDs distinguish files even when their display names match.
    names = {source["document_id"]: source["filename"] for source in payload["sources"]}
    return {document["id"]: names.get(document["id"], document.get("name", document["id"]))
            for document in payload["documents"]}


def _review_metadata(payload):
    return {key: value for key, value in payload.items() if key not in {"content", "sources", "documents"}}


def render_review_markdown(payload, *, include_sources=True):
    """Render only trusted Markdown templates, treating generated text as plain data."""
    content = _validated_review(payload)
    plain = _learning_markdown_text
    parts = [f"# {plain(payload['title'])}"]
    if payload["tool_id"] == "compare":
        names = _review_document_names(payload)
        for aspect in content["aspects"]:
            parts.append(f"## {plain(aspect['aspect'])}")
            for finding in aspect["findings"]:
                status = "Found" if finding["status"] == "found" else "Not found in selected source"
                parts.extend([f"### {plain(names[finding['document_id']])}", f"**Status:** {status}",
                              plain(finding["text"]),
                              f"**References:** {citation_label(payload, finding['source_ids']) or 'None'}"])
            synthesis = aspect["synthesis"]
            parts.extend(["### Comparison", plain(synthesis["text"]),
                          f"**References:** {citation_label(payload, synthesis['source_ids']) or 'None'}"])
    else:
        parts.append("These assessments concern only the selected source snapshot. "
                     "Not found means the selected sources did not establish the claim; it does not mean the claim is false.")
        labels = {"supported": "Supported", "contradicted": "Contradicted", "mixed": "Mixed",
                  "not_found": "Not found in selected sources"}
        stances = {"supports": "Supports", "contradicts": "Contradicts", "context": "Context"}
        for index, claim in enumerate(content["claims"], 1):
            parts.extend([f"## Claim {index}", plain(claim["text"]), f"**Assessment:** {labels[claim['verdict']]}",
                          f"**Rationale:** {plain(claim['rationale'])}"])
            for evidence in claim["evidence"]:
                parts.extend([f"**{stances[evidence['stance']]} — {citation_label(payload, [evidence['source_id']])}**",
                              plain(evidence["quote"])])
            if not claim["evidence"]:
                parts.append("No cited passage establishes this claim in the selected sources.")
    if include_sources:
        parts.append("## Sources")
        for number, source in enumerate(payload["sources"], 1):
            metadata = {key: value for key, value in source.items() if key != "text"}
            parts.extend([f"### [{number}] {plain(source['filename'])}",
                          f"**Location:** {plain(source['locator'])}", plain(source["text"]),
                          "**Source metadata:** " + plain(json.dumps(metadata, ensure_ascii=False, allow_nan=False))])
        parts.extend(["## Document snapshot",
                      plain(json.dumps(payload["documents"], ensure_ascii=False, allow_nan=False)),
                      "## Provenance", *[plain(line) for line in _provenance(payload)],
                      "**Output metadata:** " + plain(json.dumps(_review_metadata(payload), ensure_ascii=False,
                                                                  allow_nan=False))])
    return "\n\n".join(parts) + "\n"


def _review_csv(payload):
    content = _validated_review(payload)
    if payload["tool_id"] == "compare":
        names = _review_document_names(payload)
        headers = ["Aspect", "Document ID", "Source document", "Status", "Finding or comparison", "References", "Source IDs"]
        rows = []
        for aspect in content["aspects"]:
            for finding in aspect["findings"]:
                rows.append([aspect["aspect"], finding["document_id"], names[finding["document_id"]], finding["status"],
                             finding["text"], citation_label(payload, finding["source_ids"]),
                             json.dumps(finding["source_ids"], ensure_ascii=False)])
            synthesis = aspect["synthesis"]
            rows.append([aspect["aspect"], "", "", "comparison", synthesis["text"],
                         citation_label(payload, synthesis["source_ids"]),
                         json.dumps(synthesis["source_ids"], ensure_ascii=False)])
    else:
        headers = ["Claim ID", "Claim", "Assessment", "Rationale", "Evidence stance", "Verbatim passage", "Reference", "Source ID"]
        rows = []
        for claim in content["claims"]:
            prefix = [claim["id"], claim["text"], claim["verdict"], claim["rationale"]]
            for evidence in claim["evidence"]:
                rows.append([*prefix, evidence["stance"], evidence["quote"],
                             citation_label(payload, [evidence["source_id"]]), evidence["source_id"]])
            if not claim["evidence"]:
                rows.append([*prefix, "", "", "", ""])
    buffer = StringIO(newline="")
    writer = csv.writer(buffer)
    for row in [headers, *rows]:
        writer.writerow([_csv_safe(value) for value in row])
    writer.writerow([])
    writer.writerow(["Sources"])
    writer.writerow(["Citation", "Source ID", "Document ID", "Source document", "Location", "Evidence", "Complete source snapshot"])
    for number, source in enumerate(payload["sources"], 1):
        writer.writerow([str(number), *[_csv_safe(source[key]) for key in ("id", "document_id", "filename", "locator", "text")],
                         _csv_safe(json.dumps(source, ensure_ascii=False, allow_nan=False))])
    writer.writerow([])
    writer.writerow(["Provenance"])
    scope = ("Scope: Assessments concern only the selected source snapshot; not found does not mean false."
             if payload["tool_id"] == "evidence_check" else "Scope: Comparisons concern only the selected source snapshot.")
    for line in [scope, *_provenance(payload),
                 "Documents: " + json.dumps(payload["documents"], ensure_ascii=False, allow_nan=False),
                 "Output metadata: " + json.dumps(_review_metadata(payload), ensure_ascii=False, allow_nan=False)]:
        writer.writerow([_csv_safe(line)])
    return "csv", buffer.getvalue().encode("utf-8-sig")


def _quiz_questions(payload):
    from .learning import validate_learning_content
    content = validate_learning_content(
        "quizzes", payload["content"], allowed_source_ids={source["id"] for source in payload["sources"]})
    return content["questions"]


def _quiz_aiken(payload):
    """Export only importable question blocks; Aiken has no feedback fields."""
    blocks = []
    for question in _quiz_questions(payload):
        # CR/LF, tabs, and Unicode separators cannot introduce extra records.
        prompt = " ".join(question["question"].split())
        # Moodle's Aiken importer treats these prefixes as record syntax even
        # without a following space, and skips lines shorter than two bytes.
        if len(prompt) < 2 or re.match(r"^(?:[A-Za-z][).]|(?i:ANSWER:))", prompt):
            prompt = "Question: " + prompt
        lines = [prompt]
        lines.extend(f"{letter}. {' '.join(option.split())}"
                     for letter, option in zip("ABCD", question["options"]))
        lines.append(f"ANSWER: {'ABCD'[question['correct_option']]}")
        blocks.append("\n".join(lines))
    return ("\n\n".join(blocks) + "\n").encode("utf-8")


def _quiz_moodle_xml(payload):
    """Serialize plain text through XML nodes, never through generated markup."""
    root = ElementTree.Element("quiz")
    source_numbers = {source["id"]: (number, source) for number, source in enumerate(payload["sources"], 1)}

    def text_node(parent, tag, value, **attributes):
        value = str(value)
        if _INVALID_XML.search(value):
            raise ValueError("Moodle XML export contains unsupported control characters. Export JSON instead.")
        element = ElementTree.SubElement(parent, tag, attributes)
        ElementTree.SubElement(element, "text").text = value
        return element

    for number, question in enumerate(_quiz_questions(payload), 1):
        item = ElementTree.SubElement(root, "question", {"type": "multichoice"})
        text_node(item, "name", f"{payload['title']} — Question {number}")
        text_node(item, "questiontext", question["question"], format="plain_text")
        feedback = [question["explanation"], "References:"]
        for source_id in question["source_ids"]:
            citation, source = source_numbers[source_id]
            feedback.append(f"[{citation}] {source['filename']} — {source['locator']}\n{source['text']}")
        text_node(item, "generalfeedback", "\n\n".join(feedback), format="plain_text")
        for tag, value in (("defaultgrade", "1"), ("penalty", "0"), ("single", "true"),
                           ("shuffleanswers", "0"), ("answernumbering", "ABCD")):
            ElementTree.SubElement(item, tag).text = value
        for index, option in enumerate(question["options"]):
            text_node(item, "answer", option, fraction="100" if index == question["correct_option"] else "0",
                      format="plain_text")
    ElementTree.indent(root, space="  ")
    return ElementTree.tostring(root, encoding="utf-8", xml_declaration=True) + b"\n"


def _csv_export(payload, headers, rows):
    buffer = StringIO(newline="")
    writer = csv.writer(buffer)
    for row in [headers, *rows]:
        writer.writerow([_csv_safe(value) for value in row])
    writer.writerow([])
    writer.writerow(["Sources"])
    writer.writerow(["Citation", "Source document", "Location", "Evidence"])
    for number, source in enumerate(payload["sources"], 1):
        writer.writerow([str(number), _csv_safe(source["filename"]), _csv_safe(source["locator"]), _csv_safe(source["text"])])
    writer.writerow([])
    writer.writerow(["Provenance"])
    for line in _provenance(payload):
        writer.writerow([_csv_safe(line)])
    return "csv", buffer.getvalue().encode("utf-8-sig")


def _plot_csv(payload):
    """Keep validated numeric columns numeric; escape all untrusted text."""
    from .plot import validate_plot
    content = validate_plot(payload["content"], sources=payload["sources"])
    buffer = StringIO(newline="")
    writer = csv.writer(buffer)
    writer.writerow(["Series", "Label", "X", "Y", "X label", "Y label", "X unit", "Y unit", "References", "Evidence"])
    for series in content["series"]:
        for point in series["points"]:
            writer.writerow([
                _csv_safe(series["name"]), _csv_safe(point["label"]),
                "" if point["x"] is None else str(point["x"]), str(point["y"]),
                *[_csv_safe(content[key]) for key in ("x_label", "y_label", "x_unit", "y_unit")],
                citation_label(payload, point["source_ids"]), _csv_safe(point["evidence"]),
            ])
    writer.writerow([])
    writer.writerow(["Limitations", "References"])
    for limitation in content["limitations"]:
        writer.writerow([_csv_safe(limitation["text"]), citation_label(payload, limitation["source_ids"])])
    writer.writerow([])
    writer.writerow(["Sources"])
    writer.writerow(["Citation", "Source ID", "Document ID", "Source document", "Location", "Evidence", "Source snapshot"])
    for number, source in enumerate(payload["sources"], 1):
        writer.writerow([str(number), *[_csv_safe(source[key]) for key in
                         ("id", "document_id", "filename", "locator", "text")],
                         _csv_safe(json.dumps(source, ensure_ascii=False, allow_nan=False))])
    writer.writerow([])
    writer.writerow(["Provenance"])
    for line in [f"Title: {payload['title']}", f"Chart type: {content['chart_type']}", *_provenance(payload),
                 "Documents: " + json.dumps(payload["documents"], ensure_ascii=False, allow_nan=False)]:
        writer.writerow([_csv_safe(line)])
    return "csv", buffer.getvalue().encode("utf-8-sig")


def export_bytes(payload, format):
    format = str(format).lower().lstrip(".")
    if format == "json":
        return "json", json.dumps(payload, ensure_ascii=False, allow_nan=False, indent=2).encode("utf-8")
    tool = payload["tool_id"]
    if tool in ("compare", "evidence_check"):
        if format in ("md", "markdown"):
            return "md", render_review_markdown(payload).encode("utf-8")
        if format == "csv":
            return _review_csv(payload)
    if tool == "plot":
        if format in ("png", "svg", "pdf"):
            from .plot_render import plot_bytes
            return format, plot_bytes(payload, format)
        if format == "csv":
            return _plot_csv(payload)
    if tool == "graphics" and format in ("md", "markdown"):
        from .artifacts import _markdown_text
        plain = _markdown_text
        content = payload["content"]
        parts = [f"# {plain(payload['title'])}", "## Caption",
                 plain(content["caption"]) + " " + citation_label(payload, content["source_ids"]),
                 "## Visual brief", plain(content["prompt"])]
        if content.get("revised_prompt"):
            parts.extend(["## Image model revised prompt", plain(content["revised_prompt"])])
        parts.extend([source_appendix(payload), "## Source details"])
        for number, source in enumerate(payload["sources"], 1):
            parts.append(f"### [{number}] {plain(source['filename'])} — {plain(source['locator'])}")
            parts.append(plain(source["text"]))
        parts.extend(["## Provenance", *[plain(line) for line in _provenance(payload)]])
        return "md", ("\n\n".join(parts) + "\n").encode("utf-8")
    if tool == "presentation" and format in ("pptx", "md", "markdown"):
        from .presentation_render import presentation_pptx_bytes, presentation_markdown
        if format == "pptx":
            return "pptx", presentation_pptx_bytes(payload)
        return "md", presentation_markdown(payload).encode("utf-8")
    if tool == "quizzes" and format == "aiken":
        return "txt", _quiz_aiken(payload)
    if tool == "quizzes" and format == "moodle_xml":
        return "xml", _quiz_moodle_xml(payload)
    if tool in ("flashcards", "quizzes"):
        if format in ("md", "markdown"):
            return "md", _learning_markdown(payload).encode("utf-8")
        if format == "csv":
            headers, rows = _learning_export_table(payload)
            return _csv_export(payload, headers, rows)
    if tool == "mind_maps" and format == "svg":
        from .mindmap import render_mind_map_svg
        return "svg", render_mind_map_svg(payload).encode("utf-8")
    if tool == "mind_maps" and format in ("md", "markdown"):
        from .mindmap import mind_map_outline
        from .artifacts import _markdown_text
        parts = [f"# {_markdown_text(payload['title'])}", ""]
        for depth, node in mind_map_outline(payload):
            label = _markdown_text(node["label"]).replace("\n", " ")
            description = _markdown_text(node["description"]).replace("\n", " ")
            citations = citation_label(payload, node["source_ids"])
            parts.append("  " * depth + f"- **{label}**" + (f" — {description}" if description else "")
                         + (f" {citations}" if citations else ""))
        parts.extend(["", source_appendix(payload), "", "## Provenance", "", "\n\n".join(_provenance(payload))])
        return "md", ("\n".join(parts) + "\n").encode("utf-8")
    if tool == "summary" and format in ("md", "markdown"):
        text = render_summary(payload) + "\n\n## Provenance\n\n" + "\n\n".join(_provenance(payload)) + "\n"
        return "md", text.encode("utf-8")
    if tool == "summary" and format == "docx":
        try:
            from docx import Document
        except ImportError as exc:
            raise RuntimeError("DOCX export requires the optional python-docx package. Markdown and JSON remain available.") from exc
        document = Document()
        # Text is always written as text; model output never becomes XML or code.
        for line in render_summary(payload).splitlines():
            line = _xml_text(line)
            if line.startswith("# "):
                document.add_heading(line[2:], level=0)
            elif line.startswith("## "):
                document.add_heading(line[3:], level=1)
            elif line.strip():
                document.add_paragraph(line)
        document.add_heading("Provenance", level=1)
        for line in _provenance(payload):
            document.add_paragraph(_xml_text(line))
        buffer = BytesIO()
        document.save(buffer)
        return "docx", buffer.getvalue()
    if tool == "spreadsheet" and format == "csv":
        headers, rows = _spreadsheet_export_table(payload)
        return _csv_export(payload, headers, rows)
    if tool == "spreadsheet" and format == "xlsx":
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Alignment, Font
        except ImportError as exc:
            raise RuntimeError("XLSX export requires the optional openpyxl package. CSV and JSON remain available.") from exc
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Spreadsheet"
        headers, rows = _spreadsheet_export_table(payload)

        def write_rows(target, values):
            for row_index, row in enumerate(values, 1):
                for column_index, value in enumerate(row, 1):
                    cell = target.cell(row_index, column_index)
                    cell.value = _xml_text(value, spreadsheet=True)
                    cell.data_type = "s"  # Never evaluate extracted or generated formulas.
                    cell.alignment = Alignment(wrap_text=True, vertical="top")
                    if row_index == 1:
                        cell.font = Font(bold=True)
            target.freeze_panes = "A2"
            for column in target.columns:
                target.column_dimensions[column[0].column_letter].width = 32

        write_rows(sheet, [headers, *rows])
        sheet.auto_filter.ref = sheet.dimensions
        sources = workbook.create_sheet("Sources")
        write_rows(sources, [["Citation", "Source document", "Location", "Evidence"], *[
            [str(number), source["filename"], source["locator"], source["text"]]
            for number, source in enumerate(payload["sources"], 1)]])
        provenance = workbook.create_sheet("Provenance")
        write_rows(provenance, [["Provenance"], *[[line] for line in _provenance(payload)]])
        buffer = BytesIO()
        workbook.save(buffer)
        return "xlsx", buffer.getvalue()
    raise ValueError(f"{format!r} is not an available export format for {tool}.")
