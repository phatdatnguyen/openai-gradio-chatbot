"""Exercise Studio and Plugins over localhost with temporary data and fake models.

Run explicitly: python tests/manual_tools_smoke.py
"""

from functools import partial
from collections import Counter
import base64
import csv
from io import BytesIO
import json
import os
from pathlib import Path
import re
import sys
import tempfile
from urllib.parse import urljoin
from xml.etree import ElementTree

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))


def main():
    with tempfile.TemporaryDirectory(prefix="studio-rpc-smoke-") as directory:
        root = Path(directory)
        os.environ["RAG_DATA_DIR"] = str(root / "data")
        os.environ["GRADIO_TEMP_DIR"] = str(root / "gradio")
        os.environ["GRADIO_ANALYTICS_ENABLED"] = "False"
        os.environ["OPENAI_API_KEY"] = "smoke-test-not-a-real-key"

        from gradio_client import Client
        import httpx
        from qdrant_client import QdrantClient
        from chatbot import api, config, history_store
        from plugins import manager, ui as plugins_ui
        from rag_chatbot import chatbot as rag_ui, workspaces
        from rag_chatbot.config import Settings
        from rag_chatbot.service import RAGService
        from rag_chatbot.store import LibraryStore
        from rag_chatbot.tools import artifacts, graphics_ui, learning_ui, plot_ui, presentation_ui, review_ui, service as tool_service, ui as tools_ui

        def prohibit_client():
            raise AssertionError("Smoke tests must never access the paid API client.")

        config.get_client = prohibit_client
        history_store.HISTORY_DIR = str(root / "chat-history")
        manager.SETTINGS_PATH = root / "preferences" / "plugins.json"
        services, requests = {}, []

        def get_service(workspace):
            if workspace not in services:
                settings = Settings(data_dir=workspaces.workspace_path(workspace),
                                    embedding_dimensions=8, chunk_size=80, chunk_overlap=10)
                store = LibraryStore(settings, vector_client=QdrantClient(":memory:"))
                services[workspace] = RAGService(settings, store=store,
                    embedder=lambda texts: [[1.0] + [0.0] * 7 for _ in texts])
            return services[workspace]

        def generate(model, history, schema, **kwargs):
            requests.append(kwargs["schema_name"])
            prompt = json.loads(history[-1]["content"])
            if kwargs["schema_name"] == "tool_plot_columns":
                assert prompt["table"]["row_count"] == 150
                assert len(prompt["table"]["sample_rows"]) == 3
                assert prompt["table"]["columns"] == [{"index": index, "name": name}
                    for index, name in enumerate(["Id", "sepal_length", "sepal_width", "species"])]
                return {"status": "complete", "title": "Iris sepal measurements", "chart_type": "scatter",
                        "x_column": 1, "y_column": 2, "group_column": 3, "reason": ""}
            evidence = prompt["evidence"]
            if kwargs["schema_name"] == "tool_evidence_notes":
                return {"notes": [{"source_id": item["source_id"], "text": item["text"]}
                                  for item in evidence]}
            if kwargs["schema_name"] == "tool_summary":
                return {"title": "Research summary", "sections": [{"heading": "Findings",
                        "paragraphs": evidence["notes"]}]}
            if kwargs["schema_name"] == "tool_spreadsheet_row":
                return {"cells": [{"column": column, "value": "Measured result",
                        "source_ids": evidence["notes"][0]["source_ids"]}
                        for column in prompt["requested_columns"]]}
            if kwargs["schema_name"] == "tool_presentation":
                citations = evidence["notes"][0]["source_ids"]
                return {"title": "XRD presentation", "slides": [
                    {"title": "Measured peak" if index == 0 else f"Sample Alpha finding {index + 1}",
                     "bullets": ["Sample Alpha has an XRD peak at 26.5 degrees."],
                     "notes": f"Explain the measured result on slide {index + 1}.", "source_ids": citations}
                    for index in range(prompt["requested_slide_count"])]}
            if kwargs["schema_name"] == "tool_graphics":
                return {"title": "XRD graphic", "prompt": "Illustrate the XRD peak measured at 26.5 degrees for sample Alpha.",
                        "caption": "Sample Alpha has an XRD peak at 26.5 degrees.",
                        "source_ids": evidence["notes"][0]["source_ids"]}
            if kwargs["schema_name"] == "tool_plot":
                return {"status": "complete", "title": "Measured peak position", "chart_type": "bar",
                        "x_label": "Sample", "y_label": "Peak position", "x_unit": "", "y_unit": "degrees",
                        "series": [{"name": "XRD", "points": [
                            {"label": "Alpha", "x": None, "y": 26.5,
                             "source_ids": evidence["notes"][0]["source_ids"],
                             "evidence": "Peak position is 26.5 degrees."}]}], "limitations": []}
            if kwargs["schema_name"] == "tool_compare":
                findings = []
                for document in evidence["documents"]:
                    note = next(item for item in evidence["notes"]
                                if set(item["source_ids"]).intersection(document["source_ids"]))
                    findings.append({"document_id": document["id"], "status": "found",
                                     "text": note["text"], "source_ids": note["source_ids"]})
                return {"title": "XRD comparison", "aspects": [{"aspect": "Reported peak",
                    "findings": findings, "synthesis": {"text": "The documents report different peak positions.",
                        "source_ids": [source_id for item in findings for source_id in item["source_ids"]]}}]}
            if kwargs["schema_name"] == "tool_evidence_check":
                first = next(item for item in evidence["notes"] if "26.5" in item["text"])
                second = next(item for item in evidence["notes"] if "30.0" in item["text"])
                support = {"source_id": first["source_ids"][0], "quote": "Peak position is 26.5 degrees.", "stance": "supports"}
                conflict = {"source_id": second["source_ids"][0], "quote": "Peak position is 30.0 degrees.", "stance": "contradicts"}
                return {"title": "XRD evidence review", "claims": [
                    {"id": claim["id"], "verdict": verdict, "rationale": reason, "evidence": items}
                    for claim, verdict, reason, items in zip(prompt["requested_claims"],
                        ["supported", "contradicted", "not_found", "mixed"],
                        ["A peak measurement is reported.", "The reported peak differs from the claim.",
                         "The selected documents do not report costs.", "The sources disagree about the peak."],
                        [[support], [{**support, "stance": "contradicts"}], [], [support, conflict]])]}
            if kwargs["schema_name"] == "tool_mind_map":
                return {"title": "XRD research map", "nodes": [
                    {"id": "1", "parent_id": "", "label": "XRD research",
                     "description": "", "source_ids": []},
                    {"id": "1.1 peak", "parent_id": "1", "label": "Peak measurement",
                     "description": "Sample Alpha has a peak at 26.5 degrees.",
                     "source_ids": evidence["notes"][0]["source_ids"]}]}
            if kwargs["schema_name"] == "tool_flashcards":
                assert prompt["requested_count"] == 2
                citations = evidence["notes"][0]["source_ids"]
                return {"title": "XRD study cards", "cards": [
                    {"front": "Where is the measured peak?", "back": "26.5 degrees.", "source_ids": citations},
                    {"front": "Which sample was measured?", "back": "Sample Alpha.", "source_ids": citations}]}
            if kwargs["schema_name"] == "tool_quizzes":
                assert prompt["requested_count"] == 2
                citations = evidence["notes"][0]["source_ids"]
                return {"title": "XRD knowledge check", "questions": [
                    {"question": "Where is the measured peak?", "options": ["20 degrees", "26.5 degrees", "40 degrees", "60 degrees"],
                     "correct_option": 1, "explanation": "The measurement puts the peak at 26.5 degrees.", "source_ids": citations},
                    {"question": "Which sample was measured?", "options": ["Alpha", "Beta", "Gamma", "Delta"],
                     "correct_option": 0, "explanation": "The source names sample Alpha.", "source_ids": citations}]}
            raise AssertionError(f"Unexpected model request: {kwargs['schema_name']}")

        api.generate_structured = generate

        def generate_image(prompt, image_path, image_model, image_size, image_quality,
                           image_background, image_output_format, image_output_compression,
                           image_moderation, image_input_fidelity):
            from PIL import Image
            assert image_model == config.DEFAULT_IMAGE_MODEL == "gpt-image-2.5-sunburst"
            assert image_path is None and "26.5" in prompt
            requests.append("image_generation")
            data = BytesIO()
            with Image.new("RGB", (64, 32), "navy") as image:
                image.save(data, format=image_output_format.upper())
            return {"b64_json": base64.b64encode(data.getvalue()).decode("ascii"),
                    "output_format": image_output_format, "revised_prompt": None}

        api.generate_image_asset = generate_image
        rag_ui.get_service = tool_service.get_service = get_service
        import webui
        app = webui.demo
        client = None
        try:
            app.queue().launch(server_name="127.0.0.1", prevent_thread_lock=True,
                               quiet=True, inbrowser=False, share=False, show_error=True, css=webui.APP_CSS)
            client = Client(app.local_url, verbose=False, download_files=root / "downloads")

            def block_for(function, tool_id=None):
                for block in app.fns.values():
                    if tool_id is None and block.fn is function:
                        return block
                    if (tool_id is not None and isinstance(block.fn, partial)
                            and block.fn.func is function and block.fn.args == (tool_id,)):
                        return block
                raise AssertionError(f"Missing callback: {function.__name__} / {tool_id}")

            def call(function, *args, tool_id=None):
                block = block_for(function, tool_id)
                result = client.predict(*args, api_name="/" + block.api_name)
                outputs = [item for item in block.outputs if not item.skip_api]
                values = result if len(outputs) != 1 else (result,)
                return {item._id: value for item, value in zip(outputs, values)}

            call(rag_ui.on_initial_load)
            workspace = workspaces.default_workspace()
            local = get_service(workspace)
            source = root / "measurements.txt"
            source.write_text("XRD was measured on sample Alpha. Peak position is 26.5 degrees.\n",
                              encoding="utf-8")
            list(local.ingest([str(source)], "Default"))
            document_id = local.store.list_documents()[0]["id"]
            call(rag_ui.on_refresh, "Default", [], None)
            call(rag_ui.on_source_selection, [document_id])
            call(tools_ui.on_reset)

            summary = app.rag_tab.tools.controls["summary"]
            result = call(tools_ui.on_generate, "Summarize the measurement.", [document_id],
                          "Default", config.DEFAULT_MODEL, "medium", tool_id="summary")
            assert "Completed" in result[summary.status._id], result
            assert "[1]" in result[summary.preview._id]["value"]
            assert result[summary.preview._id]["visible"] is True
            assert "1 / 1" in result[summary.coverage._id]
            current = artifacts.load_current_summary(workspace)
            editing = call(tools_ui.on_summary_edit)
            assert editing[summary.editor._id]["visible"] is True
            assert editing[summary.preview._id]["visible"] is False
            editor = editing[summary.editor._id]["value"] + "\n\nUser review complete."
            saved = call(tools_ui.on_summary_save, editor)
            assert saved[summary.editor._id]["visible"] is False
            assert saved[summary.preview._id]["value"] == editor
            assert artifacts.load_current_summary(workspace)["id"] == current["id"]
            assert artifacts.list_artifacts(workspace, "summary") == []
            reference = call(tools_ui.on_reference, "0")
            assert "26.5" in reference[summary.excerpt._id]
            assert Path(reference[summary.original._id]).is_file()
            for format, suffix in [("markdown", ".md"), ("docx", ".docx"), ("json", ".json")]:
                exported = call(tools_ui.on_summary_export, "XRD summary", format)
                path = Path(exported[summary.download._id])
                assert path.name == "XRD summary" + suffix
                assert path.stat().st_size > 0
            assert artifacts.load_current_summary(workspace)["id"] == current["id"]
            assert artifacts.list_artifacts(workspace, "summary") == []

            spreadsheet = app.rag_tab.tools.controls["spreadsheet"]
            result = call(tools_ui.on_generate, "Extract values.", "Method, Result", [document_id],
                          "Default", config.DEFAULT_MODEL, "medium", tool_id="spreadsheet")
            assert "Completed" in result[spreadsheet.status._id], result
            table = result[spreadsheet.editor._id]
            table = table.get("value", table)
            grid = {key: table[key] for key in ("headers", "data")}
            assert len(grid["data"]) == 1 and grid["headers"] == ["Method", "Result"]
            current_sheet = artifacts.load_current_spreadsheet(workspace)
            assert current_sheet["content"]["columns"] == grid["headers"]
            assert artifacts.list_artifacts(workspace, "spreadsheet") == []
            reopened = call(tools_ui.on_reset)[spreadsheet.editor._id]
            assert reopened["value"]["data"] == grid["data"]
            grid["data"][0][1] = "=1+1"  # Exported as text, never an executable formula.
            saved_versions = len(artifacts.list_artifacts(workspace, "spreadsheet"))
            exported = call(tools_ui.on_spreadsheet_export, grid, "XRD spreadsheet", "xlsx")
            assert Path(exported[spreadsheet.download._id]).name == "XRD spreadsheet.xlsx"
            assert len(artifacts.list_artifacts(workspace, "spreadsheet")) == saved_versions
            from openpyxl import load_workbook
            workbook = load_workbook(exported[spreadsheet.download._id], data_only=False)
            assert workbook["Spreadsheet"]["C2"].value == "=1+1"
            assert workbook["Spreadsheet"]["C2"].data_type == "s"
            assert workbook.sheetnames == ["Spreadsheet", "Sources", "Provenance"]
            workbook.close()
            saved_sheet = call(tools_ui.on_spreadsheet_save, grid)
            assert "Spreadsheet changes saved" in saved_sheet[spreadsheet.status._id]
            assert artifacts.load_current_spreadsheet(workspace)["id"] == current_sheet["id"]
            assert artifacts.load_current_spreadsheet(workspace)["edited_rows"] == grid["data"]
            assert artifacts.list_artifacts(workspace, "spreadsheet") == []

            mind_map = app.rag_tab.tools.controls["mind_maps"]
            mapped = call(tools_ui.on_generate, tools_ui.DEFAULT_MIND_MAP_INSTRUCTIONS, [document_id],
                          "Default", config.DEFAULT_MODEL, "medium", tool_id="mind_maps")
            assert "Completed" in mapped[mind_map.status._id], mapped
            map_preview = mapped[mind_map.preview._id]
            map_preview = map_preview.get("value", "") if isinstance(map_preview, dict) else map_preview
            assert "<svg" in map_preview and "Peak measurement" in map_preview
            current_map = artifacts.load_current_mind_map(workspace)
            assert [node["id"] for node in current_map["content"]["nodes"]] == ["n1", "n2"]
            assert current_map["content"]["nodes"][1]["parent_id"] == "n1"
            assert current_map["content"]["nodes"][1]["source_ids"] == [current_map["sources"][0]["id"]]
            zoomed = call(tools_ui.on_mind_map_zoom, 1.5)
            assert "svg" in str(zoomed[mind_map.preview._id])
            for format, suffix in [("svg", ".svg"), ("markdown", ".md"), ("json", ".json")]:
                exported = call(tools_ui.on_mind_map_export, "XRD mind map", format)
                path = Path(exported[mind_map.download._id])
                assert path.name == "XRD mind map" + suffix and path.stat().st_size > 0
                if format == "svg":
                    from xml.etree import ElementTree
                    svg = ElementTree.parse(path).getroot()
                    assert svg.tag == "{http://www.w3.org/2000/svg}svg"
                    assert "Peak measurement" in "".join(svg.itertext())
                elif format == "json":
                    assert json.loads(path.read_text(encoding="utf-8"))["content"] == current_map["content"]
            assert artifacts.load_current_mind_map(workspace)["id"] == current_map["id"]
            assert artifacts.list_artifacts(workspace, "mind_maps") == []

            def value(result, component):
                output = result[component._id]
                if isinstance(output, dict) and output.get("__type__") == "update":
                    output = output.get("value", "")
                return re.sub(r"\\([\\`*_{}\[\]<>()#+.!|~-])", r"\1", output) if isinstance(output, str) else output

            def verify_plot_preview(result):
                preview = value(result, plot.preview)
                match = re.search(r"data:image/png;base64,([A-Za-z0-9+/=]+)", preview)
                assert match is not None, "Plot restore must include its rendered chart."
                with Image.open(BytesIO(base64.b64decode(match[1], validate=True))) as image:
                    image.load()
                    assert image.format == "PNG" and image.width >= 800 and image.height >= 400
                table = value(result, plot.table)
                assert "26.5" in str(table) and "Alpha" in str(table)

            def verify_graphic_preview(result):
                assert "26.5" in value(result, graphics.caption)
                graphic_file = value(result, graphics.preview)
                assert isinstance(graphic_file, dict) and graphic_file.get("url"), graphic_file
                response = httpx.get(urljoin(app.local_url, graphic_file["url"]), timeout=15)
                response.raise_for_status()
                assert response.headers["content-type"].startswith("image/")
                assert response.content == image_path.read_bytes()
                with Image.open(BytesIO(response.content)) as image:
                    image.load()
                    assert image.size == (64, 32)

            presentation = app.rag_tab.tools.controls["presentation"]
            tool_order = list(app.rag_tab.tools.controls)
            assert tool_order.index("presentation") + 1 == tool_order.index("graphics")
            assert tool_order.index("graphics") + 1 == tool_order.index("plot")
            assert tool_order.index("plot") + 1 == tool_order.index("mind_maps")
            assert presentation.slide_count.value == 10 and presentation.output_format.value == "pptx"
            deck_result = call(presentation_ui.on_generate, "Present the measurement.", 2, [document_id],
                               "Default", config.DEFAULT_MODEL, "medium")
            assert "Completed" in value(deck_result, presentation.status), deck_result
            assert "Measured peak" in value(deck_result, presentation.preview)
            assert "slide 1" in value(deck_result, presentation.notes)
            current_deck = artifacts.load_current_presentation(workspace)
            assert len(current_deck["content"]["slides"]) == current_deck["settings"]["slide_count"] == 2
            local_before = len(requests)
            next_slide = call(presentation_ui.on_navigate, None, tool_id="next")
            assert "Sample Alpha finding 2" in value(next_slide, presentation.preview)
            assert value(next_slide, presentation.progress) == "Slide 2 of 2"
            selected_slide = call(presentation_ui.on_navigate, "0", tool_id="select")
            assert value(selected_slide, presentation.progress) == "Slide 1 of 2"
            cited_slide = call(presentation_ui.on_reference, "0")
            assert "26.5 degrees" in value(cited_slide, presentation.excerpt)
            for format, suffix in (("pptx", ".pptx"), ("markdown", ".md"), ("json", ".json")):
                exported = call(presentation_ui.on_export, "XRD presentation", format)
                path = Path(exported[presentation.download._id])
                assert path.name == "XRD presentation" + suffix and path.stat().st_size > 0
                if format == "pptx":
                    from pptx import Presentation
                    deck = Presentation(path)
                    assert len(deck.slides) == 2
                    assert "Measured peak" in " ".join(shape.text for shape in deck.slides[0].shapes if shape.has_text_frame)
                    assert "measurements.txt" in deck.slides[0].notes_slide.notes_text_frame.text
                elif format == "json":
                    assert json.loads(path.read_text(encoding="utf-8"))["content"] == current_deck["content"]
            assert artifacts.load_current_presentation(workspace) == current_deck
            assert artifacts.list_artifacts(workspace, "presentation") == []
            assert len(requests) == local_before
            regenerated_deck = call(presentation_ui.on_generate, "Present the measurement again.", 3, [document_id],
                                   "Default", config.DEFAULT_MODEL, "medium")
            assert value(regenerated_deck, presentation.progress) == "Slide 1 of 3"
            assert artifacts.load_current_presentation(workspace)["id"] == current_deck["id"]

            graphics = app.rag_tab.tools.controls["graphics"]
            assert graphics.image_model.value == app.chat_tab.image_model.value == config.DEFAULT_IMAGE_MODEL
            graphic_args = ["Create a diagram of the measurement.", config.DEFAULT_IMAGE_MODEL,
                            "1024x1024", "max", "opaque", "png", 100, "auto",
                            [document_id], "Default", config.DEFAULT_MODEL, "medium"]
            graphic_result = call(graphics_ui.on_generate, *graphic_args)
            assert "Completed" in value(graphic_result, graphics.status), graphic_result
            assert "26.5" in value(graphic_result, graphics.caption)
            assert "26.5" in value(graphic_result, graphics.brief)
            current_graphic = artifacts.load_current_graphics(workspace)
            assert "image_data" not in current_graphic
            assert "b64_json" not in json.dumps(current_graphic)
            image_path = artifacts.graphics_image_path(workspace, current_graphic)
            from PIL import Image
            with Image.open(image_path) as image:
                assert image.size == (64, 32)
            local_before = len(requests)
            reference = call(graphics_ui.on_reference, "0")
            assert "26.5" in value(reference, graphics.excerpt)
            for format, suffix in (("image", ".png"), ("markdown", ".md"), ("json", ".json")):
                exported = call(graphics_ui.on_export, "XRD graphic", format)
                path = Path(exported[graphics.download._id])
                assert path.name == "XRD graphic" + suffix and path.stat().st_size > 0
                if format == "image":
                    assert path.read_bytes() == image_path.read_bytes()
                elif format == "json":
                    assert json.loads(path.read_text(encoding="utf-8"))["content"] == current_graphic["content"]
            assert artifacts.load_current_graphics(workspace) == current_graphic
            assert len(requests) == local_before

            plot = app.rag_tab.tools.controls["plot"]
            assert plot.chart_type.value == "auto" and plot.output_format.value == "png"
            plotted = call(plot_ui.on_generate, "Plot the reported peak position.", "bar", [document_id],
                           "Default", config.DEFAULT_MODEL, "medium")
            assert "Completed" in value(plotted, plot.status), plotted
            assert "data:image/png;base64," in value(plotted, plot.preview)
            plotted_table = value(plotted, plot.table)
            assert "26.5" in str(plotted_table) and "Alpha" in str(plotted_table)
            current_plot = artifacts.load_current_plot(workspace)
            assert current_plot["content"]["chart_type"] == "bar"
            assert current_plot["content"]["series"][0]["points"][0]["y"] == 26.5
            local_before = len(requests)
            plot_reference = call(plot_ui.on_reference, "0")
            assert "26.5" in value(plot_reference, plot.excerpt)
            for format in ("png", "svg", "pdf", "csv", "json"):
                exported = call(plot_ui.on_export, "XRD plot", format)
                path = Path(exported[plot.download._id])
                assert path.name == "XRD plot." + format and path.stat().st_size > 0
                if format == "png":
                    with Image.open(path) as chart:
                        assert chart.format == "PNG" and chart.width >= 800 and chart.height >= 400
                elif format == "svg":
                    chart = ElementTree.parse(path).getroot()
                    assert chart.tag == "{http://www.w3.org/2000/svg}svg"
                elif format == "pdf":
                    from pypdf import PdfReader
                    assert len(PdfReader(path).pages) == 1
                elif format == "csv":
                    with path.open(encoding="utf-8-sig", newline="") as handle:
                        rows = list(csv.reader(handle))
                    assert any("26.5" in row for row in rows) and "measurements.txt" in str(rows)
                else:
                    assert json.loads(path.read_text(encoding="utf-8"))["content"] == current_plot["content"]
            assert artifacts.load_current_plot(workspace) == current_plot
            assert artifacts.list_artifacts(workspace, "plot") == []
            assert len(requests) == local_before
            verify_plot_preview(plotted)

            flashcards = app.rag_tab.tools.controls["flashcards"]
            assert flashcards.item_count.value == 20
            cards = call(tools_ui.on_generate, "Create study cards.", [document_id], "Default",
                         config.DEFAULT_MODEL, "medium", 2, tool_id="flashcards")
            assert "Completed" in value(cards, flashcards.status), cards
            assert "Where is the measured peak" in value(cards, flashcards.preview)
            assert "26.5" not in value(cards, flashcards.answer)
            current_cards = artifacts.load_current_learning(workspace, "flashcards")
            assert current_cards["settings"]["item_count"] == 2 and len(current_cards["content"]["cards"]) == 2
            local_before = len(requests)
            revealed = call(learning_ui.on_flashcard, tool_id="reveal")
            assert "26.5" in value(revealed, flashcards.answer)
            assert "[1]" in value(revealed, flashcards.answer)
            next_card = call(learning_ui.on_flashcard, tool_id="next")
            assert "Which sample" in value(next_card, flashcards.preview)
            assert "Sample Alpha" not in value(next_card, flashcards.answer)
            call(learning_ui.on_flashcard, tool_id="previous")
            assert len(requests) == local_before

            quizzes = app.rag_tab.tools.controls["quizzes"]
            assert quizzes.output_format.value == "aiken"
            assert quizzes.item_count.value == 20
            quiz = call(tools_ui.on_generate, "Test the measurement.", [document_id], "Default",
                        config.DEFAULT_MODEL, "medium", 2, tool_id="quizzes")
            assert "Completed" in value(quiz, quizzes.status), quiz
            assert "Where is the measured peak" in value(quiz, quizzes.preview)
            assert "measurement puts" not in value(quiz, quizzes.feedback)
            assert quiz[quizzes.options._id]["choices"] == [
                ["A. 20 degrees", "0"], ["B. 26.5 degrees", "1"], ["C. 40 degrees", "2"], ["D. 60 degrees", "3"]]
            current_quiz = artifacts.load_current_learning(workspace, "quizzes")
            assert current_quiz["settings"]["item_count"] == 2 and len(current_quiz["content"]["questions"]) == 2
            local_before = len(requests)
            checked = call(learning_ui.on_quiz, "1", tool_id="check")
            assert "Correct" in value(checked, quizzes.feedback)
            assert "26.5" in value(checked, quizzes.feedback) and "[1]" in value(checked, quizzes.feedback)
            assert "Answer: B. 26.5 degrees" in value(checked, quizzes.feedback)
            repeated = call(learning_ui.on_quiz, "0", tool_id="check")
            assert value(repeated, quizzes.score) == value(checked, quizzes.score)
            next_question = call(learning_ui.on_quiz, None, tool_id="next")
            assert "Which sample" in value(next_question, quizzes.preview)
            wrong = call(learning_ui.on_quiz, "1", tool_id="check")
            assert "Not quite" in value(wrong, quizzes.feedback) and "Alpha" in value(wrong, quizzes.feedback)
            assert "1" in value(wrong, quizzes.score) and "2" in value(wrong, quizzes.score)
            restarted = call(learning_ui.on_quiz, None, tool_id="restart")
            assert "Where is the measured peak" in value(restarted, quizzes.preview)
            assert value(restarted, quizzes.feedback) == ""
            assert len(requests) == local_before
            for tool_id, control, current_output in (("flashcards", flashcards, current_cards), ("quizzes", quizzes, current_quiz)):
                formats = [("markdown", ".md"), ("csv", ".csv"), ("json", ".json")]
                if tool_id == "quizzes":
                    formats += [("aiken", ".txt"), ("moodle_xml", ".xml")]
                for format, suffix in formats:
                    exported = call(learning_ui.on_export, "XRD " + tool_id, format, tool_id=tool_id)
                    path = Path(exported[control.download._id])
                    assert path.name == "XRD " + tool_id + suffix and path.stat().st_size > 0
                    if format == "json":
                        assert json.loads(path.read_text(encoding="utf-8"))["content"] == current_output["content"]
                    elif format == "aiken":
                        contents = path.read_text(encoding="utf-8")
                        assert "\nB. 26.5 degrees\n" in contents and "\nANSWER: B\n" in contents
                        assert "\nANSWER: A\n" in contents and "The measurement puts" not in contents
                    elif format == "moodle_xml":
                        questions = ElementTree.parse(path).getroot().findall("question")
                        assert len(questions) == 2 and all(question.get("type") == "multichoice" for question in questions)
                        assert questions[0].findall("answer")[1].get("fraction") == "100"
                        assert "26.5 degrees" in questions[0].findtext("generalfeedback/text")
                assert artifacts.load_current_learning(workspace, tool_id) == current_output
                assert artifacts.list_artifacts(workspace, tool_id) == []
            assert len(requests) == local_before
            call(learning_ui.on_quiz, "1", tool_id="check")
            regenerated = call(tools_ui.on_generate, "Test the measurement again.", [document_id], "Default",
                               config.DEFAULT_MODEL, "medium", 2, tool_id="quizzes")
            assert "Completed" in value(regenerated, quizzes.status)
            assert value(regenerated, quizzes.feedback) == ""
            assert value(regenerated, quizzes.score) == value(restarted, quizzes.score)
            assert artifacts.load_current_learning(workspace, "quizzes")["id"] == current_quiz["id"]

            compare = app.rag_tab.tools.controls["compare"]
            evidence_check = app.rag_tab.tools.controls["evidence_check"]
            review_before = len(requests)
            invalid = call(review_ui.on_generate, "Compare sources", "", [document_id], "Default",
                           config.DEFAULT_MODEL, "medium", tool_id="compare")
            assert "2" in value(invalid, compare.status) and len(requests) == review_before
            invalid = call(review_ui.on_generate, "Check claims", "", [document_id], "Default",
                           config.DEFAULT_MODEL, "medium", tool_id="evidence_check")
            assert "claims" in value(invalid, evidence_check.status) and len(requests) == review_before
            second_source = root / "replication.txt"
            second_source.write_text("XRD was measured on sample Alpha. Peak position is 30.0 degrees.\n", encoding="utf-8")
            list(local.ingest([str(second_source)], "Default"))
            review_ids = [document["id"] for document in local.store.list_documents()]
            call(rag_ui.on_refresh, "Default", [document_id], None)
            call(rag_ui.on_source_selection, review_ids)
            claim_lines = "\n".join(["An XRD measurement is reported.", "The only reported peak is 40 degrees.",
                                    "Costs are reported.", "Sample Alpha has a peak at 26.5 degrees."])
            review_results = {}
            review_controls = {"compare": compare, "evidence_check": evidence_check}
            for tool_id, control in review_controls.items():
                generated = call(review_ui.on_generate, "Review the XRD findings.",
                                 claim_lines if tool_id == "evidence_check" else "", review_ids,
                                 "Default", config.DEFAULT_MODEL, "medium", tool_id=tool_id)
                assert "Completed" in value(generated, control.status), generated
                assert "26.5" in str(value(generated, control.table))
                assert "26.5" in value(generated, control.preview)
                current_review = artifacts.load_current_review(workspace, tool_id)
                review_results[tool_id] = current_review
                assert current_review["coverage"]["documents"] == 2
                local_before = len(requests)
                reference = call(review_ui.on_reference, "0", tool_id=tool_id)
                assert "Peak position" in value(reference, control.excerpt)
                for format, suffix in (("md", ".md"), ("csv", ".csv"), ("json", ".json")):
                    exported = call(review_ui.on_export, "XRD " + tool_id, format, tool_id=tool_id)
                    path = Path(exported[control.download._id])
                    assert path.name == "XRD " + tool_id + suffix and path.stat().st_size > 0
                    if format == "json":
                        assert json.loads(path.read_text(encoding="utf-8"))["content"] == current_review["content"]
                    else:
                        assert "replication.txt" in path.read_text(encoding="utf-8-sig").replace("\\.", ".")
                assert artifacts.load_current_review(workspace, tool_id) == current_review
                assert artifacts.list_artifacts(workspace, tool_id) == []
                assert len(requests) == local_before
            assert [item["verdict"] for item in review_results["evidence_check"]["content"]["claims"]] == [
                "supported", "contradicted", "not_found", "mixed"]
            assert [item["text"] for item in review_results["evidence_check"]["content"]["claims"]] == claim_lines.splitlines()
            call(rag_ui.on_source_selection, [document_id])

            before = len(requests)
            plugins = app.plugins_tab
            refreshed = call(plugins_ui.on_refresh)
            plugin_table = refreshed[plugins["catalog"]._id]
            names = [row[0] for row in plugin_table["data"]]
            assert names.index("Presentation") + 1 == names.index("Graphics")
            assert names.index("Graphics") + 1 == names.index("Plot")
            assert names.index("Plot") + 1 == names.index("Mind Map")
            summary_row = next(row for row in plugin_table["data"] if row[0] == "Summary")
            assert type(summary_row[1]) is bool
            summary_row[1] = False
            toggled = call(plugins_ui.on_catalog_input, plugin_table)
            assert "Summary disabled" in toggled[plugins["status"]._id]
            disabled = call(tools_ui.on_generate, "Summary", [document_id], "Default",
                            config.DEFAULT_MODEL, "medium", tool_id="summary")
            assert "disabled" in disabled[summary.status._id]
            assert len(requests) == before
            summary_row[1] = True
            call(plugins_ui.on_catalog_input, plugin_table)
            presentation_row = next(row for row in plugin_table["data"] if row[0] == "Presentation")
            presentation_row[1] = False
            call(plugins_ui.on_catalog_input, plugin_table)
            disabled = call(presentation_ui.on_generate, "Present", 2, [document_id], "Default",
                            config.DEFAULT_MODEL, "medium")
            assert "disabled" in value(disabled, presentation.status) and len(requests) == before
            presentation_row[1] = True
            call(plugins_ui.on_catalog_input, plugin_table)
            graphics_row = next(row for row in plugin_table["data"] if row[0] == "Graphics")
            graphics_row[1] = False
            call(plugins_ui.on_catalog_input, plugin_table)
            disabled = call(graphics_ui.on_generate, *graphic_args)
            assert "disabled" in value(disabled, graphics.status) and len(requests) == before
            graphics_row[1] = True
            call(plugins_ui.on_catalog_input, plugin_table)
            plot_row = next(row for row in plugin_table["data"] if row[0] == "Plot")
            plot_row[1] = False
            call(plugins_ui.on_catalog_input, plugin_table)
            disabled = call(plot_ui.on_generate, "Plot", "bar", [document_id], "Default",
                            config.DEFAULT_MODEL, "medium")
            assert "disabled" in value(disabled, plot.status) and len(requests) == before
            plot_row[1] = True
            call(plugins_ui.on_catalog_input, plugin_table)
            mind_map_row = next(row for row in plugin_table["data"] if row[0] == "Mind Map")
            mind_map_row[1] = False
            call(plugins_ui.on_catalog_input, plugin_table)
            disabled = call(tools_ui.on_generate, "Mind Map", [document_id], "Default",
                            config.DEFAULT_MODEL, "medium", tool_id="mind_maps")
            assert "disabled" in disabled[mind_map.status._id]
            assert len(requests) == before
            mind_map_row[1] = True
            call(plugins_ui.on_catalog_input, plugin_table)
            for tool_id, name, control in (("flashcards", "Flashcards", flashcards), ("quizzes", "Quizzes", quizzes)):
                row = next(row for row in plugin_table["data"] if row[0] == name)
                row[1] = False
                call(plugins_ui.on_catalog_input, plugin_table)
                disabled = call(tools_ui.on_generate, name, [document_id], "Default",
                                config.DEFAULT_MODEL, "medium", tool_id=tool_id)
                assert "disabled" in value(disabled, control.status)
                row[1] = True
                call(plugins_ui.on_catalog_input, plugin_table)
            assert len(requests) == before

            for tool_id, control in review_controls.items():
                row = next(row for row in plugin_table["data"] if row[0] == tools_ui.TOOL_LABELS[tool_id])
                row[1] = False
                call(plugins_ui.on_catalog_input, plugin_table)
                disabled = call(review_ui.on_generate, "Review", claim_lines, review_ids,
                                "Default", config.DEFAULT_MODEL, "medium", tool_id=tool_id)
                assert "disabled" in value(disabled, control.status) and len(requests) == before
                row[1] = True
                call(plugins_ui.on_catalog_input, plugin_table)

            call(rag_ui.on_prepare_workspace_switch)
            call(rag_ui.on_open_workspace, "Other")
            reset = call(tools_ui.on_workspace_reset)
            call(rag_ui.on_finish_workspace_switch)
            assert reset[summary.preview._id]["value"] == ""
            assert reset[summary.filename._id] == "Summary"
            assert reset[spreadsheet.editor._id]["value"]["data"] == []
            assert reset[spreadsheet.filename._id] == "Spreadsheet"
            assert value(reset, presentation.preview) == "" and value(reset, presentation.notes) == ""
            assert value(reset, graphics.preview) is None
            assert value(reset, graphics.caption) == "" and value(reset, graphics.brief) == ""
            assert value(reset, plot.preview) == "" and value(reset, plot.limitations) == ""
            assert value(reset, plot.table)["data"] == []
            for control in review_controls.values():
                assert value(reset, control.preview) == "" and value(reset, control.table)["data"] == []
            assert "Peak measurement" not in str(reset[mind_map.preview._id])
            assert reset[mind_map.filename._id] == "Mind Map"
            assert value(reset, flashcards.preview) == "" and value(reset, flashcards.answer) == ""
            assert value(reset, quizzes.preview) == "" and value(reset, quizzes.feedback) == ""
            call(rag_ui.on_prepare_workspace_switch)
            call(rag_ui.on_open_workspace, workspace)
            restored = call(tools_ui.on_workspace_reset)
            call(rag_ui.on_finish_workspace_switch)
            assert restored[summary.preview._id]["value"] == editor
            assert restored[summary.editor._id]["visible"] is False
            assert restored[spreadsheet.editor._id]["value"]["data"] == grid["data"]
            assert restored[spreadsheet.editor._id]["headers"] == grid["headers"]
            assert "[1]" in restored[spreadsheet.preview._id]
            assert "Measured peak" in value(restored, presentation.preview)
            assert value(restored, presentation.progress) == "Slide 1 of 3"
            restored_slide = call(presentation_ui.on_navigate, None, tool_id="next")
            assert value(restored_slide, presentation.progress) == "Slide 2 of 3"
            restored_deck_export = call(presentation_ui.on_export, "Restored presentation", "json")
            assert Path(restored_deck_export[presentation.download._id]).name == "Restored presentation.json"
            assert "26.5" in value(restored, graphics.caption)
            assert value(restored, graphics.preview) is not None
            restored_graphic_export = call(graphics_ui.on_export, "Restored graphic", "image")
            assert Path(restored_graphic_export[graphics.download._id]).name == "Restored graphic.png"
            restored_graphic_reference = call(graphics_ui.on_reference, "0")
            assert "26.5" in value(restored_graphic_reference, graphics.excerpt)
            assert "data:image/png;base64," in value(restored, plot.preview)
            assert "26.5" in str(value(restored, plot.table))
            restored_plot_export = call(plot_ui.on_export, "Restored plot", "json")
            assert Path(restored_plot_export[plot.download._id]).name == "Restored plot.json"
            restored_plot_reference = call(plot_ui.on_reference, "0")
            assert "26.5" in value(restored_plot_reference, plot.excerpt)
            assert "Peak measurement" in str(restored[mind_map.preview._id])
            assert "Where is the measured peak" in value(restored, flashcards.preview)
            assert value(restored, flashcards.answer) == ""
            assert "Where is the measured peak" in value(restored, quizzes.preview)
            assert value(restored, quizzes.feedback) == ""
            assert value(restored, quizzes.score) == value(restarted, quizzes.score)
            restored_reveal = call(learning_ui.on_flashcard, tool_id="reveal")
            assert "26.5" in value(restored_reveal, flashcards.answer)
            restored_check = call(learning_ui.on_quiz, "1", tool_id="check")
            assert "Correct" in value(restored_check, quizzes.feedback)
            restored_export = call(learning_ui.on_export, "Restored quiz", "json", tool_id="quizzes")
            assert Path(restored_export[quizzes.download._id]).name == "Restored quiz.json"
            assert len(requests) == before

            # A new client has no prior component or session State values. Exercise the
            # page-load callbacks and fetch the actual preview URL a browser receives.
            previous_session = client.session_hash
            client.close()
            client = Client(app.local_url, verbose=False, download_files=False)
            assert client.session_hash != previous_session
            call(rag_ui.on_initial_load)
            reopened = call(tools_ui.on_reset)
            for tool_id, control in review_controls.items():
                assert "26.5" in value(reopened, control.preview)
                assert "26.5" in str(value(reopened, control.table))
                assert artifacts.load_current_review(workspace, tool_id) == review_results[tool_id]
                reference = call(review_ui.on_reference, "0", tool_id=tool_id)
                assert "Peak position" in value(reference, control.excerpt)
            verify_graphic_preview(reopened)
            verify_plot_preview(reopened)
            # Opening either visual tab rehydrates only its persisted output.
            reopened_graphics_tab = call(tools_ui.on_visual_tool_open, tool_id="graphics")
            assert "Restored" in value(reopened_graphics_tab, graphics.status)
            verify_graphic_preview(reopened_graphics_tab)
            reopened_plot_tab = call(tools_ui.on_visual_tool_open, tool_id="plot")
            assert "Restored" in value(reopened_plot_tab, plot.status)
            verify_plot_preview(reopened_plot_tab)
            reopened_graphic_reference = call(graphics_ui.on_reference, "0")
            assert "26.5" in value(reopened_graphic_reference, graphics.excerpt)
            reopened_plot_reference = call(plot_ui.on_reference, "0")
            assert "26.5" in value(reopened_plot_reference, plot.excerpt)
            assert artifacts.load_current_graphics(workspace) == current_graphic
            assert artifacts.load_current_plot(workspace) == current_plot
            assert len(requests) == before, "Restoration and preview must never regenerate results."

            # The reported Iris workflow goes through real indexing, the queued
            # Plot UI callback, rendering, persistence and download serving.
            iris_path = root / "IrisFlower.csv"
            iris_rows = [[index + 1, round([4.9, 5.9, 6.7][index % 3] + (index // 3 % 5) / 10, 1),
                          round([3.3, 2.7, 3.0][index % 3] + (index // 3 % 5) / 10, 1),
                          ["setosa", "versicolor", "virginica"][index % 3]] for index in range(150)]
            with iris_path.open("w", encoding="utf-8", newline="") as handle:
                writer = csv.writer(handle)
                writer.writerow(["Id", "sepal_length", "sepal_width", "species"])
                writer.writerows(iris_rows)
            list(local.ingest([str(iris_path)], "Default"))
            iris_id = next(document["id"] for document in local.store.list_documents()
                           if document["name"] == "IrisFlower.csv")
            call(rag_ui.on_refresh, "Default", [document_id], None)
            call(rag_ui.on_source_selection, [iris_id])
            iris_result = call(plot_ui.on_generate,
                "Please help me make a scatter plot for this dataset, x-axis is sepal length and "
                "y-axis is sepal width, colored by species", "auto", [iris_id], "Default",
                config.DEFAULT_MODEL, "medium")
            assert "Completed" in value(iris_result, plot.status), iris_result
            assert "data:image/png;base64," in value(iris_result, plot.preview)
            assert len(value(iris_result, plot.table)["data"]) == 150
            assert requests[before:] == ["tool_plot_columns"]
            iris_saved = artifacts.load_current_plot(workspace)
            assert iris_saved["id"] == current_plot["id"]
            assert iris_saved["settings"]["data_mode"] == "direct_table"
            assert iris_saved["settings"]["row_count"] == 150
            assert Counter((series["name"], point["x"], point["y"])
                           for series in iris_saved["content"]["series"] for point in series["points"]) == Counter(
                               (row[3], row[1], row[2]) for row in iris_rows)
            iris_export = call(plot_ui.on_export, "Iris scatter", "json")
            download_file = value(iris_export, plot.download)
            response = httpx.get(urljoin(app.local_url, download_file["url"]), timeout=15)
            response.raise_for_status()
            assert response.json() == iris_saved
            iris_reopened = call(tools_ui.on_visual_tool_open, tool_id="plot")
            assert "data:image/png;base64," in value(iris_reopened, plot.preview)
            assert len(value(iris_reopened, plot.table)["data"]) == 150
            assert requests[before:] == ["tool_plot_columns"]
            print("Studio RPC smoke passed: generation, citations, Summary/edit/restore, Spreadsheet/save/restore, Presentation/slides/restore, Graphics/image/HTTP restore, Plot/data/PNG restore, 150-row Iris CSV/single-mapping/species/export/restore, fresh-session and tab-reopen restore, Mind Map/zoom/restore, Flashcards/reveal/restore, Quizzes/score/restart/restore, Compare/matrix/restore, Evidence Check/claims/quotes/restore, exports, checkbox plugins, workspace isolation.")
        finally:
            if client is not None:
                client.close()
            app.close()
            for service in services.values():
                service.store.close()


if __name__ == "__main__":
    main()
