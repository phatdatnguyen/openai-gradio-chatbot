"""Source-location and OCR regressions; never call a model or use credentials."""

import os
import struct
import sys
import uuid
from pathlib import Path

import pytest
import tiktoken
from docx import Document
from openpyxl import Workbook
from PIL import Image
from pptx import Presentation
from pptx.util import Inches
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from chatbot.readers import EmptyDocumentError
from rag_chatbot import ingestion
from rag_chatbot.schemas import SourceElement


def write_pdf(path, page_texts):
    writer = PdfWriter()
    for text in page_texts:
        page = writer.add_blank_page(width=200, height=200)
        if text is None:
            continue
        font = DictionaryObject({
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        })
        page[NameObject("/Resources")] = DictionaryObject({
            NameObject("/Font"): DictionaryObject({NameObject("/F1"): font}),
        })
        content = DecodedStreamObject()
        content.set_data(b"BT /F1 12 Tf 20 100 Td (" + text.encode("ascii") + b") Tj ET")
        page[NameObject("/Contents")] = content
    with path.open("wb") as file:
        writer.write(file)


def test_pdf_keeps_each_page_and_preserves_short_text(tmp_path):
    path = tmp_path / "paper.PDF"
    write_pdf(path, ["Short", "Different page"])
    elements = ingestion.extract_elements(path)
    assert [(element.locator, element.text) for element in elements] == [
        ("Page 1", "Short"), ("Page 2", "Different page"),
    ]


def test_mixed_pdf_cannot_silently_skip_scanned_pages(tmp_path):
    path = tmp_path / "mixed.pdf"
    write_pdf(path, ["Native text", None])
    with pytest.raises(ingestion.IngestionError, match=r"page\(s\) 2.*Enable image/OCR"):
        ingestion.extract_elements(path)
    assert not (tmp_path / ".mixed.pdf.rag-assets").exists()


def test_mixed_pdf_renders_only_textless_pages_for_vision(tmp_path):
    path = tmp_path / "mixed.pdf"
    write_pdf(path, ["Native text", None])
    original = path.read_bytes()
    calls = []

    def describe(image_path):
        calls.append(Path(image_path))
        with Image.open(image_path) as image:
            assert image.format == "PNG"
            assert max(image.size) <= ingestion.MAX_IMAGE_EDGE
        return "OCR: result 42. Figure: a chart."

    elements = ingestion.extract_elements(path, vision_describe=describe)
    assert len(calls) == 1
    assert elements[0] == SourceElement("Native text", "Page 1")
    assert elements[1].locator == "Page 2"
    assert elements[1].kind == "image"
    assert elements[1].image_path == str(calls[0])
    assert "result 42" in elements[1].text
    assert calls[0].is_file()
    assert calls[0].parent.parent == tmp_path
    assert path.read_bytes() == original


def test_images_require_explicit_vision_processing(tmp_path):
    path = tmp_path / "scan.png"
    Image.new("RGB", (4, 4)).save(path)
    with pytest.raises(ingestion.IngestionError, match="Enable it"):
        ingestion.extract_elements(path)


def test_image_normalization_keeps_original_and_uses_white_transparency(tmp_path):
    path = tmp_path / "transparent.png"
    Image.new("RGBA", (5, 5), (0, 0, 0, 0)).save(path)
    original = path.read_bytes()

    def describe(image_path):
        with Image.open(image_path) as image:
            assert image.mode == "RGB"
            assert image.getpixel((0, 0)) == (255, 255, 255)
        return "A transparent source image."

    elements = ingestion.extract_elements(path, vision_describe=describe)
    assert elements[0].locator == "Image"
    assert elements[0].kind == "image"
    assert Path(elements[0].image_path).is_file()
    assert path.read_bytes() == original


def test_multiframe_tiff_keeps_frame_locations(tmp_path):
    path = tmp_path / "scans.tiff"
    first = Image.new("RGB", (4, 4), "red")
    second = Image.new("RGB", (4, 4), "blue")
    first.save(path, save_all=True, append_images=[second])
    elements = ingestion.extract_elements(path, vision_describe=lambda image_path: Path(image_path).stem)
    assert [element.locator for element in elements] == ["Frame 1", "Frame 2"]
    assert len({element.image_path for element in elements}) == 2


@pytest.mark.parametrize("description", [None, "", "   ", ["invalid"]])
def test_empty_or_invalid_vision_output_fails_instead_of_indexing_empty_data(tmp_path, description):
    path = tmp_path / "scan.png"
    Image.new("RGB", (4, 4)).save(path)
    with pytest.raises(ingestion.IngestionError, match="no usable text"):
        ingestion.extract_elements(path, vision_describe=lambda _: description)


def test_docx_keeps_sections_nested_tables_and_paragraph_locations(tmp_path):
    path = tmp_path / "report.docx"
    document = Document()
    document.add_heading("Results", level=1)
    document.add_heading("First experiment", level=3)
    document.add_paragraph("Original experiment")
    document.add_heading("Second experiment", level=3)
    document.add_paragraph("Observed result")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Accuracy"
    table.cell(0, 1).text = "98%"
    table.cell(0, 1).add_table(rows=1, cols=1).cell(0, 0).text = "Nested note"
    document.save(path)

    elements = ingestion.extract_elements(path)
    observed = next(element for element in elements if "Observed result" in element.text)
    assert "Results / Second experiment" in observed.locator
    assert "First experiment" not in observed.locator
    table_element = next(element for element in elements if element.kind == "table")
    assert "Table 1" in table_element.locator and "row 1" in table_element.locator
    assert "Accuracy" in table_element.text and "98%" in table_element.text
    assert "Nested note" in table_element.text


def test_docx_embedded_rasters_are_opt_in_located_and_described_once(tmp_path):
    picture = tmp_path / "figure.png"
    Image.new("RGB", (5, 5), "red").save(picture)
    document = Document()
    document.add_heading("Results", level=1)
    document.add_paragraph("Native explanation")
    document.add_picture(str(picture))
    table = document.add_table(rows=1, cols=1)
    table.cell(0, 0).paragraphs[0].add_run().add_picture(str(picture))
    nested = table.cell(0, 0).add_table(rows=1, cols=1)
    nested.cell(0, 0).paragraphs[0].add_run().add_picture(str(picture))
    path = tmp_path / "illustrated.docx"
    document.save(path)
    assert all(element.kind != "image" for element in ingestion.extract_elements(path))
    calls = []

    def describe(image_path):
        calls.append(image_path)
        return "A red scientific diagram."

    elements = ingestion.extract_elements(path, vision_describe=describe)
    pictures = [element for element in elements if element.kind == "image"]
    assert len(pictures) == 3
    assert len(calls) == 1
    assert len({element.locator for element in pictures}) == 3
    assert all("Section Results" in element.locator for element in pictures)
    assert any("Table 1" in element.locator for element in pictures)
    assert any("nested table" in element.locator for element in pictures)
    assert all(Path(element.image_path).is_file() for element in pictures)


def test_xlsx_keeps_sheet_names_physical_cells_leading_zero_strings_and_formulas(tmp_path):
    path = tmp_path / "measurements.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Measurements"
    sheet["A3"], sheet["B3"] = "Sample ID", "Result"
    sheet["A4"], sheet["B4"] = "00123", 42
    sheet["B5"] = "=B4*2"
    workbook.create_sheet("Notes")["C7"] = "Keep this note"
    workbook.save(path)

    elements = ingestion.extract_elements(path)
    measured = next(element for element in elements if "Measurements" in element.locator)
    assert "rows 3-5" in measured.locator
    assert "A4: 00123" in measured.text and "B4: 42" in measured.text
    assert "B5: =B4*2 (cached value unavailable)" in measured.text
    note = next(element for element in elements if "Notes" in element.locator)
    assert "C7: Keep this note" in note.text


def test_xlsx_row_groups_keep_column_context_without_losing_rows(tmp_path, monkeypatch):
    monkeypatch.setattr(ingestion, "TABLE_ROWS_PER_ELEMENT", 2)
    path = tmp_path / "rows.xlsx"
    workbook = Workbook()
    for row in [("Quantity",), (10,), (20,), (30,)]:
        workbook.active.append(row)
    workbook.save(path)
    elements = ingestion.extract_elements(path)
    assert len(elements) == 2
    assert "rows 3-4" in elements[1].locator
    assert "Quantity" in elements[1].text
    assert "A3: 20" in elements[1].text and "A4: 30" in elements[1].text


def test_xls_keeps_legacy_cell_coordinates(tmp_path):
    def record(code, payload):
        return struct.pack("<HH", code, len(payload)) + payload

    path = tmp_path / "legacy.xls"
    path.write_bytes(
        record(0x0009, struct.pack("<HH", 7, 0x0010))
        + record(0x0042, struct.pack("<H", 1252))
        + record(0x0000, struct.pack("<HHHH", 0, 2, 0, 1))
        + record(0x0004, struct.pack("<HH3sB", 0, 0, b"\x00" * 3, 7) + b"Revenue")
        + record(0x0003, struct.pack("<HH3sd", 1, 0, b"\x00" * 3, 123.0))
        + record(0x000A, b"")
    )
    elements = ingestion.extract_elements(path)
    assert "A1: Revenue" in elements[0].text
    assert "A2: 123" in elements[0].text


def test_pptx_keeps_slides_grouped_text_tables_and_speaker_notes(tmp_path):
    path = tmp_path / "presentation.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    table = slide.shapes.add_table(1, 2, Inches(1), Inches(1), Inches(4), Inches(1)).table
    table.cell(0, 0).text = "Revenue"
    table.cell(0, 1).text = "$456"
    group = slide.shapes.add_group_shape()
    group.shapes.add_textbox(Inches(1), Inches(3), Inches(4), Inches(1)).text = "Grouped annotation"
    slide.notes_slide.notes_text_frame.text = "Explain the experimental limits"
    presentation.slides.add_slide(presentation.slide_layouts[6]).shapes.add_textbox(
        Inches(1), Inches(1), Inches(4), Inches(1),
    ).text = "Second slide"
    presentation.save(path)

    elements = ingestion.extract_elements(path)
    assert any(element.locator.startswith("Slide 1") and "Grouped annotation" in element.text for element in elements)
    assert any(element.kind == "table" and "$456" in element.text for element in elements)
    assert any(element.kind == "notes" and "experimental limits" in element.text for element in elements)
    assert any(element.locator.startswith("Slide 2") and "Second slide" in element.text for element in elements)


def test_pptx_embedded_rasters_keep_slide_and_group_locations(tmp_path):
    picture = tmp_path / "diagram.png"
    Image.new("RGB", (5, 5), "blue").save(picture)
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.shapes.add_textbox(Inches(1), Inches(1), Inches(2), Inches(1)).text = "Native text"
    slide.shapes.add_picture(str(picture), Inches(1), Inches(2))
    group = slide.shapes.add_group_shape()
    group.shapes.add_picture(str(picture), Inches(3), Inches(3))
    path = tmp_path / "illustrated.pptx"
    presentation.save(path)
    assert len(ingestion.extract_elements(path)) == 1
    calls = []

    def describe(image_path):
        calls.append(image_path)
        return "A blue diagram"

    elements = ingestion.extract_elements(path, vision_describe=describe)
    pictures = [element for element in elements if element.kind == "image"]
    assert len(pictures) == 2 and len(calls) == 1
    assert all(element.locator.startswith("Slide 1") for element in pictures)
    assert len({element.locator for element in pictures}) == 2
    assert all(Path(element.image_path).is_file() for element in pictures)


def test_office_numeric_table_rows_keep_label_context(tmp_path):
    document = Document()
    document.add_heading("Trial A", level=1)
    table = document.add_table(rows=2, cols=1)
    table.cell(0, 0).text = "Accuracy"
    table.cell(1, 0).text = "98%"
    path = tmp_path / "table.docx"
    document.save(path)
    element = next(element for element in ingestion.extract_elements(path) if "row 2" in element.locator)
    assert "Accuracy" in element.text and "98%" in element.text and "Trial A" in element.text

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = "Trial B"
    table = slide.shapes.add_table(2, 1, Inches(1), Inches(2), Inches(4), Inches(2)).table
    table.cell(0, 0).text = "Latency"
    table.cell(1, 0).text = "12 ms"
    path = tmp_path / "table.pptx"
    presentation.save(path)
    element = next(element for element in ingestion.extract_elements(path) if "table row 2" in element.locator)
    assert "Latency" in element.text and "12 ms" in element.text and "Trial B" in element.text


@pytest.mark.parametrize("suffix, body, expected", [
    (".txt", "Plain text", "Plain text"),
    (".md", "# Heading\n\nMarkdown body", "Markdown body"),
    (".html", "<h1>Title</h1><p>Page body</p>", "Page body"),
    (".csv", "Sample,Value\n00123,42", "A2: 00123"),
    (".tsv", "Sample\tValue\n00123\t42", "B2: 42"),
])
def test_text_and_tabular_formats(tmp_path, suffix, body, expected):
    path = tmp_path / f"source{suffix}"
    path.write_text(body, encoding="utf-8-sig")
    assert expected in ingestion.extract_elements(path)[0].text


def test_empty_source_and_unsupported_legacy_office_fail_clearly(tmp_path):
    empty = tmp_path / "empty.txt"
    empty.write_text("   ", encoding="utf-8")
    with pytest.raises(EmptyDocumentError):
        ingestion.extract_elements(empty)
    legacy = tmp_path / "legacy.doc"
    legacy.write_bytes(b"not a supported modern document")
    with pytest.raises(ingestion.IngestionError, match="DOCX/PPTX"):
        ingestion.extract_elements(legacy)


def test_pdf_page_limit_is_checked_before_ocr(tmp_path, monkeypatch):
    monkeypatch.setattr(ingestion, "MAX_PAGES", 1)
    path = tmp_path / "too-many.pdf"
    write_pdf(path, [None, None])
    with pytest.raises(ingestion.IngestionError, match="limited to 1 pages"):
        ingestion.extract_elements(path, vision_describe=lambda _: pytest.fail("OCR must not run"))


def test_file_and_office_expansion_limits(tmp_path, monkeypatch):
    path = tmp_path / "large.txt"
    path.write_text("a" * 20, encoding="utf-8")
    monkeypatch.setattr(ingestion, "MAX_FILE_BYTES", 10)
    with pytest.raises(ingestion.IngestionError, match="Files are limited"):
        ingestion.extract_elements(path)
    monkeypatch.setattr(ingestion, "MAX_FILE_BYTES", 1_000_000)
    monkeypatch.setattr(ingestion, "MAX_EXPANDED_OFFICE_BYTES", 10)
    path = tmp_path / "expanded.docx"
    Document().save(path)
    with pytest.raises(ingestion.IngestionError, match="expands beyond"):
        ingestion.extract_elements(path)


def test_workbook_and_image_limits(tmp_path, monkeypatch):
    monkeypatch.setattr(ingestion, "MAX_SPREADSHEET_CELLS", 3)
    workbook = Workbook()
    workbook.active["B2"] = "fourth cell"
    path = tmp_path / "cells.xlsx"
    workbook.save(path)
    with pytest.raises(ingestion.IngestionError, match="limited to 3 cells"):
        ingestion.extract_elements(path)
    monkeypatch.setattr(ingestion, "MAX_IMAGE_PIXELS", 3)
    path = tmp_path / "pixels.png"
    Image.new("RGB", (2, 2)).save(path)
    with pytest.raises(ingestion.IngestionError, match="pixel limit"):
        ingestion.extract_elements(path, vision_describe=lambda _: pytest.fail("OCR must not run"))


def test_chunking_is_token_bounded_stable_and_keeps_source_identity():
    original = "Tiếng Việt: nghiên cứu dữ liệu. " * 80 + "<|endoftext|>"
    elements = [
        SourceElement(original, "Page 3", "image", "/managed/page-3.png"),
        SourceElement("Separate table", "Sheet Data · rows 1-2", "table"),
    ]
    chunks = ingestion.chunk_elements(elements, "document-a", chunk_size=40, chunk_overlap=8)
    encoding = tiktoken.get_encoding("cl100k_base")
    assert len(chunks) > 3
    assert all(len(encoding.encode(chunk.text, disallowed_special=())) <= 40 for chunk in chunks)
    assert all(chunk.document_id == "document-a" for chunk in chunks)
    assert all(uuid.UUID(chunk.id).version == 5 for chunk in chunks)
    assert chunks == ingestion.chunk_elements(elements, "document-a", chunk_size=40, chunk_overlap=8)
    assert all(chunk.kind == "image" and chunk.image_path == "/managed/page-3.png" for chunk in chunks[:-1])
    assert chunks[-1].locator == "Sheet Data · rows 1-2"
    assert chunks[-1].kind == "table"
    assert elements[0].text == original
    other = ingestion.chunk_elements(elements, "document-b", chunk_size=40, chunk_overlap=8)
    assert {chunk.id for chunk in chunks}.isdisjoint(chunk.id for chunk in other)


@pytest.mark.parametrize("size, overlap", [(0, 0), (-1, 0), (10, -1), (10, 10), (10, 11)])
def test_invalid_chunk_settings_fail_clearly(size, overlap):
    with pytest.raises(ValueError):
        ingestion.chunk_elements([], "doc", chunk_size=size, chunk_overlap=overlap)
