"""Actual document fixtures exercise safe local preview without Office/API calls."""

from concurrent.futures import ThreadPoolExecutor
import json
from pathlib import Path
import struct
import subprocess
import sys
import zipfile

from docx import Document
from openpyxl import Workbook
from PIL import Image
from pptx import Presentation
from pptx.util import Inches
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject
import pytest

from rag_chatbot import viewer


@pytest.fixture
def library(tmp_path):
    root = tmp_path / "Alpha"
    (root / "sources" / "document-id").mkdir(parents=True)
    return root


@pytest.fixture(autouse=True)
def no_office_launch(monkeypatch):
    from rag_chatbot import office_preview

    def unavailable(*args, **kwargs):
        raise office_preview.PreviewUnavailable("No Office converter is available in this test.")

    monkeypatch.setattr(office_preview, "get_pdf_preview", unavailable)


def source(library, name):
    return library / "sources" / "document-id" / name


def pdf(path, pages=2):
    writer = PdfWriter()
    for number in range(pages):
        page = writer.add_blank_page(width=200 + number * 100, height=200)
        font = DictionaryObject({NameObject("/Type"): NameObject("/Font"),
                                 NameObject("/Subtype"): NameObject("/Type1"),
                                 NameObject("/BaseFont"): NameObject("/Helvetica")})
        page[NameObject("/Resources")] = DictionaryObject({
            NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)})})
        content = DecodedStreamObject()
        content.set_data(f"BT /F1 16 Tf 20 100 Td (Readable page {number + 1}) Tj ET".encode())
        page[NameObject("/Contents")] = writer._add_object(content)
    with path.open("wb") as handle:
        writer.write(handle)
    return path


def test_module_import_does_not_initialize_format_libraries_or_storage():
    command = "import sys,json; import rag_chatbot.viewer; print(json.dumps([name for name in ('pypdfium2','openpyxl','docx','pptx','win32com','rag_chatbot.service') if name in sys.modules]))"
    result = subprocess.run([sys.executable, "-c", command], check=True, capture_output=True, text=True)
    assert json.loads(result.stdout) == []


def test_pdf_page_navigation_renders_owned_pixels_at_bounded_size(library, monkeypatch):
    path = pdf(source(library, "study.PDF"))
    original = path.read_bytes()
    monkeypatch.setattr(viewer, "MAX_IMAGE_EDGE", 240)
    first = viewer.preview_document(path, data_dir=library)
    second = viewer.preview_document(path, data_dir=library, unit=2)
    assert first.unit == 1 and first.unit_count == 2 and first.unit_label == "Page"
    assert second.unit == 2 and second.image.width > second.image.height
    assert max(first.image.size) <= 240 and max(second.image.size) <= 240
    assert first.image.mode == "RGB"
    assert first.image.getextrema()[0][0] < 255  # rendered text, not an empty page
    assert path.read_bytes() == original
    first.image.close()
    second.image.close()
    clamped = viewer.preview_document(path, data_dir=library, unit=999)
    assert clamped.unit == 2
    clamped.image.close()


def test_pdf_concurrent_readers_share_render_lock(library):
    path = pdf(source(library, "parallel.pdf"))
    with ThreadPoolExecutor(max_workers=4) as pool:
        results = list(pool.map(lambda page: viewer.preview_document(path, data_dir=library, unit=page), [1, 2] * 4))
    assert all(result.image is not None for result in results)
    for result in results:
        result.image.close()


def test_encrypted_damaged_or_too_many_pdf_pages_give_safe_status(library, monkeypatch):
    path = source(library, "locked.pdf")
    writer = PdfWriter()
    writer.add_blank_page(100, 100)
    writer.encrypt("secret")
    with path.open("wb") as handle:
        writer.write(handle)
    result = viewer.preview_document(path, data_dir=library)
    assert result.image is None and "password-protected" in result.status
    path.write_bytes(b"not a PDF")
    assert "could not be rendered" in viewer.preview_document(path, data_dir=library).status
    pdf(path)
    monkeypatch.setattr(viewer, "MAX_PAGES", 1)
    assert "1 to 1 pages" in viewer.preview_document(path, data_dir=library).status


def test_images_normalize_exif_transparency_and_keep_original(library):
    path = source(library, "rotated.jpg")
    image = Image.new("RGB", (40, 20), "blue")
    exif = Image.Exif()
    exif[274] = 6
    image.save(path, exif=exif)
    before = path.read_bytes()
    result = viewer.preview_document(path, data_dir=library)
    assert result.image.size == (20, 40) and result.image.mode == "RGB"
    assert path.read_bytes() == before
    result.image.close()
    path = source(library, "transparent.png")
    Image.new("RGBA", (10, 10), (0, 0, 0, 0)).save(path)
    result = viewer.preview_document(path, data_dir=library)
    assert result.image.getpixel((0, 0)) == (255, 255, 255)
    result.image.close()


def test_multiframe_images_and_pixel_frame_bounds(library, monkeypatch):
    path = source(library, "frames.tiff")
    with Image.new("RGB", (10, 10), "red") as red, Image.new("RGB", (10, 10), "blue") as blue:
        red.save(path, save_all=True, append_images=[blue])
    result = viewer.preview_document(path, data_dir=library, unit=2)
    assert result.unit_count == 2 and result.unit_label == "Frame"
    assert result.image.getpixel((0, 0)) == (0, 0, 255)
    result.image.close()
    monkeypatch.setattr(viewer, "MAX_IMAGE_FRAMES", 1)
    assert "at most 1 frames" in viewer.preview_document(path, data_dir=library).status
    monkeypatch.setattr(viewer, "MAX_IMAGE_FRAMES", 100)
    monkeypatch.setattr(viewer, "MAX_IMAGE_PIXELS", 99)
    assert "99 pixels" in viewer.preview_document(path, data_dir=library).status


def test_xlsx_selects_sheet_and_paginates_physical_rows_without_running_formulas(library):
    path = source(library, "results.xlsx")
    book = Workbook()
    book.active.title = "Measurements"
    book.active.append(["Identifier", "Value", "Formula"])
    for number in range(1, 55):
        book.active.append([f"{number:05}", number, f"=B{number + 1}*2"])
    notes = book.create_sheet("Notes")
    notes["C7"] = "This sheet is different."
    book.save(path)
    book.close()
    first = viewer.preview_document(path, data_dir=library)
    second = viewer.preview_document(path, data_dir=library, unit=2)
    assert first.sheets == ["Measurements", "Notes"] and first.sheet == "Measurements"
    assert first.headers == ["Row", "A", "B", "C"]
    assert first.table[1] == [2, "00001", "1", "=B2*2"]
    assert first.unit_count == 2 and len(first.table) == 50
    assert second.table[0][0] == 51 and len(second.table) == 5
    note_preview = viewer.preview_document(path, data_dir=library, sheet="Notes", unit=9)
    assert note_preview.unit == 1 and note_preview.table[6] == [7, "", "", "This sheet is different."]
    invalid = viewer.preview_document(path, data_dir=library, sheet="Missing")
    assert not invalid.table and "available worksheet" in invalid.status


def test_xlsx_limits_rows_columns_and_cell_strings(library, monkeypatch):
    path = source(library, "wide.xlsx")
    book = Workbook()
    for number in range(5):
        book.active.append(["longer than limit", number, "extra column"])
    book.save(path)
    monkeypatch.setattr(viewer, "MAX_SHEET_ROWS", 3)
    monkeypatch.setattr(viewer, "MAX_SHEET_COLUMNS", 2)
    monkeypatch.setattr(viewer, "ROWS_PER_PAGE", 2)
    monkeypatch.setattr(viewer, "MAX_CELL_CHARS", 4)
    result = viewer.preview_document(path, data_dir=library, unit=999)
    assert result.unit == 2 and result.unit_count == 2
    assert result.table == [[3, "long…", "2"]]
    assert result.headers == ["Row", "A", "B"]
    assert "limited to 3 rows and 2 columns" in result.status


def test_xlsx_streaming_export_without_dimensions_still_has_bounded_visible_cells(library, monkeypatch):
    path = source(library, "streaming.xlsx")
    book = Workbook(write_only=True)
    sheet = book.create_sheet("Streamed")
    for row in (("Identifier", "Value"), ("00123", 42), ("00456", 51), ("00789", 63)):
        sheet.append(row)
    book.save(path)
    result = viewer.preview_document(path, data_dir=library)
    assert result.headers == ["Row", "A", "B"]
    assert result.table[1] == [2, "00123", "42"]
    monkeypatch.setattr(viewer, "MAX_SHEET_ROWS", 2)
    limited = viewer.preview_document(path, data_dir=library)
    assert len(limited.table) == 2 and "at least 3" in limited.status


def test_legacy_xls_preview_reads_actual_biff_data(library):
    def record(code, payload):
        return struct.pack("<HH", code, len(payload)) + payload

    path = source(library, "legacy.xls")
    path.write_bytes(record(0x0009, struct.pack("<HH", 7, 0x0010))
                     + record(0x0042, struct.pack("<H", 1252))
                     + record(0x0000, struct.pack("<HHHH", 0, 2, 0, 1))
                     + record(0x0004, struct.pack("<HH3sB", 0, 0, b"\x00" * 3, 7) + b"Revenue")
                     + record(0x0003, struct.pack("<HH3sd", 1, 0, b"\x00" * 3, 123.0))
                     + record(0x000A, b""))
    result = viewer.preview_document(path, data_dir=library)
    assert result.sheets and result.headers == ["Row", "A"]
    assert result.table == [[1, "Revenue"], [2, "123.0"]]


def test_corrupt_compound_xls_reports_preview_status(library):
    path = source(library, "corrupt.xls")
    path.write_bytes(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 504)
    result = viewer.preview_document(path, data_dir=library)
    assert not result.table and "spreadsheet could not be opened" in result.status


def test_csv_and_tsv_keep_headers_as_real_rows_and_bound_preview(library, monkeypatch):
    path = source(library, "values.csv")
    path.write_text('"Identifier","Value"\n"00123","=HYPERLINK(\"\"https://example.com\"\")"\n"2","3"\n', encoding="utf-8")
    monkeypatch.setattr(viewer, "MAX_SHEET_ROWS", 2)
    result = viewer.preview_document(path, data_dir=library)
    assert result.table[0] == [1, "Identifier", "Value"]
    assert result.table[1][1] == "00123"
    assert "=HYPERLINK" in result.table[1][2] and "CSV preview limit" in result.status
    path = source(library, "values.tsv")
    path.write_text("left\tright\n1\t2", encoding="utf-8")
    assert viewer.preview_document(path, data_dir=library).table[1] == [2, "1", "2"]


def test_docx_structural_preview_escapes_text_tables_and_renders_only_local_images(library, monkeypatch):
    path = source(library, "article.docx")
    image = source(library, "image.png")
    Image.new("RGB", (20, 10), "green").save(image)
    document = Document()
    document.add_heading("Scientific <heading>", level=1)
    document.add_paragraph('<script>alert("bad")</script> & evidence')
    document.add_table(rows=1, cols=2).cell(0, 0).text = '<img src="https://bad.example/track">'
    document.add_picture(str(image))
    document.save(path)
    result = viewer.preview_document(path, data_dir=library)
    assert result.image is None and result.unit_label == "Section page"
    assert "<h1>Scientific &lt;heading&gt;</h1>" in result.html
    assert "&lt;script&gt;" in result.html and "<script>" not in result.html
    assert '<img src="https://' not in result.html
    assert 'src="data:image/jpeg;base64,' in result.html
    assert "<table" in result.html and "do not match printed page" in result.status
    monkeypatch.setattr(viewer, "BLOCKS_PER_PAGE", 1)
    second = viewer.preview_document(path, data_dir=library, unit=2)
    assert second.unit_count == 4 and "&lt;script&gt;" in second.html and "Scientific" not in second.html


def test_pptx_fallback_keeps_slide_navigation_grouped_text_and_images(library):
    path = source(library, "slides.pptx")
    image = source(library, "chart.png")
    Image.new("RGB", (20, 10), "blue").save(image)
    document = Presentation()
    slide = document.slides.add_slide(document.slide_layouts[6])
    group = slide.shapes.add_group_shape()
    group.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(1)).text = "Grouped <script>bad</script>"
    group.shapes.add_picture(str(image), Inches(1), Inches(2))
    table = slide.shapes.add_table(1, 2, Inches(1), Inches(3), Inches(4), Inches(1)).table
    table.cell(0, 0).text = "Metric"
    table.cell(0, 1).text = "42"
    document.slides.add_slide(document.slide_layouts[6]).shapes.add_textbox(
        Inches(1), Inches(1), Inches(3), Inches(1)).text = "Second slide"
    document.save(path)
    first = viewer.preview_document(path, data_dir=library)
    assert first.unit_count == 2 and first.unit_label == "Slide"
    assert "Grouped &lt;script&gt;bad&lt;/script&gt;" in first.html
    assert "data:image/jpeg;base64," in first.html and "<td" in first.html
    second = viewer.preview_document(path, data_dir=library, unit=2)
    assert "Second slide" in second.html and "Grouped" not in second.html


@pytest.mark.parametrize("suffix", [".doc", ".docx", ".ppt", ".pptx"])
def test_office_converter_output_is_rendered_as_pdf_when_available(library, monkeypatch, suffix):
    from rag_chatbot import office_preview

    path = source(library, "document" + suffix)
    if suffix == ".docx":
        Document().save(path)
    elif suffix == ".pptx":
        Presentation().save(path)
    else:
        path.write_bytes(b"legacy Office placeholder")
    converted = pdf(source(library, "converted.pdf"))
    calls = []

    def convert(candidate, *, data_dir):
        calls.append((candidate, data_dir))
        return converted

    monkeypatch.setattr(office_preview, "get_pdf_preview", convert)
    result = viewer.preview_document(path, data_dir=library, unit=2)
    assert result.image is not None and result.unit == 2 and result.unit_count == 2
    assert "Converted Office preview" in result.status and calls == [(path, library)]
    result.image.close()


@pytest.mark.parametrize("suffix", [".doc", ".ppt"])
def test_legacy_office_without_converter_has_clear_download_fallback(library, suffix):
    path = source(library, "legacy" + suffix)
    path.write_bytes(b"legacy Office placeholder")
    result = viewer.preview_document(path, data_dir=library)
    assert not result.html and result.image is None
    assert "Office preview unavailable" in result.status and "Download the original" in result.status


def test_office_converted_output_cannot_escape_workspace(library, tmp_path, monkeypatch):
    from rag_chatbot import office_preview

    path = source(library, "article.docx")
    Document().save(path)
    foreign = pdf(tmp_path / "foreign.pdf")
    monkeypatch.setattr(office_preview, "get_pdf_preview", lambda *args, **kwargs: foreign)
    result = viewer.preview_document(path, data_dir=library)
    assert result.image is None and "workspace's sources" in result.status


def test_raw_html_is_plain_text_without_scripts_remote_images_or_styles(library):
    path = source(library, "website.html")
    path.write_text('<h1>Study &amp; results</h1><script>alert("bad")</script><style>body{color:red}</style>'
                    '<p>Measured value 42.</p><img src="https://tracker.example/pixel">'
                    '<iframe src="https://example.com"></iframe>', encoding="utf-8")
    result = viewer.preview_document(path, data_dir=library)
    assert "Study & results" in result.text and "Measured value 42" in result.text
    assert "alert(" not in result.text and "color:red" not in result.text
    assert not result.html and not result.image and "https://" not in result.text


def test_text_pages_decode_unicode_and_honor_byte_character_limits(library, monkeypatch):
    path = source(library, "notes.txt")
    path.write_text("αβγδεζηθικλ", encoding="utf-16")
    monkeypatch.setattr(viewer, "TEXT_CHARS_PER_PAGE", 4)
    monkeypatch.setattr(viewer, "MAX_TEXT_CHARS", 9)
    result = viewer.preview_document(path, data_dir=library, unit=2)
    assert result.text == "εζηθ" and result.unit_count == 3
    assert "text preview limit" in result.status
    monkeypatch.setattr(viewer, "MAX_TEXT_BYTES", 12)
    result = viewer.preview_document(path, data_dir=library, unit=1)
    assert "text preview limit" in result.status


def test_sources_only_containment_rejects_sibling_workspace_and_chat_files(library, tmp_path):
    sibling = tmp_path / "Beta" / "sources"
    sibling.mkdir(parents=True)
    foreign = sibling / "secret.txt"
    foreign.write_text("Do not show this in Alpha", encoding="utf-8")
    chat = library / "chats" / "private.txt"
    chat.parent.mkdir()
    chat.write_text("Chat state is not a source", encoding="utf-8")
    for path in (foreign, chat, "../Beta/sources/secret.txt", "sources/../../Beta/sources/secret.txt"):
        result = viewer.preview_document(path, data_dir=library)
        assert not result.text and "workspace's sources" in result.status
    local = source(library, "local.txt")
    local.write_text("Allowed source", encoding="utf-8")
    assert viewer.preview_document(local.relative_to(library), data_dir=library).text == "Allowed source"


def test_source_symlink_escape_is_rejected(library, tmp_path):
    foreign = tmp_path / "outside.txt"
    foreign.write_text("Private", encoding="utf-8")
    linked = source(library, "linked.txt")
    try:
        linked.symlink_to(foreign)
    except OSError:
        pytest.skip("Creating symlinks requires privileges on this Windows host.")
    result = viewer.preview_document(linked, data_dir=library)
    assert not result.text and "workspace's sources" in result.status


def test_missing_unsupported_oversized_and_bad_unit_fail_with_cleared_preview(library, monkeypatch):
    assert "unavailable" in viewer.preview_document(source(library, "missing.pdf"), data_dir=library).status
    path = source(library, "archive.bin")
    path.write_bytes(b"12345")
    assert "no built-in preview" in viewer.preview_document(path, data_dir=library).status
    monkeypatch.setattr(viewer, "MAX_FILE_BYTES", 4)
    assert "100 MiB" in viewer.preview_document(path, data_dir=library).status
    monkeypatch.setattr(viewer, "MAX_FILE_BYTES", 100)
    text = source(library, "text.txt")
    text.write_text("Content", encoding="utf-8")
    assert "valid page number" in viewer.preview_document(text, data_dir=library, unit="bad").status


def test_archive_expansion_limit_applies_before_office_or_workbook_parsing(library, monkeypatch):
    from rag_chatbot import office_preview

    monkeypatch.setattr(viewer, "MAX_EXPANDED_OFFICE_BYTES", 20)
    monkeypatch.setattr(office_preview, "get_pdf_preview", lambda *args, **kwargs: pytest.fail("Converter must not run"))
    for suffix in (".docx", ".pptx", ".xlsx"):
        path = source(library, "bomb" + suffix)
        with zipfile.ZipFile(path, "w", compression=zipfile.ZIP_DEFLATED) as archive:
            archive.writestr("large.xml", "x" * 100)
        assert "expanded Office file" in viewer.preview_document(path, data_dir=library).status


@pytest.mark.parametrize("suffix", [".png", ".xlsx", ".xls", ".docx", ".pptx"])
def test_damaged_formats_report_status_instead_of_raising(library, suffix):
    path = source(library, "damaged" + suffix)
    path.write_bytes(b"corrupt format")
    result = viewer.preview_document(path, data_dir=library)
    assert result.image is None and not result.html and not result.table
    assert "Preview unavailable" in result.status
