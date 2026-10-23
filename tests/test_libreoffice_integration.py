"""Opt-in Linux/macOS LibreOffice conversions; no network or OpenAI calls.

Install LibreOffice Writer and Impress, then run with RUN_LIBREOFFICE_SMOKE=1.
An explicitly requested run fails when the converter is missing or unusable.
"""

import os
from pathlib import Path
import sys

from docx import Document
from pptx import Presentation
from pypdf import PdfReader
import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from rag_chatbot import office_preview


pytestmark = [
    pytest.mark.skipif(os.name == "nt", reason="Exercises the Linux/macOS LibreOffice converter"),
    pytest.mark.skipif(os.environ.get("RUN_LIBREOFFICE_SMOKE") != "1",
                       reason="Set RUN_LIBREOFFICE_SMOKE=1 to run real LibreOffice conversion"),
]


@pytest.mark.parametrize("extension", [".docx", ".pptx"])
def test_real_libreoffice_preview_preserves_source_and_reuses_cache(tmp_path, monkeypatch, extension):
    workspace = tmp_path / "Managed workspace ü"
    directory = workspace / "sources" / "smoke-document"
    directory.mkdir(parents=True)
    original = directory / f"original{extension}"
    markers = ["Preview smoke page one", "Preview smoke page two"]
    if extension == ".docx":
        document = Document()
        document.add_paragraph(markers[0])
        document.add_page_break()
        document.add_paragraph(markers[1])
        document.save(original)
    else:
        presentation = Presentation()
        for marker in markers:
            slide = presentation.slides.add_slide(presentation.slide_layouts[0])
            slide.shapes.title.text = marker
        presentation.save(original)

    original_bytes = original.read_bytes()
    original_mtime = original.stat().st_mtime_ns
    preview = office_preview.get_pdf_preview(original, data_dir=workspace)
    assert preview.parent == directory / ".preview"
    assert preview.is_file() and preview.stat().st_size > 0
    with preview.open("rb") as handle:
        pages = PdfReader(handle).pages
        assert len(pages) == 2
        for page, marker in zip(pages, markers):
            assert float(page.mediabox.width) > 0 and float(page.mediabox.height) > 0
            assert marker in " ".join(page.extract_text().split())
    preview_bytes = preview.read_bytes()
    preview_mtime = preview.stat().st_mtime_ns

    def unexpected_conversion(*args, **kwargs):
        pytest.fail("The second preview must reuse the cached PDF without launching LibreOffice")

    monkeypatch.setattr(office_preview, "_run_worker", unexpected_conversion)
    assert office_preview.get_pdf_preview(original, data_dir=workspace) == preview
    assert preview.read_bytes() == preview_bytes
    assert preview.stat().st_mtime_ns == preview_mtime
    assert original.read_bytes() == original_bytes
    assert original.stat().st_mtime_ns == original_mtime
    assert not list(preview.parent.glob(".converting-*"))


def test_generated_studio_presentation_renders_all_slides_and_text(tmp_path):
    from rag_chatbot.tools.exports import export_bytes

    workspace = tmp_path / "Studio export"
    directory = workspace / "sources" / "generated-deck"
    directory.mkdir(parents=True)
    original = directory / "Presentation.pptx"
    titles = ["XRD measurement from the selected source", "Scope and limitations of the reported result"]
    bullets = [
        ["The XRD peak was measured at 26.5 degrees.", "The source names the measured sample Alpha."],
        [f"Finding {index}: The reported result applies to the measured sample; additional experiments are needed to compare other samples."
         for index in range(1, 5)],
    ]
    payload = {
        "tool_id": "presentation", "title": "Studio export smoke", "model": "offline-model",
        "content": {"slides": [{"title": title, "bullets": points,
                                "notes": "Speaker-only explanation.", "source_ids": ["source-1"]}
                               for title, points in zip(titles, bullets)]},
        "sources": [{"id": "source-1", "document_id": "document-1", "filename": "measurements.txt",
                     "locator": "Section 1", "text": "Sample Alpha: 26.5 degrees.",
                     "source_path": "sources/measurements.txt"}],
        "documents": [{"id": "document-1"}], "coverage": {"documents": 1, "chunks": 1, "processed_chunks": 1},
        "settings": {"slide_count": 2},
    }
    extension, data = export_bytes(payload, "pptx")
    assert extension == "pptx"
    original.write_bytes(data)
    preview = office_preview.get_pdf_preview(original, data_dir=workspace)
    pages = PdfReader(preview).pages
    assert len(pages) == len(titles)
    for page, title, points in zip(pages, titles, bullets):
        text = " ".join(page.extract_text().split())
        assert title in text
        assert all(point in text for point in points)
        assert "Speaker-only explanation" not in text
        assert float(page.mediabox.width) / float(page.mediabox.height) == pytest.approx(16 / 9, rel=0.01)
    assert original.read_bytes() == data
