"""Studio persistence, safe paths, and offline export regressions."""

import builtins
import copy
import csv
from io import BytesIO, StringIO
import json
from pathlib import Path

import pytest

from rag_chatbot.tools import artifacts
from rag_chatbot.tools.exports import export_bytes
from rag_chatbot.workspaces import create_workspace, workspace_path


@pytest.fixture
def workspace(tmp_path, monkeypatch):
    monkeypatch.setenv("RAG_DATA_DIR", str(tmp_path / "data"))
    create_workspace("Research")
    return "Research"


@pytest.fixture
def summary(workspace):
    return {
        "tool_id": "summary", "title": "Selected research", "model": "gpt-6-luna",
        "content": {"sections": [{"heading": "Findings", "paragraphs": [
            {"text": "The sample has a peak.", "source_ids": ["chunk-one"]}]}]},
        "sources": [{"id": "chunk-one", "document_id": "doc-one", "filename": "paper.pdf",
                     "locator": "Page 2", "text": "A peak was measured.",
                     "source_path": str(workspace_path(workspace) / "sources" / "paper.pdf")}],
        "documents": [{"id": "doc-one", "name": "paper.pdf", "content_hash": "abc", "fingerprint": "v1", "generation": "g1"}],
        "coverage": {"documents": 1, "chunks": 1, "processed_chunks": 1}, "settings": {"instructions": ""},
    }


@pytest.fixture
def spreadsheet(summary):
    value = copy.deepcopy(summary)
    value.update(tool_id="spreadsheet", title="Results", content={"columns": ["Observation"], "rows": [
        {"document_id": "doc-one", "filename": "paper.pdf", "cells": [
            {"column": "Observation", "value": "Peak measured", "source_ids": ["chunk-one"]}]}]})
    return value


def test_save_load_new_versions_and_scoped_listing(workspace, summary):
    first = artifacts.save_artifact(workspace, summary)
    edited = artifacts.save_artifact(workspace, {**first, "edited_text": "My edited summary [1]."})
    assert first["id"] != edited["id"]
    assert first["version"] == 1
    assert first["sources"][0]["source_path"] == "sources/paper.pdf"
    assert Path(summary["sources"][0]["source_path"]).is_absolute()
    assert artifacts.load_artifact(workspace, first["id"]) == first
    assert [item["id"] for item in artifacts.list_artifacts(workspace)] == [edited["id"], first["id"]]
    assert len(artifacts.list_artifacts(workspace, "summary")) == 2
    assert artifacts.list_artifacts(workspace, "spreadsheet") == []
    create_workspace("Other")
    assert artifacts.list_artifacts("Other") == []
    with pytest.raises(FileNotFoundError):
        artifacts.load_artifact("Other", first["id"])


@pytest.mark.parametrize("identifier", ["../chats", "/etc/passwd", "C:\\temp\\test", "not-a-uuid", ""])
def test_rejects_artifact_path_traversal(workspace, identifier):
    with pytest.raises(ValueError, match="identifier"):
        artifacts.load_artifact(workspace, identifier)


def test_rejects_outside_source_and_invalid_citations(workspace, summary, tmp_path):
    summary["sources"][0]["source_path"] = str(tmp_path / "outside.pdf")
    with pytest.raises(ValueError, match="workspace"):
        artifacts.save_artifact(workspace, summary)
    summary["sources"][0]["source_path"] = "sources/paper.pdf"
    summary["content"]["sections"][0]["paragraphs"][0]["source_ids"] = ["forged-source"]
    with pytest.raises(ValueError, match="unknown source"):
        artifacts.save_artifact(workspace, summary)


def test_load_validates_identity_format_size_and_skips_broken_entry(workspace, summary, monkeypatch):
    saved = artifacts.save_artifact(workspace, summary)
    path = workspace_path(workspace) / "artifacts" / saved["id"] / "manifest.json"
    changed = {**saved, "id": "00000000-0000-0000-0000-000000000000"}
    path.write_text(json.dumps(changed), encoding="utf-8")
    with pytest.raises(ValueError, match="does not match"):
        artifacts.load_artifact(workspace, saved["id"])
    assert artifacts.list_artifacts(workspace) == []
    path.write_text(json.dumps(saved), encoding="utf-8")
    monkeypatch.setattr(artifacts, "MAX_MANIFEST_BYTES", 10)
    with pytest.raises(ValueError, match="size"):
        artifacts.load_artifact(workspace, saved["id"])


@pytest.mark.parametrize("constant", ["NaN", "Infinity", "-Infinity"])
def test_loaded_manifest_rejects_nonfinite_json_metadata(workspace, summary, constant):
    saved = artifacts.save_artifact(workspace, summary)
    path = workspace_path(workspace) / "artifacts" / saved["id"] / "manifest.json"
    text = json.dumps(saved).replace('"settings": {"instructions": ""}', f'"settings": {{"bad": {constant}}}')
    path.write_text(text, encoding="utf-8")
    with pytest.raises(ValueError, match="JSON"):
        artifacts.load_artifact(workspace, saved["id"])
    assert artifacts.list_artifacts(workspace) == []


def test_loaded_manifest_handles_excessive_nesting(workspace, summary):
    saved = artifacts.save_artifact(workspace, summary)
    path = workspace_path(workspace) / "artifacts" / saved["id"] / "manifest.json"
    path.write_text("[" * 2000 + "0" + "]" * 2000, encoding="utf-8")
    with pytest.raises(ValueError, match="JSON"):
        artifacts.load_artifact(workspace, saved["id"])
    assert artifacts.list_artifacts(workspace) == []


@pytest.mark.parametrize("corruption", [
    {"tool_id": []}, {"tool_id": {}}, {"tool_id": None},
    {"created_at": []}, {"created_at": {}}, {"created_at": 123}, {"created_at": "not a timestamp"},
    {"sources": {}}, {"sources": [{"id": []}]}, {"documents": [{"id": {}}]},
    {"content": {"sections": [{"heading": "Section", "paragraphs": [None]}]}},
    {"content": {"sections": [{"heading": "Section", "paragraphs": [
        {"text": "Claim", "source_ids": [["chunk-one"]]}]}]}},
])
def test_corrupt_entry_does_not_hide_valid_saved_output(workspace, summary, corruption):
    good = artifacts.save_artifact(workspace, summary)
    bad = artifacts.save_artifact(workspace, summary)
    path = workspace_path(workspace) / "artifacts" / bad["id"] / "manifest.json"
    path.write_text(json.dumps({**bad, **corruption}), encoding="utf-8")
    with pytest.raises(ValueError):
        artifacts.load_artifact(workspace, bad["id"])
    assert [entry["id"] for entry in artifacts.list_artifacts(workspace)] == [good["id"]]
    assert artifacts.load_artifact(workspace, good["id"]) == good


@pytest.mark.parametrize("field", ["sources", "documents"])
def test_invalid_snapshot_shape_is_rejected_when_saving(workspace, summary, field):
    summary[field] = None
    with pytest.raises(ValueError, match="lists"):
        artifacts.save_artifact(workspace, summary)


def test_cross_document_spreadsheet_citations_rejected(workspace, spreadsheet):
    spreadsheet["documents"].append({"id": "doc-other"})
    spreadsheet["content"]["rows"][0]["document_id"] = "doc-other"
    with pytest.raises(ValueError, match="own source"):
        artifacts.save_artifact(workspace, spreadsheet)


def test_invalid_artifact_payload_raises_clear_error(workspace):
    with pytest.raises(ValueError, match="object"):
        artifacts.save_artifact(workspace, [])
    with pytest.raises(ValueError, match="JSON"):
        artifacts.save_artifact(workspace, {"bad": float("nan")})
    assert artifacts.list_artifacts(workspace) == []


def test_artifact_directory_link_rejected_before_write(workspace, summary, tmp_path):
    target = workspace_path(workspace) / "artifacts"
    outside = tmp_path / "outside"
    outside.mkdir()
    try:
        target.symlink_to(outside, target_is_directory=True)
    except OSError:
        pytest.skip("Creating symlinks requires privileges on this Windows installation.")
    with pytest.raises(ValueError, match="linked"):
        artifacts.save_artifact(workspace, summary)
    assert list(outside.iterdir()) == []


def test_linked_manifest_and_export_rejected(workspace, summary, tmp_path):
    saved = artifacts.save_artifact(workspace, summary)
    directory = workspace_path(workspace) / "artifacts" / saved["id"]
    external = tmp_path / "external.txt"
    external.write_text("untouched", encoding="utf-8")
    output = directory / "output.md"
    try:
        output.symlink_to(external)
    except OSError:
        pytest.skip("Creating symlinks requires privileges on this Windows installation.")
    with pytest.raises(ValueError, match="linked"):
        artifacts.export_artifact(workspace, saved["id"], "md")
    assert external.read_text(encoding="utf-8") == "untouched"
    output.unlink()
    manifest = directory / "manifest.json"
    manifest.unlink()
    manifest.symlink_to(external)
    with pytest.raises(ValueError, match="linked"):
        artifacts.load_artifact(workspace, saved["id"])


def test_summary_export_has_citations_provenance_and_edits(workspace, summary):
    saved = artifacts.save_artifact(workspace, summary)
    text = artifacts.export_artifact(workspace, saved["id"], "markdown").read_text(encoding="utf-8")
    assert "The sample has a peak. [1]" in text
    assert "paper\\.pdf" in text
    assert "Page 2" in text and "gpt-6-luna" in text and "processed_chunks" in text
    edited = artifacts.save_artifact(workspace, {**saved, "edited_text": "Edited observation [1]."})
    assert artifacts.render_summary(edited, include_sources=False) == "Edited observation [1]."
    content = artifacts.export_artifact(workspace, edited["id"], "md").read_text(encoding="utf-8")
    assert "Edited by the user" in content and "## Sources" in content
    exported = artifacts.export_artifact(workspace, edited["id"], "json")
    assert json.loads(exported.read_text(encoding="utf-8"))["documents"] == edited["documents"]


def test_spreadsheet_edits_shape_and_csv_formula_injection(workspace, spreadsheet):
    for value in ("=1+2", "  +SUM(1,2)", "@SUM(A1)", "-5", "\tplain"):
        saved = artifacts.save_artifact(workspace, {**spreadsheet, "edited_rows": [[value]]})
        exported = artifacts.export_artifact(workspace, saved["id"], "csv")
        rows = list(csv.reader(StringIO(exported.read_text(encoding="utf-8-sig"))))
        assert rows[1][1] == "'" + value
        assert rows[1][2] == "[1]"
        assert artifacts.spreadsheet_table(saved) == (["Observation"], [[value]])
        assert any("Edited by the user" in row[0] for row in rows if row)
    with pytest.raises(ValueError, match="row and column"):
        artifacts.save_artifact(workspace, {**spreadsheet, "edited_rows": [["a", "b"]]})


def test_xlsx_has_literal_strings_and_source_sheet(workspace, spreadsheet):
    from openpyxl import load_workbook
    spreadsheet["content"]["rows"][0]["cells"][0]["value"] = '=HYPERLINK("https://example.com","click")'
    saved = artifacts.save_artifact(workspace, spreadsheet)
    output = artifacts.export_artifact(workspace, saved["id"], "xlsx")
    workbook = load_workbook(output, data_only=False)
    assert workbook["Spreadsheet"]["B2"].data_type == "s"
    assert workbook["Spreadsheet"]["B2"].value.startswith("=HYPERLINK")
    assert workbook["Spreadsheet"]["C2"].value == "[1]"
    assert workbook["Sources"]["D2"].value == "A peak was measured."
    assert "processed_chunks" in workbook["Provenance"]["A4"].value


def test_docx_and_missing_optional_exporters(workspace, summary, spreadsheet, monkeypatch):
    from docx import Document
    saved = artifacts.save_artifact(workspace, summary)
    output = artifacts.export_artifact(workspace, saved["id"], "docx")
    paragraphs = [paragraph.text for paragraph in Document(output).paragraphs]
    assert "The sample has a peak. [1]" in paragraphs
    original_import = builtins.__import__

    def missing(name, *args, **kwargs):
        if name in ("docx", "openpyxl"):
            raise ImportError("simulated missing optional dependency")
        return original_import(name, *args, **kwargs)

    monkeypatch.setattr(builtins, "__import__", missing)
    with pytest.raises(RuntimeError, match="python-docx"):
        export_bytes(saved, "docx")
    with pytest.raises(RuntimeError, match="openpyxl"):
        export_bytes(spreadsheet, "xlsx")
    assert export_bytes(saved, "md")[0] == "md"
    assert export_bytes(spreadsheet, "csv")[0] == "csv"


def test_office_export_rejects_text_it_cannot_preserve(summary, spreadsheet):
    spreadsheet["edited_rows"] = [["x" * 32768]]
    with pytest.raises(ValueError, match="32,767"):
        export_bytes(spreadsheet, "xlsx")
    assert "x" * 32768 in export_bytes(spreadsheet, "csv")[1].decode("utf-8-sig")
    spreadsheet["edited_rows"] = [["illegal\x00character"]]
    with pytest.raises(ValueError, match="control"):
        export_bytes(spreadsheet, "xlsx")
    summary["edited_text"] = "illegal\x00character"
    with pytest.raises(ValueError, match="control"):
        export_bytes(summary, "docx")


def test_current_summary_repeated_save_keeps_one_snapshot(workspace, summary):
    assert artifacts.load_current_summary(workspace) is None
    assert not (workspace_path(workspace) / "artifacts").exists()
    first = artifacts.save_current_summary(workspace, summary)
    second = artifacts.save_current_summary(workspace, {**first, "edited_text": "Edited finding [1]."})
    third = artifacts.save_current_summary(workspace, {**summary, "title": "Regenerated finding"})
    assert first["id"] == second["id"] == third["id"]
    assert first["created_at"] == second["created_at"] == third["created_at"]
    assert first["updated_at"] <= second["updated_at"] <= third["updated_at"]
    assert "edited_text" not in third
    assert third["sources"][0]["source_path"] == "sources/paper.pdf"
    assert artifacts.load_current_summary(workspace) == third
    directory = workspace_path(workspace) / "artifacts"
    assert [path.name for path in directory.iterdir()] == ["summary"]
    assert [path.name for path in (directory / "summary").iterdir()] == ["current.json"]
    assert artifacts.list_artifacts(workspace) == []
    create_workspace("Other")
    assert artifacts.load_current_summary("Other") is None


def test_current_summary_migrates_latest_legacy_once_and_preserves_files(workspace, summary, spreadsheet):
    old = artifacts.save_artifact(workspace, summary)
    latest = artifacts.save_artifact(workspace, {**summary, "title": "Most recent Summary"})
    table = artifacts.save_artifact(workspace, spreadsheet)
    root = workspace_path(workspace) / "artifacts"
    legacy_bytes = {value["id"]: (root / value["id"] / "manifest.json").read_bytes()
                    for value in (old, latest, table)}
    current = artifacts.load_current_summary(workspace)
    assert current["id"] == latest["id"]
    assert current["title"] == latest["title"]
    assert current["updated_at"] == latest["created_at"]
    updated = artifacts.save_current_summary(workspace, {**current, "edited_text": "My current Summary."})
    assert artifacts.load_current_summary(workspace) == updated
    for identifier, original in legacy_bytes.items():
        assert (root / identifier / "manifest.json").read_bytes() == original
    assert len(artifacts.list_artifacts(workspace, "summary")) == 2
    assert artifacts.load_artifact(workspace, table["id"]) == table


def test_export_summary_custom_filename_changes_no_saved_content(workspace, summary):
    current = artifacts.save_current_summary(workspace, summary)
    current_path = workspace_path(workspace) / "artifacts" / "summary" / "current.json"
    before = current_path.read_bytes()
    edited = {**current, "edited_text": "Unsaved export text [1]."}
    exported = artifacts.export_summary(workspace, edited, "json", "My findings.JSON")
    assert exported.name == "My findings.json"
    content = json.loads(exported.read_text(encoding="utf-8"))
    assert content["edited_text"] == edited["edited_text"]
    assert content["sources"][0]["source_path"] == "sources/paper.pdf"
    assert current_path.read_bytes() == before
    default = artifacts.export_summary(workspace, current, "markdown")
    assert default.name == "Summary.md"
    renamed = artifacts.export_summary(workspace, current, "docx", "Results.md")
    assert renamed.name == "Results.docx"
    assert artifacts.list_artifacts(workspace) == []
    assert "edited_text" not in artifacts.load_current_summary(workspace)


def test_export_without_save_never_migrates_or_creates_current_summary(workspace, summary):
    legacy = artifacts.save_artifact(workspace, summary)
    output = artifacts.export_summary(workspace, summary, "md", "Temporary")
    directory = workspace_path(workspace) / "artifacts"
    assert output.read_text(encoding="utf-8").startswith("# Selected research")
    assert not (directory / "summary" / "current.json").exists()
    assert [item["id"] for item in artifacts.list_artifacts(workspace)] == [legacy["id"]]


@pytest.mark.parametrize("filename", ["../escape", "sub/name", "sub\\name", "C:escape", ".hidden",
                                     "CON", "con.md", "NUL.txt", "LPT¹", "COM3.docx", "a?b", "bad\x00name",
                                     "bad\x7fname", "trail.", " name", "name ", "", "a" * 101,
                                     "界" * 84, "😀" * 63])
def test_summary_export_rejects_unsafe_or_overlong_filenames(workspace, summary, filename):
    with pytest.raises(ValueError, match="filename"):
        artifacts.export_summary(workspace, summary, "docx", filename)
    assert not (workspace_path(workspace) / "artifacts").exists()


def test_summary_export_accepts_portable_unicode_and_replaces_known_suffix(workspace, summary):
    for filename, expected in [("Kết quả", "Kết quả.md"), ("Cafe\u0301.JSON", "Café.md"),
                               ("Research.v1", "Research.v1.md"), ("😀" * 62, "😀" * 62 + ".md")]:
        result = artifacts.export_summary(workspace, summary, "md", filename)
        assert result.name == expected
        assert result.is_file()


def test_failed_current_summary_save_keeps_previous_contents(workspace, summary, monkeypatch):
    first = artifacts.save_current_summary(workspace, summary)
    manifest = workspace_path(workspace) / "artifacts" / "summary" / "current.json"
    before = manifest.read_bytes()
    original_replace = artifacts.os.replace

    def failed_replace(source, target):
        if Path(target) == manifest:
            raise OSError("Simulated disk failure")
        return original_replace(source, target)

    monkeypatch.setattr(artifacts.os, "replace", failed_replace)
    with pytest.raises(OSError, match="disk failure"):
        artifacts.save_current_summary(workspace, {**first, "edited_text": "Unsaved edit"})
    assert manifest.read_bytes() == before
    assert [path.name for path in manifest.parent.iterdir()] == ["current.json"]


def test_summary_cancellation_before_publish_preserves_previous_summary(workspace, summary):
    first = artifacts.save_current_summary(workspace, summary)
    manifest = workspace_path(workspace) / "artifacts" / "summary" / "current.json"
    before = manifest.read_bytes()
    checks = []

    def cancelled_before_publish():
        checks.append(True)
        return len(checks) > 1

    with pytest.raises(RuntimeError, match="cancelled"):
        artifacts.save_current_summary(workspace, {**first, "edited_text": "Late result"},
                                       is_cancelled=cancelled_before_publish)
    assert len(checks) == 2
    assert manifest.read_bytes() == before
    assert [path.name for path in manifest.parent.iterdir()] == ["current.json"]


def test_current_summary_rejects_wrong_tool_and_corrupt_manifest(workspace, summary, spreadsheet):
    with pytest.raises(ValueError, match="Only a Summary"):
        artifacts.save_current_summary(workspace, spreadsheet)
    with pytest.raises(ValueError, match="Only a Summary"):
        artifacts.export_summary(workspace, spreadsheet, "json")
    artifacts.save_current_summary(workspace, summary)
    current = workspace_path(workspace) / "artifacts" / "summary" / "current.json"
    current.write_text('{"broken":', encoding="utf-8")
    with pytest.raises(ValueError, match="JSON"):
        artifacts.load_current_summary(workspace)
    with pytest.raises(ValueError, match="JSON"):
        artifacts.save_current_summary(workspace, summary)
    assert current.read_text(encoding="utf-8") == '{"broken":'


def test_current_summary_and_named_export_reject_links(workspace, summary, tmp_path):
    current = artifacts.save_current_summary(workspace, summary)
    directory = workspace_path(workspace) / "artifacts" / "summary"
    outside = tmp_path / "outside.md"
    outside.write_text("unchanged", encoding="utf-8")
    exports = directory / "exports"
    exports.mkdir()
    linked = exports / "Summary.md"
    try:
        linked.symlink_to(outside)
    except OSError:
        pytest.skip("Creating symlinks requires privileges on this Windows installation.")
    with pytest.raises(ValueError, match="linked"):
        artifacts.export_summary(workspace, current, "md")
    assert outside.read_text(encoding="utf-8") == "unchanged"
    manifest = directory / "current.json"
    manifest.unlink()
    manifest.symlink_to(outside)
    with pytest.raises(ValueError, match="linked"):
        artifacts.load_current_summary(workspace)
    with pytest.raises(ValueError, match="linked"):
        artifacts.save_current_summary(workspace, summary)


def test_current_spreadsheet_roundtrip_edits_and_workspace_isolation(workspace, spreadsheet):
    assert artifacts.load_current_spreadsheet(workspace) is None
    assert not (workspace_path(workspace) / "artifacts").exists()
    original = copy.deepcopy(spreadsheet)
    generated = artifacts.save_current_spreadsheet(workspace, spreadsheet)
    edited = artifacts.save_current_spreadsheet(workspace, {**generated, "edited_rows": [["Edited finding"]]})
    assert artifacts.load_current_spreadsheet(workspace) == edited
    assert artifacts.spreadsheet_table(edited) == (["Observation"], [["Edited finding"]])
    assert edited["content"] == generated["content"]
    assert edited["sources"] == generated["sources"]
    assert edited["documents"] == generated["documents"]
    assert edited["sources"][0]["source_path"] == "sources/paper.pdf"
    assert spreadsheet == original
    create_workspace("Other")
    assert artifacts.load_current_spreadsheet("Other") is None
    assert not (workspace_path("Other") / "artifacts").exists()
    other = artifacts.save_current_spreadsheet("Other", {**edited, "title": "Other result"})
    assert other["id"] != edited["id"]
    assert artifacts.load_current_spreadsheet(workspace) == edited
    assert artifacts.load_current_spreadsheet("Other") == other


def test_current_spreadsheet_replacement_keeps_identity_without_versions(workspace, spreadsheet):
    first = artifacts.save_current_spreadsheet(workspace, spreadsheet)
    edited = artifacts.save_current_spreadsheet(workspace, {**first, "edited_rows": [["Edited finding"]]})
    regenerated = artifacts.save_current_spreadsheet(workspace, {**spreadsheet, "title": "New result"})
    assert first["id"] == edited["id"] == regenerated["id"]
    assert first["created_at"] == edited["created_at"] == regenerated["created_at"]
    assert first["updated_at"] <= edited["updated_at"] <= regenerated["updated_at"]
    assert "edited_rows" not in regenerated
    assert artifacts.load_current_spreadsheet(workspace) == regenerated
    assert artifacts.list_artifacts(workspace) == []
    directory = workspace_path(workspace) / "artifacts"
    assert [path.name for path in directory.iterdir()] == ["spreadsheet"]
    assert [path.name for path in (directory / "spreadsheet").iterdir()] == ["current.json"]


def test_current_spreadsheet_adopts_latest_saved_table_once_preserving_versions(workspace, spreadsheet, summary):
    first = artifacts.save_artifact(workspace, spreadsheet)
    latest = artifacts.save_artifact(workspace, {**spreadsheet, "title": "Latest table", "edited_rows": [["Saved edit"]]})
    other_tool = artifacts.save_artifact(workspace, summary)
    root = workspace_path(workspace) / "artifacts"
    legacy_bytes = {value["id"]: (root / value["id"] / "manifest.json").read_bytes()
                    for value in (first, latest, other_tool)}
    current = artifacts.load_current_spreadsheet(workspace)
    assert current == {**latest, "updated_at": latest["created_at"]}
    assert artifacts.spreadsheet_table(current) == (["Observation"], [["Saved edit"]])
    edited = artifacts.save_current_spreadsheet(workspace, {**current, "edited_rows": [["Current edit"]]})
    artifacts.save_artifact(workspace, {**spreadsheet, "title": "Later archived version"})
    assert artifacts.load_current_spreadsheet(workspace) == edited
    assert edited["id"] == latest["id"]
    for identifier, original in legacy_bytes.items():
        assert (root / identifier / "manifest.json").read_bytes() == original
    assert len(artifacts.list_artifacts(workspace, "spreadsheet")) == 3


def test_current_spreadsheet_save_adopts_identity_without_copying_legacy_first(workspace, spreadsheet):
    saved = artifacts.save_artifact(workspace, spreadsheet)
    current = artifacts.save_current_spreadsheet(workspace, {**spreadsheet, "title": "New generation"})
    assert current["id"] == saved["id"]
    assert current["created_at"] == saved["created_at"]
    assert current["title"] == "New generation"
    assert artifacts.load_artifact(workspace, saved["id"]) == saved
    assert artifacts.load_current_spreadsheet(workspace) == current


@pytest.mark.parametrize("change", [
    {"tool_id": "summary"}, {"edited_rows": [["extra", "cell"]]},
    {"content": {"columns": ["Observation"], "rows": [
        {"document_id": "doc-one", "filename": "paper.pdf", "cells": [
            {"column": "Observation", "value": "Peak", "source_ids": ["unknown"]}]}]}},
])
def test_current_spreadsheet_invalid_save_does_not_create_storage(workspace, spreadsheet, change):
    with pytest.raises(ValueError):
        artifacts.save_current_spreadsheet(workspace, {**spreadsheet, **change})
    assert not (workspace_path(workspace) / "artifacts").exists()


def test_current_spreadsheet_cancellation_before_start_does_not_adopt_legacy(workspace, spreadsheet):
    saved = artifacts.save_artifact(workspace, spreadsheet)
    with pytest.raises(ValueError, match="callable"):
        artifacts.save_current_spreadsheet(workspace, spreadsheet, is_cancelled=True)
    with pytest.raises(RuntimeError, match="cancelled"):
        artifacts.save_current_spreadsheet(workspace, spreadsheet, is_cancelled=lambda: True)
    assert not (workspace_path(workspace) / "artifacts" / "spreadsheet").exists()
    assert artifacts.load_artifact(workspace, saved["id"]) == saved


def test_current_spreadsheet_cancelled_publish_preserves_previous_contents(workspace, spreadsheet):
    first = artifacts.save_current_spreadsheet(workspace, spreadsheet)
    manifest = workspace_path(workspace) / "artifacts" / "spreadsheet" / "current.json"
    before = manifest.read_bytes()
    checks = []

    def cancelled_before_publish():
        checks.append(True)
        return len(checks) > 1

    with pytest.raises(RuntimeError, match="cancelled"):
        artifacts.save_current_spreadsheet(workspace, {**first, "edited_rows": [["Late edit"]]},
                                           is_cancelled=cancelled_before_publish)
    assert len(checks) == 2
    assert manifest.read_bytes() == before
    assert [path.name for path in manifest.parent.iterdir()] == ["current.json"]


def test_current_spreadsheet_failed_replace_preserves_previous_contents(workspace, spreadsheet, monkeypatch):
    first = artifacts.save_current_spreadsheet(workspace, spreadsheet)
    manifest = workspace_path(workspace) / "artifacts" / "spreadsheet" / "current.json"
    before = manifest.read_bytes()
    original_replace = artifacts.os.replace

    def failed_replace(source, target):
        if Path(target) == manifest:
            raise OSError("Simulated disk failure")
        return original_replace(source, target)

    monkeypatch.setattr(artifacts.os, "replace", failed_replace)
    with pytest.raises(OSError, match="disk failure"):
        artifacts.save_current_spreadsheet(workspace, {**first, "edited_rows": [["Unsaved edit"]]})
    assert manifest.read_bytes() == before
    assert [path.name for path in manifest.parent.iterdir()] == ["current.json"]


@pytest.mark.parametrize("corruption", ["json", "version", "wrong_tool"])
def test_current_spreadsheet_corrupt_manifest_is_not_overwritten_or_replaced_by_legacy(
        workspace, spreadsheet, summary, corruption):
    artifacts.save_artifact(workspace, spreadsheet)
    current = artifacts.save_current_spreadsheet(workspace, spreadsheet)
    manifest = workspace_path(workspace) / "artifacts" / "spreadsheet" / "current.json"
    if corruption == "json":
        text = '{"broken":'
    elif corruption == "version":
        text = json.dumps({**current, "version": 9})
    else:
        text = json.dumps(artifacts.save_artifact(workspace, summary))
    manifest.write_text(text, encoding="utf-8")
    for operation in (lambda: artifacts.load_current_spreadsheet(workspace),
                      lambda: artifacts.save_current_spreadsheet(workspace, spreadsheet)):
        with pytest.raises(ValueError):
            operation()
        assert manifest.read_text(encoding="utf-8") == text


def test_current_spreadsheet_rejects_linked_manifest(workspace, spreadsheet, tmp_path):
    artifacts.save_current_spreadsheet(workspace, spreadsheet)
    manifest = workspace_path(workspace) / "artifacts" / "spreadsheet" / "current.json"
    outside = tmp_path / "outside.json"
    outside.write_text("unchanged", encoding="utf-8")
    manifest.unlink()
    try:
        manifest.symlink_to(outside)
    except OSError:
        pytest.skip("Creating symlinks requires privileges on this Windows installation.")
    with pytest.raises(ValueError, match="linked"):
        artifacts.load_current_spreadsheet(workspace)
    with pytest.raises(ValueError, match="linked"):
        artifacts.save_current_spreadsheet(workspace, spreadsheet)
    assert outside.read_text(encoding="utf-8") == "unchanged"


def test_direct_spreadsheet_export_does_not_replace_current_table(workspace, spreadsheet):
    current = artifacts.save_current_spreadsheet(workspace, spreadsheet)
    manifest = workspace_path(workspace) / "artifacts" / "spreadsheet" / "current.json"
    before = manifest.read_bytes()
    exported = artifacts.export_spreadsheet(workspace, {**current, "edited_rows": [["Export-only edit"]]}, "json")
    assert json.loads(exported.read_text(encoding="utf-8"))["edited_rows"] == [["Export-only edit"]]
    assert manifest.read_bytes() == before
    assert artifacts.load_current_spreadsheet(workspace) == current
    assert artifacts.list_artifacts(workspace, "spreadsheet") == []


def test_direct_spreadsheet_export_uses_edits_without_creating_saved_version(workspace, spreadsheet):
    edited = {**spreadsheet, "edited_rows": [["Edited result"]]}
    output = artifacts.export_spreadsheet(workspace, edited, "json")
    assert output.name == "Spreadsheet.json"
    content = json.loads(output.read_text(encoding="utf-8"))
    assert content["edited_rows"] == [["Edited result"]]
    assert content["sources"][0]["source_path"] == "sources/paper.pdf"
    assert "id" not in content
    assert "edited_rows" not in spreadsheet
    assert Path(spreadsheet["sources"][0]["source_path"]).is_absolute()
    assert artifacts.list_artifacts(workspace) == []
    root = workspace_path(workspace) / "artifacts"
    assert [path.name for path in root.iterdir()] == ["spreadsheet"]
    assert [path.name for path in (root / "spreadsheet").iterdir()] == ["exports"]


def test_direct_spreadsheet_export_preserves_saved_history_and_current_summary(workspace, spreadsheet, summary):
    saved = artifacts.save_artifact(workspace, spreadsheet)
    current = artifacts.save_current_summary(workspace, summary)
    root = workspace_path(workspace) / "artifacts"
    original_table = (root / saved["id"] / "manifest.json").read_bytes()
    original_summary = (root / "summary" / "current.json").read_bytes()
    output = artifacts.export_spreadsheet(workspace, {**saved, "edited_rows": [["Updated value"]]},
                                        "csv", "Findings.xlsx")
    assert output.name == "Findings.csv"
    rows = list(csv.reader(StringIO(output.read_text(encoding="utf-8-sig"))))
    assert rows[1][1] == "Updated value"
    assert rows[1][2] == "[1]"
    assert (root / saved["id"] / "manifest.json").read_bytes() == original_table
    assert (root / "summary" / "current.json").read_bytes() == original_summary
    assert artifacts.load_current_summary(workspace) == current
    assert [entry["id"] for entry in artifacts.list_artifacts(workspace)] == [saved["id"]]


def test_direct_spreadsheet_exports_keep_formula_protection_and_citations(workspace, spreadsheet):
    from openpyxl import load_workbook
    edited = {**spreadsheet, "edited_rows": [["=SUM(1,2)"]]}
    csv_output = artifacts.export_spreadsheet(workspace, edited, "csv", "Research.CSV")
    rows = list(csv.reader(StringIO(csv_output.read_text(encoding="utf-8-sig"))))
    assert csv_output.name == "Research.csv"
    assert rows[1][1] == "'=SUM(1,2)"
    assert rows[1][2] == "[1]"
    assert any("Edited by the user" in row[0] for row in rows if row)
    xlsx_output = artifacts.export_spreadsheet(workspace, edited, "xlsx", "Results.json")
    workbook = load_workbook(xlsx_output, data_only=False)
    assert xlsx_output.name == "Results.xlsx"
    assert workbook["Spreadsheet"]["B2"].value == "=SUM(1,2)"
    assert workbook["Spreadsheet"]["B2"].data_type == "s"
    assert workbook["Spreadsheet"]["C2"].value == "[1]"
    assert workbook["Sources"]["D2"].value == "A peak was measured."
    assert artifacts.list_artifacts(workspace, "spreadsheet") == []


@pytest.mark.parametrize("filename", ["../escape", "sub/name", "sub\\name", "C:escape", "NUL.csv", "CON.xlsx",
                                     ".hidden", "LPT¹", "bad\x00name", "trailing.", " name", "", "a" * 101,
                                     "界" * 84, "😀" * 63])
def test_direct_spreadsheet_export_rejects_unsafe_filenames(workspace, spreadsheet, filename):
    with pytest.raises(ValueError, match="filename"):
        artifacts.export_spreadsheet(workspace, spreadsheet, "xlsx", filename)
    assert not (workspace_path(workspace) / "artifacts").exists()


def test_direct_spreadsheet_export_preserves_unicode_and_validates_shape(workspace, spreadsheet, summary):
    output = artifacts.export_spreadsheet(workspace, spreadsheet, "csv", "Kết quả.JSON")
    assert output.name == "Kết quả.csv"
    with pytest.raises(ValueError, match="Only a Spreadsheet"):
        artifacts.export_spreadsheet(workspace, summary, "json")
    with pytest.raises(ValueError, match="row and column"):
        artifacts.export_spreadsheet(workspace, {**spreadsheet, "edited_rows": [["extra", "cell"]]}, "csv")
    with pytest.raises(ValueError, match="export format"):
        artifacts.export_spreadsheet(workspace, spreadsheet, "md")


def test_direct_spreadsheet_export_rejects_linked_destination(workspace, spreadsheet, tmp_path):
    output = artifacts.export_spreadsheet(workspace, spreadsheet, "csv")
    outside = tmp_path / "outside.csv"
    outside.write_text("unchanged", encoding="utf-8")
    output.unlink()
    try:
        output.symlink_to(outside)
    except OSError:
        pytest.skip("Creating symlinks requires privileges on this Windows installation.")
    with pytest.raises(ValueError, match="linked"):
        artifacts.export_spreadsheet(workspace, spreadsheet, "csv")
    assert outside.read_text(encoding="utf-8") == "unchanged"
