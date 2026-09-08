"""Regression coverage for persisted images, document text, and image payloads."""

import base64
import io
import json
import os
import struct
import sys

import pytest
from docx import Document
from PIL import Image
from pptx import Presentation
from pptx.util import Inches

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import history_store
import messages
import readers


@pytest.fixture
def history_dir(tmp_path, monkeypatch):
    directory = tmp_path / "history"
    directory.mkdir()
    monkeypatch.setattr(history_store, "HISTORY_DIR", str(directory))
    return directory


@pytest.mark.parametrize("name", ["My generated image", "Chart (final)", "Vietnamese chart 01"])
def test_generated_images_reload_with_punctuated_history_names(history_dir, name):
    inline = "data:image/png;base64," + base64.b64encode(b"image bytes").decode()
    original = [{"role": "assistant", "content": f"Result: ![Image]({inline})"}]

    history_store.save_history(original, name)
    stored = json.loads((history_dir / f"{name}.json").read_text(encoding="utf8"))
    assert inline not in stored[0]["content"]
    loaded, _ = history_store.load_history(name)
    assert loaded == original


@pytest.mark.parametrize("reference", ["../private.txt", "assets/../../private.txt", "assets/..\\..\\private.txt"])
def test_history_image_cannot_read_files_outside_assets(history_dir, reference):
    (history_dir.parent / "private.txt").write_text("private content", encoding="utf8")
    saved = [{"role": "user", "content": "image", "image_url": reference}]
    (history_dir / "untrusted.json").write_text(json.dumps(saved), encoding="utf8")

    loaded, _ = history_store.load_history("untrusted")
    assert loaded == [{"role": "user", "content": "image"}]


def test_history_image_cannot_read_an_absolute_file_path(history_dir):
    private = history_dir.parent / "private.png"
    private.write_bytes(b"private content")
    assert history_store._read_asset_as_data_url(str(private)) is None


def test_external_urls_with_asset_paths_are_not_treated_as_local_files(history_dir):
    local_asset = history_dir / "assets" / "chat"
    local_asset.mkdir(parents=True)
    (local_asset / "0000.png").write_bytes(b"local image")
    remote = "https://example.test/assets/chat/0000.png"
    original = [{"role": "user", "content": f"![Image]({remote})", "image_url": remote}]
    history_store.save_history(original, "remote")

    loaded, _ = history_store.load_history("remote")
    assert loaded == original


@pytest.mark.parametrize("failure_stage", ["serialization", "replacement"])
def test_failed_history_overwrite_preserves_previous_json_and_images(history_dir, monkeypatch, failure_stage):
    def data_url(payload):
        return "data:image/png;base64," + base64.b64encode(payload).decode()

    original = [{"role": "user", "content": "original", "image_url": data_url(b"first image")}]
    replacement = [{"role": "user", "content": "replacement", "image_url": data_url(b"second image")}]
    history_store.save_history(original, "saved")
    previous_json = (history_dir / "saved.json").read_bytes()

    def fail_serialization(value, file, **kwargs):
        file.write("partial JSON")
        raise OSError("simulated disk error")

    def fail_replacement(source, destination):
        raise OSError("simulated replacement error")

    if failure_stage == "serialization":
        monkeypatch.setattr(history_store.json, "dump", fail_serialization)
    else:
        monkeypatch.setattr(history_store.os, "replace", fail_replacement)
    with pytest.raises(OSError, match="simulated"):
        history_store.save_history(replacement, "saved")

    assert (history_dir / "saved.json").read_bytes() == previous_json
    assert not list(history_dir.glob(".history-*.tmp"))
    loaded, _ = history_store.load_history("saved")
    assert loaded == original


def test_overwriting_images_reuses_identical_assets_and_keeps_new_images(history_dir):
    first = "data:image/png;base64," + base64.b64encode(b"first image").decode()
    second = "data:image/png;base64," + base64.b64encode(b"second image").decode()
    original = [{"role": "assistant", "content": f"![Image]({first})"}]
    updated = [{"role": "assistant", "content": f"![Image]({second})"}]
    history_store.save_history(original, "updated")
    history_store.save_history(original, "updated")
    target = history_dir / "assets" / "updated"
    assert len(list(target.iterdir())) == 1

    history_store.save_history(updated, "updated")
    history_store.save_history(updated, "updated")
    assert len(list(target.iterdir())) == 2
    assert (target / "0000_0.png").read_bytes() == b"first image"
    loaded, _ = history_store.load_history("updated")
    assert loaded == updated


@pytest.mark.parametrize("entry", [None, 42, [], {"role": "tool", "content": "x"}, {"content": []}, {"image_url": {"url": "x"}}])
def test_invalid_history_messages_fail_before_loading(history_dir, entry):
    (history_dir / "broken.json").write_text(json.dumps([entry]), encoding="utf8")
    with pytest.raises(ValueError, match="invalid message at position 1"):
        history_store.load_history("broken")


@pytest.mark.parametrize("name, expected", [("CON", "_CON"), ("nul.txt", "_nul.txt"), ("LPT1", "_LPT1"), ("name. ", "name"), (" .. ", "Chat history")])
def test_history_names_are_valid_on_windows(name, expected):
    assert history_store.sanitize_history_file_name(name) == expected


def test_cmyk_jpeg_upload_is_converted_to_rgb_png(tmp_path):
    path = tmp_path / "print-photo.jpg"
    Image.new("CMYK", (2, 2), (10, 20, 30, 40)).save(path)

    data_url = messages.image_file_to_data_url(path)
    with Image.open(io.BytesIO(base64.b64decode(data_url.split(",", 1)[1]))) as converted:
        assert converted.format == "PNG"
        assert converted.mode == "RGB"
        assert converted.size == (2, 2)


@pytest.mark.parametrize("prepare", [messages.prepare_chat_messages, messages.prepare_responses_input])
def test_generated_images_stay_in_history_but_not_text_api_payloads(prepare):
    inline = "data:image/png;base64," + base64.b64encode(b"image bytes" * 1000).decode()
    generated = f"Generated image.\n\n![Image]({inline})"
    history = [
        {"role": "assistant", "content": generated},
        {"role": "user", "content": "Describe this upload", "image_url": inline},
    ]

    payload = prepare(history)
    assert payload[0]["content"] == "Generated image.\n\n[Generated image]"
    assert history[0]["content"] == generated
    assert inline in json.dumps(payload[1])
    assert inline in messages.replace_history_content(history)[0]["content"]


def test_word_tables_are_extracted_in_document_order(tmp_path):
    document = Document()
    document.add_paragraph("Before table")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Revenue"
    table.cell(0, 1).text = "$123"
    nested = table.cell(0, 1).add_table(rows=1, cols=1)
    nested.cell(0, 0).text = "Nested detail"
    document.add_paragraph("After table")
    path = tmp_path / "report.docx"
    document.save(path)

    text = readers.read_document(path)
    assert text.index("Before table") < text.index("Revenue") < text.index("After table")
    assert "$123" in text and "Nested detail" in text


def test_powerpoint_tables_and_grouped_text_are_extracted(tmp_path):
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    table = slide.shapes.add_table(1, 2, Inches(1), Inches(1), Inches(4), Inches(1)).table
    table.cell(0, 0).text = "Revenue"
    table.cell(0, 1).text = "$456"
    group = slide.shapes.add_group_shape()
    group.shapes.add_textbox(Inches(1), Inches(3), Inches(4), Inches(1)).text = "Grouped annotation"
    path = tmp_path / "report.pptx"
    presentation.save(path)

    text = readers.read_document(path)
    assert "Revenue" in text and "$456" in text and "Grouped annotation" in text


def test_utf8_bom_is_removed_and_bom_only_file_is_empty(tmp_path):
    path = tmp_path / "bom.txt"
    path.write_text("content", encoding="utf-8-sig")
    assert readers.read_document(path) == "content"
    path.write_text("", encoding="utf-8-sig")
    with pytest.raises(readers.EmptyDocumentError):
        readers.read_document(path)


def test_legacy_xls_document_can_be_read(tmp_path):
    # A minimal BIFF2 worksheet exercises the xlrd dependency without requiring
    # a separate legacy Excel writer in the development environment.
    def record(code, payload):
        return struct.pack("<HH", code, len(payload)) + payload

    path = tmp_path / "legacy.xls"
    path.write_bytes(
        record(0x0009, struct.pack("<HH", 7, 0x0010))
        + record(0x0042, struct.pack("<H", 1252))
        + record(0x0000, struct.pack("<HHHH", 0, 2, 0, 1))
        + record(0x0004, struct.pack("<HH3sB", 0, 0, b"\x00" * 3, 7) + b"Revenue")
        + record(0x0003, struct.pack("<HH3sd", 1, 0, b"\x00" * 3, 123.0))
        + record(0x000A, b"")
    )

    text = readers.read_document(path)
    assert "Revenue" in text and "123" in text


def test_short_pdf_text_is_not_discarded_as_a_scan(tmp_path):
    from pypdf import PdfWriter
    from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

    writer = PdfWriter()
    page = writer.add_blank_page(width=200, height=200)
    font = DictionaryObject({
        NameObject("/Type"): NameObject("/Font"),
        NameObject("/Subtype"): NameObject("/Type1"),
        NameObject("/BaseFont"): NameObject("/Helvetica"),
    })
    page[NameObject("/Resources")] = DictionaryObject({
        NameObject("/Font"): DictionaryObject({NameObject("/F1"): font}),
    })
    content = DecodedStreamObject()
    content.set_data(b"BT /F1 12 Tf 20 100 Td (Hello) Tj ET")
    page[NameObject("/Contents")] = content
    path = tmp_path / "short.pdf"
    with path.open("wb") as file:
        writer.write(file)

    assert readers.read_document(path).strip() == "Hello"
